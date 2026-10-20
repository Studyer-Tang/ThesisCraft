"""Recognize logical Word fields across runs without rewriting their OOXML."""

from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn

from word_formatter.academic.audit import audit
from word_formatter.academic.fields import add_contents, number_captions
from word_formatter.academic.layout import setup_styles
from word_formatter.academic.structure import scan
from word_formatter.academic.templates import load_template
from word_formatter.academic.workflow import run
from word_formatter.academic.xmlutil import element, field_instructions, has_field


def split_field(paragraph, pieces, cached="1"):
    paragraph.add_run()._r.append(element("w:fldChar", fldCharType="begin"))
    for piece in pieces:
        instruction = element("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = piece
        paragraph.add_run()._r.append(instruction)
    paragraph.add_run()._r.append(element("w:fldChar", fldCharType="separate"))
    paragraph.add_run(cached)
    paragraph.add_run()._r.append(element("w:fldChar", fldCharType="end"))


def roundtrip(doc):
    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)
    return Document(stream)


class LogicalFieldTests(unittest.TestCase):
    def setUp(self):
        self.template = load_template("master")

    def test_split_bookmark_and_citation_instructions_are_audited(self):
        doc = Document()
        split_field(doc.add_paragraph(), [" RE", 'F "missing_', 'bookmark" \\h '])
        split_field(doc.add_paragraph(), [" AD", "DIN ZO", "TERO_ITEM CSL_CITATION "])
        doc = roundtrip(doc)
        before = doc.element.xml
        findings, _ = audit(doc, self.template)
        self.assertEqual(
            ["找不到书签：missing_bookmark"],
            [issue["message"] for issue in findings if issue["code"] == "broken-bookmark"],
        )
        self.assertIn("citation-manager", {issue["code"] for issue in findings})
        self.assertEqual(before, doc.element.xml)

    def test_split_existing_bookmark_does_not_report_a_partial_name(self):
        doc = Document()
        paragraph = doc.add_paragraph("引用目标")
        paragraph._p.append(element("w:bookmarkStart", id=1, name="whole_target"))
        paragraph._p.append(element("w:bookmarkEnd", id=1))
        split_field(doc.add_paragraph(), [" REF whole_", "target \\h "])
        findings, _ = audit(roundtrip(doc), self.template)
        self.assertFalse(any(issue["code"] == "broken-bookmark" for issue in findings))

    def test_distinct_fields_are_never_joined_into_one_instruction(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        split_field(paragraph, [" TO"])
        split_field(paragraph, ['C \\o "1-3" '])
        self.assertFalse(has_field(roundtrip(doc), "TOC \\o"))

    def test_orphan_and_unclosed_instructions_are_not_complete_fields(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        orphan = element("w:instrText")
        orphan.text = " REF not_a_field "
        paragraph.add_run()._r.append(orphan)
        paragraph.add_run()._r.append(element("w:fldChar", fldCharType="begin"))
        unclosed = element("w:instrText")
        unclosed.text = ' TOC \\o "1-3" '
        paragraph.add_run()._r.append(unclosed)
        doc = roundtrip(doc)
        self.assertFalse(has_field(doc, "TOC \\o"))
        self.assertFalse(any(i["code"] == "broken-bookmark" for i in audit(doc, self.template)[0]))

    def test_nested_fields_do_not_turn_dynamic_arguments_into_bookmark_names(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run()._r.append(element("w:fldChar", fldCharType="begin"))
        instruction = element("w:instrText")
        instruction.text = " REF "
        paragraph.add_run()._r.append(instruction)
        split_field(paragraph, [' IF 1 = 1 "chosen_target" "other_target" '], "chosen_target")
        tail = element("w:instrText")
        tail.text = " \\h "
        paragraph.add_run()._r.append(tail)
        paragraph.add_run()._r.append(element("w:fldChar", fldCharType="separate"))
        paragraph.add_run("1")
        paragraph.add_run()._r.append(element("w:fldChar", fldCharType="end"))
        # The static inner REF remains auditable even inside another field's code.
        outer = doc.add_paragraph()
        outer.add_run()._r.append(element("w:fldChar", fldCharType="begin"))
        split_field(outer, [" re", "f missing_nested "])
        outer.add_run()._r.append(element("w:fldChar", fldCharType="end"))
        doc = roundtrip(doc)
        before = doc.element.xml
        missing = [i["message"] for i in audit(doc, self.template)[0]
                   if i["code"] == "broken-bookmark"]
        self.assertEqual(["找不到书签：missing_nested"], missing)
        self.assertEqual(before, doc.element.xml)

    def test_prefixed_dynamic_bookmark_argument_is_not_checked_as_a_shorter_name(self):
        for quoted in (False, True):
            with self.subTest(quoted=quoted):
                doc = Document()
                target = doc.add_paragraph("动态引用目标")
                target._p.append(element("w:bookmarkStart", id=1, name="prefix_target"))
                target._p.append(element("w:bookmarkEnd", id=1))
                paragraph = doc.add_paragraph()
                paragraph.add_run()._r.append(element("w:fldChar", fldCharType="begin"))
                prefix = element("w:instrText")
                prefix.text = ' REF "prefix_' if quoted else " REF prefix_"
                paragraph.add_run()._r.append(prefix)
                split_field(paragraph, [' IF 1 = 1 "target" "other" '], "target")
                suffix = element("w:instrText")
                suffix.text = '" \\h ' if quoted else " \\h "
                paragraph.add_run()._r.append(suffix)
                paragraph.add_run()._r.append(element("w:fldChar", fldCharType="separate"))
                paragraph.add_run("1")
                paragraph.add_run()._r.append(element("w:fldChar", fldCharType="end"))
                findings, _ = audit(roundtrip(doc), self.template)
                self.assertFalse(any(i["code"] == "broken-bookmark" for i in findings))

    def test_literal_braces_are_not_truncated_into_a_static_bookmark_argument(self):
        doc = Document()
        for argument in ('"prefix_{}"', "prefix_{}"):
            paragraph = doc.add_paragraph()
            paragraph._p.append(element("w:fldSimple", instr=f"REF {argument} \\h"))
        # Braces cannot name a static Word bookmark. Do not pretend the prefix
        # was the full target; this checker does not evaluate dynamic arguments.
        findings, _ = audit(roundtrip(doc), self.template)
        self.assertFalse(any(i["code"] == "broken-bookmark" for i in findings))

    def test_simple_fields_and_fields_without_cached_results_still_work(self):
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph._p.append(element("w:fldSimple", instr=" pageref missing_simple \\h "))
        paragraph.add_run()._r.append(element("w:fldChar", fldCharType="begin"))
        for text in (" se", "q PSChapter \\h "):
            instruction = element("w:instrText")
            instruction.text = text
            paragraph.add_run()._r.append(instruction)
        paragraph.add_run()._r.append(element("w:fldChar", fldCharType="end"))
        doc = roundtrip(doc)
        self.assertTrue(has_field(doc, "SEQ PSChapter"))
        missing = [i["message"] for i in audit(doc, self.template)[0]
                   if i["code"] == "broken-bookmark"]
        self.assertEqual(["找不到书签：missing_simple"], missing)

    def test_existing_split_toc_is_not_inserted_again(self):
        doc = Document()
        setup_styles(doc, self.template)
        split_field(doc.add_paragraph(), [" TO", 'C \\o "1-3" \\h '], "现有目录")
        doc.add_paragraph("第1章 绪论", "Heading 1")
        doc = roundtrip(doc)
        self.template["contents"].update(toc=True, figures=False, tables=False)
        add_contents(doc, self.template, [])
        self.assertEqual(1, sum(node.get(qn("w:fldCharType")) == "begin"
                                for node in doc.element.iter(qn("w:fldChar"))))
        self.assertEqual([" TO", 'C \\o "1-3" \\h '],
                         [n.text for n in doc.element.iter(qn("w:instrText"))])

    def test_existing_split_chapter_counter_is_not_inserted_again(self):
        doc = Document()
        paragraph = doc.add_paragraph("第1章 绪论", "Heading 1")
        split_field(paragraph, [" SE", "Q PS", "Chapter \\h "], "")
        doc = roundtrip(doc)
        self.template["numbering"]["by_chapter"] = True
        number_captions(doc, scan(doc), self.template, [], [])
        self.assertEqual(1, sum(node.get(qn("w:fldCharType")) == "begin"
                                for node in doc.element.iter(qn("w:fldChar"))))
        self.assertEqual([" SE", "Q PS", "Chapter \\h "],
                         [n.text for n in doc.element.iter(qn("w:instrText"))])

    def test_office_dropped_manager_field_keeps_the_saved_copy(self):
        instruction = ' ADDIN ZOTERO_ITEM CSL_CITATION {"id":"safe"} '
        for kind in ("split", "simple", "duplicate"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "source.docx"
                doc = Document()
                doc.add_paragraph("必须保留的正文。")
                for _ in range(2 if kind == "duplicate" else 1):
                    paragraph = doc.add_paragraph()
                    if kind == "simple":
                        simple = element("w:fldSimple", instr=instruction)
                        result, text = element("w:r"), element("w:t")
                        text.text = "[1]"
                        result.append(text)
                        simple.append(result)
                        paragraph._p.append(simple)
                    else:
                        pieces = [instruction[:3], instruction[3:]] if kind == "split" else [instruction]
                        split_field(paragraph, pieces, "[1]")
                doc.save(source)
                source_bytes = source.read_bytes()

                def drop_citation(stage, host, pdf):
                    staged = Document(stage)
                    for paragraph in staged.paragraphs:
                        if any("ZOTERO_ITEM" in code for code in field_instructions(paragraph._p)):
                            paragraph._p.getparent().remove(paragraph._p)
                            break
                    staged.save(stage)
                    return dict(success=True)

                with patch("word_formatter.academic.office_io.finalize", drop_citation):
                    report = run(source, self.template, host="word")
                self.assertFalse(report["office"]["success"])
                self.assertIn("文献管理器域", report["office"]["error"])
                output = Document(report["output"])
                self.assertEqual(2 if kind == "duplicate" else 1,
                                 sum("ZOTERO_ITEM" in code for code in field_instructions(output.element)))
                self.assertEqual(source_bytes, source.read_bytes())

    def test_office_may_repartition_preserved_manager_instruction_runs(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.docx"
            doc = Document()
            instruction = ' ADDIN ZOTERO_ITEM CSL_CITATION {"id":"safe"} '
            split_field(doc.add_paragraph(), [instruction[:3], instruction[3:]], "[1]")
            doc.save(source)
            source_bytes = source.read_bytes()

            def join_instruction(stage, host, pdf):
                staged = Document(stage)
                nodes = list(staged.element.iter(qn("w:instrText")))
                nodes[0].text = "".join(node.text or "" for node in nodes)
                for node in nodes[1:]:
                    node.getparent().remove(node)
                staged.save(stage)
                return dict(success=True)

            with patch("word_formatter.academic.office_io.finalize", join_instruction):
                report = run(source, self.template, host="word")
            self.assertTrue(report["office"]["success"])
            self.assertEqual([instruction], list(field_instructions(Document(report["output"]).element)))
            self.assertEqual(source_bytes, source.read_bytes())


if __name__ == "__main__":
    unittest.main()
