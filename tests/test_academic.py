"""Thesis behavior and safety tests using real OOXML documents."""

from copy import deepcopy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from word_formatter.academic.templates import (
    load_template,
    validate_template,
    save_template,
)
from word_formatter.academic.workflow import run
from word_formatter.academic.structure import scan
from word_formatter.academic.fields import substitute, field_nodes
from word_formatter.academic.audit import inventory, compare_inventory
from word_formatter.academic.bibliography import load_references


def fixture():
    doc = Document()
    doc.add_paragraph("测试论文")
    doc.add_paragraph("摘要")
    doc.add_paragraph("中文摘要内容。")
    doc.add_paragraph("关键词：排版，测试，论文")
    doc.add_paragraph("Abstract")
    doc.add_paragraph("An abstract.")
    doc.add_paragraph("第1章 绪论", "Heading 1")
    doc.add_paragraph("1.1 研究背景", "Heading 2")
    doc.add_paragraph("正文包含引用 {{cite:demo}} 和图号 {{ref:fig:sample}}。")
    doc.add_paragraph("{{fig:sample}} 研究流程")
    p = doc.add_paragraph()
    math = OxmlElement("m:oMath")
    r = OxmlElement("m:r")
    t = OxmlElement("m:t")
    t.text = "x=1"
    r.append(t)
    math.append(r)
    p._p.append(math)
    table = doc.add_table(rows=4, cols=2)
    for ri, row in enumerate(table.rows):
        for ci, c in enumerate(row.cells):
            c.text = f"数据{ri}-{ci}"
    doc.add_paragraph("第2章 结论", "Heading 1")
    doc.add_paragraph("第二章正文。")
    doc.add_paragraph("参考文献")
    doc.add_paragraph("致谢")
    doc.add_paragraph("感谢指导。")
    doc.add_paragraph("附录 A 补充材料")
    doc.add_paragraph("补充内容。")
    return doc


class AcademicTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.source = self.root / "source.docx"
        fixture().save(self.source)
        self.template = load_template("pku-master")

    def test_presets_roundtrip_and_validation(self):
        for key in ("pku-master", "pku-doctor", "bachelor", "master", "doctor"):
            t = load_template(key)
            save_template(t, self.root / (key + ".json"))
            self.assertEqual(t, load_template(self.root / (key + ".json")))
        self.assertEqual(self.template["page"]["top"], 3.0)
        self.assertEqual(self.template["styles"]["h3"]["size"], 13)
        for path, value in [("schema_version", 2), ("degree", "bad")]:
            with self.assertRaises(ValueError):
                validate_template({path: value})
        broken = deepcopy(self.template)
        broken["styles"]["body"]["size"] = float("nan")
        with self.assertRaises(ValueError):
            validate_template(broken)

    def test_format_source_unchanged_fields_sections_and_formula(self):
        before = self.source.read_bytes()
        result = run(self.source, self.template)
        self.assertEqual(before, self.source.read_bytes())
        self.assertTrue(result["integrity"]["passed"])
        output = Document(result["output"])
        self.assertGreater(len(output.sections), 2)
        self.assertEqual(1, len(list(output.element.iter(qn("m:oMath")))))
        codes = " ".join(n.text or "" for n in output.element.iter(qn("w:instrText")))
        self.assertIn("SEQ PSChapter", codes)
        self.assertIn("SEQ PSfig", codes)
        self.assertIn("REF PS_", codes)
        self.assertIn("TOC", codes)
        self.assertIsNone(result["report"])
        self.assertEqual({self.source, Path(result['output'])}, set(self.root.iterdir()))

    def test_host_intermediates_are_removed_on_success_and_failure(self):
        for success in (True, False):
            def finalize(stage, host, pdf):
                stage.with_suffix('.office.json').write_text('{}', encoding='utf-8')
                stage.with_suffix('.pdf').write_bytes(b'partial temporary file')
                if not success:
                    raise RuntimeError('Office unavailable')
                return {'success': True}
            before = set(self.root.iterdir())
            with patch('word_formatter.academic.office_io.finalize', side_effect=finalize):
                result = run(self.source, self.template, host='word')
            self.assertEqual({Path(result['output'])}, set(self.root.iterdir()) - before)
            self.assertIsNone(result['report'])
            self.assertEqual(not success, bool(result['warnings']))

    def test_format_failure_does_not_leave_reports_or_templates(self):
        with patch('word_formatter.academic.workflow.setup_styles', side_effect=RuntimeError('failed')):
            with self.assertRaisesRegex(RuntimeError, 'failed'):
                run(self.source, self.template)
        self.assertEqual({self.source}, set(self.root.iterdir()))

    def test_native_footnotes_and_endnotes_are_formatted_and_preserved(self):
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        from lxml import etree
        doc = Document(self.source)
        namespace = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        for kind in ('footnote', 'endnote'):
            xml = (f'<w:{kind}s xmlns:w="{namespace}">'
                   f'<w:{kind} w:id="1"><w:p><w:r><w:{kind}Ref/></w:r>'
                   f'<w:r><w:t>保留的{kind}内容</w:t></w:r></w:p></w:{kind}></w:{kind}s>')
            part = Part(PackURI(f'/word/{kind}s.xml'),
                        f'application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml',
                        xml.encode('utf-8'), doc.part.package)
            doc.part.relate_to(part, f'http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}s')
            note_run = doc.paragraphs[6].add_run()
            ref = OxmlElement(f'w:{kind}Reference'); ref.set(qn('w:id'), '1'); note_run._r.append(ref)
        doc.save(self.source)
        result = run(self.source, self.template)
        output = Document(result['output'])
        for kind in ('footnote', 'endnote'):
            part = next(p for p in output.part.package.parts if str(p.partname) == f'/word/{kind}s.xml')
            root = etree.fromstring(part.blob)
            self.assertIn(f'保留的{kind}内容', ''.join(root.itertext()))
            self.assertEqual(['1'], [n.get(qn('w:id')) for n in output.element.iter(qn(f'w:{kind}Reference'))])
            self.assertEqual('PS'+kind, next(root.iter(qn('w:pStyle'))).get(qn('w:val')))

    def test_exact_body_spacing_does_not_clip_images_or_equations(self):
        import base64
        import io
        from docx.enum.text import WD_LINE_SPACING
        png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aX1cAAAAASUVORK5CYII=')
        doc = Document(self.source)
        doc.add_paragraph().add_run().add_picture(io.BytesIO(png))
        doc.save(self.source)
        self.template['styles']['equation'].update(spacing_unit='pt', spacing=20)
        result = run(self.source, self.template)
        output = Document(result['output'])
        for paragraph in output.paragraphs:
            if next(paragraph._p.iter(qn('wp:inline')), None) is not None or next(paragraph._p.iter(qn('m:oMath')), None) is not None:
                self.assertEqual(WD_LINE_SPACING.AT_LEAST, paragraph.paragraph_format.line_spacing_rule)
        body = next(p for p in output.paragraphs if p.text.startswith('第二章正文'))
        self.assertEqual(WD_LINE_SPACING.EXACTLY, body.paragraph_format.line_spacing_rule)

    def test_pdf_is_explicit_and_does_not_overwrite_existing_pdf(self):
        existing_pdf = self.root / 'source_排版.pdf'
        existing_pdf.write_bytes(b'existing PDF')
        def finalize(stage, host, pdf):
            self.assertTrue(pdf)
            pdf_path = stage.with_suffix('.pdf')
            pdf_path.write_bytes(b'test PDF')
            return {'success': True, 'pdf': str(pdf_path)}
        with patch('word_formatter.academic.office_io.finalize', side_effect=finalize), \
                patch('word_formatter.academic.documents.inspect_pdf', return_value={}):
            result = run(self.source, self.template, host='word', pdf=True)
        self.assertEqual(b'existing PDF', existing_pdf.read_bytes())
        self.assertEqual('source_排版_2.docx', Path(result['output']).name)
        self.assertEqual({self.source, existing_pdf, Path(result['output']),
                          Path(result['office']['pdf'])}, set(self.root.iterdir()))

    def test_check_only_does_not_write_docx(self):
        result = run(self.source, self.template, check_only=True)
        self.assertIsNone(result["output"])
        self.assertEqual([], list(Path(result["report"]).parent.glob("*.docx")))

    def test_no_overwrite_repeated_runs(self):
        one = run(self.source, self.template)
        two = run(self.source, self.template)
        self.assertNotEqual(one["output"], two["output"])
        before = Path(one["output"]).read_bytes()
        again = run(one["output"], self.template)
        self.assertEqual(before, Path(one["output"]).read_bytes())
        self.assertTrue(again["integrity"]["passed"])
        original = Document(one["output"])
        updated = Document(again["output"])
        for token in ("SEQ PSfig", "TOC \\o"):
            count = lambda d: sum(
                token in (n.text or "") for n in d.element.iter(qn("w:instrText"))
            )
            self.assertEqual(count(original), count(updated))

    def test_cross_run_markers_preserve_surrounding_emphasis(self):
        d = Document()
        p = d.add_paragraph()
        p.add_run("before ").bold = True
        p.add_run("{{re")
        p.add_run("f:key}} after").italic = True
        self.assertTrue(substitute(p, 7, 18, field_nodes("REF target", "1")))
        self.assertIn("before 1 after", p.text)
        self.assertTrue(p.runs[0].bold)
        self.assertTrue(p.runs[-1].italic)

    def test_structure_override_and_multi_level(self):
        doc = fixture()
        items = scan(doc, {"7": "body"})
        self.assertEqual("body", items[7].kind)
        doc.add_paragraph("2.3.4 三级标题")
        self.assertEqual("h3", scan(doc)[-1].kind)

    def test_references_import_and_numeric_link(self):
        path = self.root / "refs.json"
        path.write_text(
            json.dumps(
                [
                    dict(
                        key="demo",
                        title="A study",
                        authors=["A. Author"],
                        year="2024",
                        type="journal",
                        journal="Testing",
                        volume="2",
                        issue="1",
                        pages="1-8",
                    )
                ]
            ),
            encoding="utf-8",
        )
        result = run(self.source, self.template, reference_path=path)
        doc = Document(result["output"])
        self.assertTrue(any("A study" in p.text for p in doc.paragraphs))
        self.assertFalse(any("{{cite:demo}}" in p.text for p in doc.paragraphs))
        self.assertTrue(
            any(
                "PSBibliography" in (n.text or "")
                for n in doc.element.iter(qn("w:instrText"))
            )
        )
        path.write_text('[{"key":"bad"}]', encoding="utf-8")
        with self.assertRaises(ValueError):
            load_references(path)

    def test_integrity_detects_deleted_formula(self):
        doc = fixture()
        before = inventory(doc)
        math = next(doc.element.iter(qn("m:oMath")))
        math.getparent().remove(math)
        self.assertFalse(compare_inventory(before, inventory(doc))["passed"])

    def test_cancel_before_output(self):
        import threading

        event = threading.Event()
        event.set()
        with self.assertRaises(InterruptedError):
            run(self.source, self.template, cancel=event)
        self.assertEqual([self.source], list(self.root.rglob("*.docx")))

    def test_split_and_merge_copies(self):
        from word_formatter.academic.documents import split_chapters, merge_chapters

        paths = split_chapters(self.source, self.root / "chapters")
        self.assertGreater(len(paths), 2)
        merge_chapters(paths, self.root / "merged.docx")
        merged = Document(self.root / "merged.docx")
        self.assertTrue(any("第二章正文" in p.text for p in merged.paragraphs))
        self.assertEqual(1, len(list(merged.element.iter(qn("m:oMath")))))

    def test_three_line_table_and_repeat_header(self):
        result = run(self.source, self.template)
        table = Document(result["output"]).tables[0]
        self.assertIsNotNone(
            table.rows[0]._tr.find(qn("w:trPr") + "/" + qn("w:tblHeader"))
        )
        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        self.assertEqual("nil", borders.find(qn("w:insideV")).get(qn("w:val")))

    def test_repeated_numbering_keeps_valid_definition_ids(self):
        result = run(self.source, self.template)
        repeated = run(result["output"], self.template)
        doc = Document(repeated["output"])
        valid = {
            n.get(qn("w:numId"))
            for n in doc.part.numbering_part.element.findall(qn("w:num"))
        }
        for p in doc.paragraphs:
            for n in p._p.iter(qn("w:numId")):
                self.assertIn(n.get(qn("w:val")), valid)

    def test_existing_field_and_new_markers_coexist(self):
        from word_formatter.academic.xmlutil import field

        d = Document()
        p = d.add_paragraph("正文 ")
        field(p, 'ADDIN ZOTERO_ITEM CSL_CITATION {"id":"safe"}', "[1]")
        p.add_run(" 新引用 {{ref:key}}。")
        start = p.text.index("{{ref:")
        self.assertTrue(substitute(p, start, start + 11, field_nodes("REF other", "2")))
        self.assertTrue(
            any("ZOTERO_ITEM" in (n.text or "") for n in p._p.iter(qn("w:instrText")))
        )

    def test_numeric_bibliography_regeneration_does_not_duplicate(self):
        refs = self.root / "refs.json"
        records = [
            dict(
                key="demo",
                authors=["Smith, Jane"],
                title="First",
                year="2024",
                type="journal",
                journal="Test",
            )
        ]
        refs.write_text(json.dumps(records), encoding="utf-8")
        first = run(self.source, self.template, reference_path=refs)
        records[0]["title"] = "Revised"
        refs.write_text(json.dumps(records), encoding="utf-8")
        second = run(first["output"], self.template, reference_path=refs)
        paragraphs = Document(second["output"]).paragraphs
        self.assertEqual(1, sum("Revised" in p.text for p in paragraphs))
        self.assertFalse(any("First" in p.text for p in paragraphs))

    def test_inline_formula_stays_in_body(self):
        d = Document()
        p = d.add_paragraph("正文中的公式：")
        math = OxmlElement("m:oMath")
        p._p.append(math)
        self.assertEqual("body", scan(d)[0].kind)

    def test_glossary_and_continuation_table(self):
        csv = self.root / "symbols.csv"
        csv.write_text("symbol,meaning,unit\nx,距离,m\ny,时间,s\n", encoding="utf-8")
        template = deepcopy(self.template)
        template["glossary_path"] = str(csv)
        template["tables"]["rows_per_part"] = 2
        result = run(self.source, template)
        doc = Document(result["output"])
        self.assertTrue(
            any(
                "距离" in c.text
                for table in doc.tables
                for row in table.rows
                for c in row.cells
            )
        )
        self.assertGreater(len(doc.tables), 2)

    def test_formal_front_template_is_filled_without_overwrite(self):
        from word_formatter.academic.layout import create_skeleton

        p = self.root / "front.docx"
        d = Document()
        d.add_paragraph("{{title}}")
        d.add_paragraph("{{author}}")
        d.save(p)
        before = p.read_bytes()
        template = deepcopy(self.template)
        template["front_matter_path"] = str(p)
        template["metadata"].update(title="正式题目", author="Test Author")
        output = self.root / "skeleton.docx"
        create_skeleton(template, output)
        self.assertEqual(before, p.read_bytes())
        self.assertEqual("正式题目", Document(output).paragraphs[0].text)

    def test_unselected_styles_preserve_existing_heading_font(self):
        from docx.shared import Pt

        d = Document(self.source)
        d.styles["Heading 1"].font.size = Pt(23)
        d.save(self.source)
        template = deepcopy(self.template)
        template["enabled"] = []
        r = run(self.source, template)
        self.assertEqual(23, Document(r["output"]).styles["Heading 1"].font.size.pt)


if __name__ == "__main__":
    unittest.main()
