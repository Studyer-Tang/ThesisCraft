import base64
from io import BytesIO
import unittest

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from word_formatter.academic.figures import format_images
from word_formatter.academic.layout import setup_styles
from word_formatter.academic.templates import load_template, validate_template

PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aX1cAAAAASUVORK5CYII=')


def picture(p, width, height):
    return p.add_run().add_picture(BytesIO(PNG), width=Cm(width), height=Cm(height))


class FigureTests(unittest.TestCase):
    def setUp(self):
        self.doc = Document()
        self.template = load_template('pku-master')
        setup_styles(self.doc, self.template)

    def test_portrait_and_landscape_use_their_own_section(self):
        self.template['figures']['max_width_cm'] = 40
        first = self.doc.sections[0]
        first.page_width, first.page_height = Cm(21), Cm(29.7)
        first.left_margin = first.right_margin = Cm(3)
        portrait = picture(self.doc.add_paragraph(), 30, 10)
        second = self.doc.add_section(WD_SECTION_START.NEW_PAGE)
        second.orientation = WD_ORIENT.LANDSCAPE
        second.page_width, second.page_height = Cm(29.7), Cm(21)
        landscape = picture(self.doc.add_paragraph(), 30, 10)
        self.assertEqual(2, format_images(self.doc, self.template))
        self.assertAlmostEqual(15, portrait.width.cm, places=2)
        self.assertAlmostEqual(23.7, landscape.width.cm, places=2)
        self.assertAlmostEqual(3, portrait.width / portrait.height, places=4)

    def test_tall_picture_reserves_caption_space_and_keeps_small_picture(self):
        p = self.doc.add_paragraph()
        tall = picture(p, 10, 50)
        p.paragraph_format.line_spacing = Pt(20)
        self.doc.add_paragraph('图 1.1 高图', 'PS caption_figure')
        english = self.doc.add_paragraph('Figure 1.1 Tall figure', 'PS caption_en')
        self.doc.add_paragraph('正文')
        small = picture(self.doc.add_paragraph(), 1, 1)
        format_images(self.doc, self.template)
        section = self.doc.sections[0]
        self.assertLess(tall.height, section.page_height - section.top_margin - section.bottom_margin - Pt(48))
        self.assertEqual(Cm(1), small.width)
        self.assertEqual(WD_ALIGN_PARAGRAPH.CENTER, p.alignment)
        self.assertEqual(0, p.paragraph_format.first_line_indent)
        self.assertEqual(WD_LINE_SPACING.AT_LEAST, p.paragraph_format.line_spacing_rule)
        self.assertTrue(p.paragraph_format.keep_with_next)
        self.assertFalse(english.paragraph_format.keep_with_next)

    def test_cell_image_stays_inside_cell_and_media_is_identical(self):
        table = self.doc.add_table(rows=1, cols=2)
        table.cell(0, 0).width = Cm(4)
        shape = picture(table.cell(0, 0).paragraphs[0], 12, 6)
        format_images(self.doc, self.template)
        self.assertLess(shape.width, Cm(4))
        self.assertAlmostEqual(shape.width / shape.height, 2, places=4)
        self.assertEqual([PNG], [p.blob for p in self.doc.part.package.parts if '/media/' in str(p.partname)])

    def test_side_by_side_images_fit_together_and_second_pass_is_stable(self):
        p = self.doc.add_paragraph()
        a, b = picture(p, 10, 5), picture(p, 10, 5)
        format_images(self.doc, self.template)
        self.assertLessEqual(a.width + b.width, Cm(14.5) + 1)
        before = self.doc.element.xml
        self.assertEqual(0, format_images(self.doc, self.template))
        self.assertEqual(before, self.doc.element.xml)

    def test_inline_text_and_floating_layout_not_repositioned(self):
        p = self.doc.add_paragraph('行内图标与正文 ')
        picture(p, 1, 1)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = Cm(1)
        floating = picture(self.doc.add_paragraph(), 5, 3)
        floating._inline.tag = qn('wp:anchor')
        before = floating._inline.xml
        format_images(self.doc, self.template)
        self.assertEqual(WD_ALIGN_PARAGRAPH.RIGHT, p.alignment)
        self.assertAlmostEqual(1, p.paragraph_format.first_line_indent.cm, places=3)
        self.assertEqual(before, floating._inline.xml)

    def test_unrelated_paragraph_not_glued_to_caption_and_options_validate(self):
        p = self.doc.add_paragraph()
        picture(p, 3, 2)
        self.doc.add_paragraph('正文不是题注')
        self.doc.add_paragraph('图 1.1 远处题注', 'PS caption_figure')
        format_images(self.doc, self.template)
        self.assertIsNone(p.paragraph_format.keep_with_next)
        for key, value in [('max_height_cm', -1), ('before_pt', 100), ('after_pt', float('nan'))]:
            with self.assertRaises(ValueError):
                validate_template({'figures': {key: value}})


if __name__ == '__main__':
    unittest.main()
