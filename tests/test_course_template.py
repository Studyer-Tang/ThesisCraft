"""Course formats exercise actual DOCX nodes, not paper-specific paragraph offsets."""

from copy import deepcopy
import tempfile
from pathlib import Path
import unittest

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Pt
from lxml import etree

from word_formatter.academic import catalog
from word_formatter.academic.audit import audit
from word_formatter.academic.course import heading_run_sizes
from word_formatter.academic.layout import apply_style, format_tables, setup_styles
from word_formatter.academic.structure import scan
from word_formatter.academic.templates import base_template, load_template, validate_template
from word_formatter.academic.workflow import run
from word_formatter.academic.xmlutil import element, field


class CourseTemplateTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.template = load_template('pku-marxism-course')

    def test_reusable_profile_has_no_paper_offsets_or_identity(self):
        template = self.template
        self.assertEqual(template['document_type'], 'course')
        self.assertEqual(template['structure_overrides'], {})
        self.assertEqual(template['metadata']['author'], '')
        self.assertEqual(template['department'], '')
        self.assertEqual(template['reference_library'], '')
        self.assertFalse(any(template['cleanup'].values()))
        self.assertEqual(template['notes']['number_format'], 'decimalEnclosedCircleChinese')
        self.assertEqual(template['notes']['restart'], 'eachPage')
        self.assertEqual(template['notes']['start'], 1)

    def test_compact_draft_has_inline_labels_identity_and_page_field(self):
        path = catalog.new_document(catalog.get_profile('pku-marxism-course'),
                                    self.root / 'draft.docx')
        doc = Document(path)
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('摘要：', text)
        self.assertIn('关键词：', text)
        self.assertIn('授课教师：________', text)
        self.assertIn('院系：________', text)
        self.assertNotIn('目录', text)
        self.assertNotIn('当证明', text)
        self.assertEqual(len(doc.sections), 1)
        self.assertEqual(doc.styles['PS subtitle'].font.italic, False)
        blocks = scan(doc, course=self.template['course'])
        self.assertEqual([b.kind for b in blocks[:6]],
                         ['title', 'subtitle', 'cover_info', 'cover_info', 'abstract', 'keywords'])
        self.assertNotIn('missing-section', [f['code'] for f in audit(doc, self.template)[0]])
        self.assertEqual([b.kind for b in blocks[-2:]], ['bibliography_heading', 'bibliography'])
        self.assertEqual([n.text.strip() for n in doc.sections[0].footer._element.iter(qn('w:instrText'))], ['PAGE'])
        for label in ('摘要：', '关键词：'):
            p = next(p for p in doc.paragraphs if p.text.startswith(label))
            self.assertTrue(p.runs[0].bold)
        self.assertAlmostEqual(doc.sections[0].left_margin.cm, 2.54, places=2)

    def test_split_and_unsplit_mixed_heading_are_audited_correctly(self):
        doc = Document()
        setup_styles(doc, self.template, initialize=True)
        p = doc.add_paragraph('一、问题：解释为何仍有必要', 'Heading 1')
        p.runs[0].font.size = Pt(14)
        self.assertEqual(heading_run_sizes(p, self.template)[p.runs[0]._r], [14, 12])
        self.assertIn('font-size', [f['code'] for f in audit(doc, self.template)[0]])
        before = p.text
        apply_style(p, 'h1', self.template)
        self.assertEqual(p.text, before)
        self.assertEqual([(r.text, r.font.size.pt) for r in p.runs],
                         [('一、问题：', 14), ('解释为何仍有必要', 12)])
        self.assertNotIn('font-size', [f['code'] for f in audit(doc, self.template)[0]])
        xml = p._p.xml
        apply_style(p, 'h1', self.template)
        self.assertEqual(p._p.xml, xml)

    def test_existing_runs_no_subtitle_and_configurable_separator(self):
        doc = Document()
        setup_styles(doc, self.template, initialize=True)
        p = doc.add_paragraph(style='Heading 1')
        for text in ('二、理论', '：', '概念的边界'):
            p.add_run(text)
        apply_style(p, 'h1', self.template)
        self.assertEqual([r.font.size.pt for r in p.runs], [14, 14, 12])
        for title in ('三、结论', '四、材料:', '五、时间:一个符号', '：没有主标题'):
            p = doc.add_paragraph(title, 'Heading 1')
            apply_style(p, 'h1', self.template)
            self.assertTrue(all(r.font.size.pt == 14 for r in p.runs))
        self.template['course']['heading_subtitle_separator'] = '——'
        p = doc.add_paragraph('六、路径——条件', 'Heading 1')
        apply_style(p, 'h1', self.template)
        self.assertEqual([r.font.size.pt for r in p.runs], [14, 12])
        self.template['course']['heading_subtitle_separator'] = ''
        p = doc.add_paragraph('七、问题：普通标题中的冒号', 'Heading 1')
        apply_style(p, 'h1', self.template)
        self.assertEqual(len(p.runs), 1)
        self.assertEqual(p.runs[0].font.size.pt, 14)

    def test_protected_heading_field_and_link_nodes_are_not_split(self):
        doc = Document()
        setup_styles(doc, self.template, initialize=True)
        p = doc.add_paragraph('一、问题：', 'Heading 1')
        field(p, 'REF kept_bookmark', '保留的域')
        link = element('w:hyperlink')
        r = element('w:r')
        t = element('w:t')
        t.text = '保留的链接'
        r.append(t)
        link.append(r)
        p._p.append(link)
        original_link, original_fields = etree.tostring(link), list(p._p.iter(qn('w:instrText')))
        apply_style(p, 'h1', self.template)
        self.assertEqual(etree.tostring(link), original_link)
        self.assertEqual(list(p._p.iter(qn('w:instrText'))), original_fields)
        self.assertEqual(heading_run_sizes(p, self.template), {})

    def test_mixed_heading_keeps_note_anchors_even_inside_a_text_run(self):
        for combined in (False, True):
            with self.subTest(combined=combined):
                doc = Document()
                setup_styles(doc, self.template, initialize=True)
                p = doc.add_paragraph(style='Heading 1')
                if combined:
                    run = p.add_run('一、问题：副标题')
                else:
                    p.add_run('一、问题：').font.size = Pt(14)
                    run = p.add_run('副标题')
                    run.font.size = Pt(12)
                run._r.append(element('w:footnoteReference', id=12))
                run._r.append(element('w:endnoteReference', id=3))
                apply_style(p, 'h1', self.template)
                self.assertEqual(p.text, '一、问题：副标题')
                self.assertEqual([r.font.size.pt for r in p.runs if r.text], [14, 12])
                self.assertEqual([n.get(qn('w:id')) for n in p._p.iter(qn('w:footnoteReference'))], ['12'])
                self.assertEqual([n.get(qn('w:id')) for n in p._p.iter(qn('w:endnoteReference'))], ['3'])
                self.assertNotIn('font-size', [f['code'] for f in audit(doc, self.template)[0]])

    def test_numbered_bibliography_is_not_a_course_heading(self):
        doc = Document()
        doc.add_paragraph('一个题目')
        doc.add_paragraph('一、引言')
        doc.add_paragraph('参考文献')
        doc.add_paragraph('1. Smith, A. A Book. 2020.')
        doc.add_paragraph('2. Another source')
        items = scan(doc, course=self.template['course'])
        self.assertEqual([b.kind for b in items[-3:]],
                         ['bibliography_heading', 'bibliography', 'bibliography'])
        self.assertEqual([b.region for b in items[-3:]], ['bibliography'] * 3)

    def test_inline_abstract_and_keyword_labels_preserve_note_anchors(self):
        doc = Document()
        setup_styles(doc, self.template, initialize=True)
        for key, label in (('abstract', '摘要：'), ('keywords', '关键词：')):
            for combined in (False, True):
                with self.subTest(key=key, combined=combined):
                    p = doc.add_paragraph()
                    text = p.add_run(label + '内容。')
                    text.font.size = Pt(9)
                    target = text if combined else p.add_run()
                    target._r.append(element('w:footnoteReference', id=9))
                    apply_style(p, key, self.template)
                    text_runs = [r for r in p.runs if r.text]
                    self.assertEqual([(r.text, r.font.size.pt) for r in text_runs],
                                     [(label, 12), ('内容。', 12)])
                    self.assertTrue(text_runs[0].bold)
                    self.assertFalse(text_runs[1].bold)
                    self.assertEqual([n.get(qn('w:id')) for n in p._p.iter(qn('w:footnoteReference'))], ['9'])

    def test_plain_course_structure_is_opt_in_and_does_not_use_offsets(self):
        doc = Document()
        for text in ('一个新题目', '姓名：', '摘要：摘要内容。', '关键词：概念；实践',
                     '一、引言：问题的出现', '正文论证。', '（一）概念', '1. 范围',
                     '（1）条件', '来源：作者归纳。', '参考文献', '一条文献'):
            doc.add_paragraph(text)
        self.assertEqual([b.kind for b in scan(doc, course=self.template['course'])],
                         ['title', 'cover_info', 'abstract', 'keywords', 'h1', 'body',
                          'h2', 'h3', 'h4', 'note', 'bibliography_heading', 'bibliography'])
        self.assertNotEqual(scan(doc)[4].kind, 'h1')
        path = self.root / 'plain.docx'
        doc.save(path)
        before = path.read_bytes()
        result = run(path, self.template)
        output = Document(result['output'])
        self.assertEqual(path.read_bytes(), before)
        self.assertTrue(result['integrity']['passed'])
        self.assertEqual([p.text for p in output.paragraphs], [p.text for p in doc.paragraphs])
        self.assertTrue(output.paragraphs[2].runs[0].bold)
        self.assertTrue(output.paragraphs[3].runs[0].bold)
        self.assertFalse(output.paragraphs[2].runs[-1].bold)
        self.assertNotIn('font-size', [f['code'] for f in result['after_issues']])

    def test_native_notes_and_section_overrides_preserve_anchors_and_links(self):
        doc = Document()
        doc.add_paragraph('一个题目')
        p = doc.add_heading('一、引言', 1)
        ref_run = p.add_run()
        ref_run._r.append(element('w:footnoteReference', id=1))
        doc.add_paragraph('正文内容。')
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        for section in doc.sections:
            props = element('w:footnotePr')
            props.append(element('w:numFmt', val='decimal'))
            props.append(element('w:numStart', val=7))
            props.append(element('w:numRestart', val='continuous'))
            section._sectPr.append(props)
        doc.add_heading('参考文献', 1)
        doc.add_paragraph('作者：《书名》，第1页。')
        ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        relns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
        content = (f'<w:footnotes xmlns:w="{ns}" xmlns:r="{relns}">'
                   '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
                   '<w:footnote w:id="1"><w:p><w:r><w:footnoteRef/></w:r>'
                   '<w:r><w:t>原有引文，第1页。</w:t></w:r>'
                   '<w:hyperlink r:id="rId1"><w:r><w:t>原文</w:t></w:r></w:hyperlink>'
                   '</w:p></w:footnote></w:footnotes>')
        part = Part(PackURI('/word/footnotes.xml'),
                    'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml',
                    content.encode(), doc.part.package)
        url = 'https://example.org/source?v=1'
        part.relate_to(url, relns + '/hyperlink', is_external=True)
        doc.part.relate_to(part, relns + '/footnotes')
        path = self.root / 'notes.docx'
        doc.save(path)
        result = run(path, self.template)
        output = Document(result['output'])
        for props in [output.settings.element.find(qn('w:footnotePr')),
                      *(s._sectPr.find(qn('w:footnotePr')) for s in output.sections)]:
            self.assertEqual(props.find(qn('w:numFmt')).get(qn('w:val')), 'decimalEnclosedCircleChinese')
            self.assertEqual(props.find(qn('w:numStart')).get(qn('w:val')), '1')
            self.assertEqual(props.find(qn('w:numRestart')).get(qn('w:val')), 'eachPage')
        part = next(p for p in output.part.package.parts if str(p.partname) == '/word/footnotes.xml')
        tree = etree.fromstring(part.blob)
        self.assertEqual(tree.xpath('//w:footnote/@w:id', namespaces={'w': ns}), ['-1', '1'])
        self.assertEqual(tree.xpath('//w:footnote[2]//w:t/text()', namespaces={'w': ns}), ['原有引文，第1页。', '原文'])
        self.assertEqual(tree.xpath('//w:footnote[2]//w:sz/@w:val', namespaces={'w': ns}), ['18', '18'])
        self.assertEqual(len(tree.xpath('//w:footnote[2]//w:keepLines', namespaces={'w': ns})), 1)
        self.assertEqual(next(iter(part.rels.values())).target_ref, url)
        self.assertEqual([n.get(qn('w:id')) for n in output.element.iter(qn('w:footnoteReference'))], ['1'])
        self.assertTrue(output.paragraphs[1].paragraph_format.keep_together)

    def test_new_settings_validate_without_changing_other_profiles(self):
        template = base_template()
        self.assertFalse(template['course']['compact_front_matter'])
        self.assertEqual(template['course']['heading_subtitle_separator'], '')
        self.assertFalse(template['notes']['keep_together'])
        for section, key, value in [('notes', 'start', 0), ('notes', 'start', True),
                                    ('course', 'heading_subtitle_size', 100),
                                    ('course', 'heading_subtitle_separator', 'x' * 9)]:
            invalid = deepcopy(template)
            invalid[section][key] = value
            with self.assertRaises(ValueError):
                validate_template(invalid)

    def test_table_caption_stays_with_table_without_joining_body_text(self):
        doc = Document()
        setup_styles(doc, self.template, initialize=True)
        body = doc.add_paragraph('正文论证。')
        caption = doc.add_paragraph('表 1 比较', 'PS caption_table')
        blank = doc.add_paragraph()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '表头'
        table.cell(1, 1).text = '材料'
        format_tables(doc, self.template)
        self.assertTrue(caption.paragraph_format.keep_with_next)
        self.assertTrue(blank.paragraph_format.keep_with_next)
        self.assertFalse(bool(body.paragraph_format.keep_with_next))
        self.assertEqual(table.cell(1, 1).text, '材料')
