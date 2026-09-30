"""Desktop routing, input locking and native scrolling regressions."""

import os
from pathlib import Path
import sys
import tempfile
import time
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from word_formatter.entry import main


class DesktopEntryTests(unittest.TestCase):
    def test_default_opens_thesis_and_forwards_source(self):
        with patch("word_formatter.academic.gui.main") as launch:
            main(["论文.docx"])
            launch.assert_called_once_with(["论文.docx"])

    def test_explicit_general_mode_preserves_legacy_desktop(self):
        with patch("word_formatter.gui.main") as launch:
            main(["--general"])
            launch.assert_called_once_with()

    def test_trackpad_and_mouse_deltas(self):
        from word_formatter.ui_common import scroll_units

        for platform, delta, expected in [
            ("darwin", 2, -2),
            ("win32", 120, -1),
            ("win32", -240, 2),
            ("win32", 30, -1),
        ]:
            with (
                self.subTest(platform=platform, delta=delta),
                patch("sys.platform", platform),
            ):
                self.assertEqual(scroll_units(SimpleNamespace(delta=delta)), expected)
        self.assertEqual(scroll_units(SimpleNamespace(num=4)), -1)
        self.assertEqual(scroll_units(SimpleNamespace(num=5)), 1)


@unittest.skipUnless(
    sys.platform in ("win32", "darwin") or os.environ.get("DISPLAY"), "Display required"
)
class DesktopWindowTests(unittest.TestCase):
    def test_simple_view_roundtrip_and_input_states(self):
        import tkinter as tk
        from word_formatter.academic.gui import AcademicWindow
        from word_formatter.academic.templates import load_template

        with (
            tempfile.TemporaryDirectory() as folder,
            patch(
                "word_formatter.academic.templates.template_path",
                return_value=Path(folder) / "default.json",
            ),
        ):
            root = tk.Tk()
            root.withdraw()
            try:
                window = AcademicWindow(root)
                self.assertEqual(window.collect(), load_template("master"))
                self.assertEqual(window.overview.winfo_manager(), "pack")
                self.assertFalse(window.notebook.winfo_manager())
                window.toggle_advanced()
                self.assertEqual(window.notebook.winfo_manager(), "pack")
                window.lock_inputs(True)
                widgets = dict(window.busy_states)
                self.assertTrue(widgets)
                self.assertTrue(all("disabled" in widget.state() for widget in widgets))
                window.lock_inputs(False)
                for widget, state in widgets.items():
                    self.assertEqual(set(widget.state()), set(state))
                window.toggle_advanced()
                self.assertEqual(window.collect(), load_template("master"))
                if sys.platform != "win32":
                    self.assertEqual(window.host.get(), "稍后手动更新")
                    self.assertFalse(window.pdf.get())
                # Exercise the queue/poll completion path used by real button clicks.
                from docx import Document

                source = Path(folder) / "含空格的论文 test.docx"
                document = Document()
                document.add_heading("第1章 绪论", 1)
                document.add_paragraph("桌面工作流测试。")
                document.save(source)
                original = source.read_bytes()
                window.source.set(str(source))
                for check_only in (True, False):
                    window.report = None
                    window.process(check_only)
                    deadline = time.monotonic() + 10
                    while window.report is None and time.monotonic() < deadline:
                        root.update()
                        time.sleep(0.01)
                    self.assertIsNotNone(window.report)
                    self.assertEqual(source.read_bytes(), original)
                    path = (
                        window.report["report"]
                        if check_only
                        else window.report["output"]
                    )
                    self.assertTrue(Path(path).is_file())
                    self.assertFalse(window.busy_states)
                    self.assertIn("完成", window.status.get())
            finally:
                root.destroy()
