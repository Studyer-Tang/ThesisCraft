"""Audit real inherited OOXML and reject lost text before publishing a copy."""

from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

from word_formatter.academic.audit import audit, inventory, compare_inventory
from word_formatter.academic.content import paragraph_text
from word_formatter.academic.layout import setup_styles, apply_style
from word_formatter.academic.templates import load_template
from word_formatter.academic.workflow import run
from word_formatter.academic.xmlutil import element, field, set_child


class FormattingAuditTests(unittest.TestCase):
    def setUp(self):
        self.template = load_template("master")
        self.doc = Document()
        setup_styles(self.doc, self.template)
        self.doc.add_paragraph("第1章 绪论", "Heading 1")
        self.paragraph = self.doc.add_paragraph(style="PS body")

    def codes(self):
        return {f["code"] for f in audit(self.doc, self.template)[0] if f["index"] == 1}

    def test_later_run_size_latin_font_and_line_spacing_are_reported(self):
        self.paragraph.add_run("正常文字 English").font.size = Pt(12)
        later = self.paragraph.add_run("后半段 Wrong font")
        later.font.size = Pt(40)
        later.font.name = "Courier New"
        self.paragraph.paragraph_format.line_spacing = 3
        self.assertTrue({"font-size", "font-latin", "line-spacing"} <= self.codes())

    def test_chinese_and_latin_slots_are_checked_only_for_present_scripts(self):
        run = self.paragraph.add_run("中文内容")
        run.font.name = "Courier New"  # Not used for these characters.
        self.assertNotIn("font-latin", self.codes())
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
        self.assertIn("font-eastasia", self.codes())
        run.text = "English only"
        run.font.name = "Times New Roman"
        self.assertNotIn("font-eastasia", self.codes())

    def test_font_alias_and_style_inheritance_do_not_create_false_findings(self):
        base = self.doc.styles.add_style("Inherited body", WD_STYLE_TYPE.PARAGRAPH)
        base.base_style = self.doc.styles["PS body"]
        leaf = self.doc.styles.add_style("Inherited leaf", WD_STYLE_TYPE.PARAGRAPH)
        leaf.base_style = base
        self.paragraph.style = leaf
        self.paragraph.paragraph_format.space_after = Pt(6)
        run = self.paragraph.add_run("中英 English")
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "SimSun")
        self.assertFalse({"font-size", "font-eastasia", "font-latin", "line-spacing"} & self.codes())

    def test_character_style_precedes_paragraph_style_and_direct_override_wins(self):
        base = self.doc.styles.add_style("Char base", WD_STYLE_TYPE.CHARACTER)
        base.font.size = Pt(18)
        base.font.name = "Courier New"
        char = self.doc.styles.add_style("Char leaf", WD_STYLE_TYPE.CHARACTER)
        char.base_style = base
        run = self.paragraph.add_run("English")
        run.style = char
        self.assertTrue({"font-size", "font-latin"} <= self.codes())
        run.font.size, run.font.name = Pt(12), "Times New Roman"
        self.assertFalse({"font-size", "font-latin"} & self.codes())

    def test_document_defaults_are_used_when_styles_do_not_define_values(self):
        style = self.doc.styles.add_style("Default values", WD_STYLE_TYPE.PARAGRAPH)
        self.paragraph.style = style
        self.paragraph.add_run("中英 English")
        defaults = self.doc.styles.element.find(qn("w:docDefaults"))
        rpr = defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        set_child(rpr, "w:sz", val=24)
        set_child(rpr, "w:rFonts", ascii="Times New Roman", hAnsi="Times New Roman", eastAsia="宋体")
        ppr = defaults.find(qn("w:pPrDefault") + "/" + qn("w:pPr"))
        set_child(ppr, "w:spacing", line=360, lineRule="auto")
        self.assertFalse({"font-size", "font-eastasia", "font-latin", "line-spacing"} & self.codes())

    def test_resolved_theme_font_has_priority_over_literal_same_layer(self):
        run = self.paragraph.add_run("English")
        fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:asciiTheme"), "minorAscii")
        self.assertIn("font-latin", self.codes())  # Default theme resolves to Cambria.
        fonts.attrib.pop(qn("w:asciiTheme"))
        self.assertNotIn("font-latin", self.codes())

    def test_east_asian_theme_supplemental_font_uses_inherited_language(self):
        run = self.paragraph.add_run("中文内容")
        props = run._r.get_or_add_rPr()
        fonts = props.get_or_add_rFonts()
        fonts.set(qn("w:eastAsiaTheme"), "minorEastAsia")
        set_child(props, "w:lang", eastAsia="zh-CN")
        theme = next(p for p in self.doc.part.package.parts if str(p.partname).startswith("/word/theme/"))
        root = etree.fromstring(theme.blob)
        minor = root.find(".//" + qn("a:minorFont"))
        for node in minor.findall(qn("a:font")):
            if node.get("script") == "Hans":
                node.set("typeface", "SimSun")
        theme._blob = etree.tostring(root)
        self.assertNotIn("font-eastasia", self.codes())
        props.remove(props.find(qn("w:lang")))
        # An unresolved theme is a manual-review item, never a guessed font error.
        self.assertNotIn("font-eastasia", self.codes())
        self.assertIn("format-review", self.codes())

    def test_hyperlink_display_text_is_audited(self):
        self.paragraph.add_run("正文 ")
        link, run = element("w:hyperlink"), element("w:r")
        props = element("w:rPr")
        props.append(element("w:sz", val=80))
        run.append(props)
        text = element("w:t")
        text.text = "Link"
        run.append(text)
        link.append(run)
        self.paragraph._p.append(link)
        self.assertIn("font-size", self.codes())

    def test_field_math_and_superscript_special_formats_are_not_body_errors(self):
        self.paragraph.add_run("正文")
        _, result = field(self.paragraph, "ADDIN ZOTERO_ITEM CSL_CITATION", "[1]")
        result.font.size, result.font.name = Pt(8), "Arial"
        superscript = self.paragraph.add_run("2")
        superscript.font.size, superscript.font.superscript = Pt(8), True
        math = element("w:r")
        props = element("w:rPr")
        props.append(element("w:oMath"))  # An omitted OnOff value means enabled.
        props.append(element("w:sz", val=80))
        math.append(props)
        text = element("w:t")
        text.text = "x"
        math.append(text)
        self.paragraph._p.append(math)
        self.assertFalse({"font-size", "font-latin"} & self.codes())
        self.assertIn("format-review", self.codes())

    def test_complex_script_size_and_font_are_reviewed_without_guessing_latin_properties(self):
        run = self.paragraph.add_run("Latin-shaped complex script")
        props = run._r.get_or_add_rPr()
        props.append(element("w:cs"))
        props.append(element("w:sz", val=80))
        run.font.name = "Courier New"
        self.assertFalse({"font-size", "font-latin"} & self.codes())
        self.assertIn("format-review", self.codes())

    def test_intrinsic_complex_script_size_does_not_use_latin_font_size(self):
        run = self.paragraph.add_run("مرحبا")
        props = run._r.get_or_add_rPr()
        props.append(element("w:sz", val=20))
        props.append(element("w:szCs", val=24))
        self.assertNotIn("font-size", self.codes())
        self.assertIn("format-review", self.codes())

    def test_field_result_spanning_paragraphs_remains_protected(self):
        self.paragraph.add_run()._r.append(element("w:fldChar", fldCharType="begin"))
        self.paragraph.add_run()._r.append(element("w:fldChar", fldCharType="separate"))
        self.paragraph.add_run("cached first").font.size = Pt(8)
        following = self.doc.add_paragraph("cached second", "PS body")
        following.runs[0].font.size = Pt(8)
        following.add_run()._r.append(element("w:fldChar", fldCharType="end"))
        findings = audit(self.doc, self.template)[0]
        self.assertFalse(any(i["code"] == "font-size" and i["index"] in (1, 2) for i in findings))

    def test_exact_and_minimum_spacing_rules_and_paper_size(self):
        self.paragraph.add_run("正文")
        self.template["styles"]["body"].update(spacing_unit="pt", spacing=20)
        apply_style(self.paragraph, "body", self.template)
        self.assertNotIn("line-spacing", self.codes())
        self.paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        self.assertIn("line-spacing", self.codes())
        math = element("m:oMath")
        self.paragraph._p.append(math)
        self.assertNotIn("line-spacing", self.codes())
        findings = audit(self.doc, self.template)[0]
        self.assertIn("paper-size", {i["code"] for i in findings})  # Blank Word file is Letter.
        self.template["page"]["paper_size"] = "preserve"
        self.assertNotIn("paper-size", {i["code"] for i in audit(self.doc, self.template)[0]})


class ContentIntegrityTests(unittest.TestCase):
    def test_deleted_and_same_length_replaced_text_both_fail(self):
        for replacement in ("", "modified"):
            doc = Document()
            paragraph = doc.add_paragraph("original")
            before = inventory(doc)
            paragraph.runs[0].text = replacement
            checked = compare_inventory(before, inventory(doc))
            self.assertFalse(checked["passed"])
            self.assertIn("body_text", checked["lost_categories"])
            self.assertEqual(1, checked["missing_body_blocks"])

    def test_run_splits_are_identical_and_duplicate_paragraph_loss_is_detected(self):
        doc = Document()
        paragraph = doc.add_paragraph("same text")
        doc.add_paragraph("same text")
        before = inventory(doc)
        paragraph.runs[0].text = "same"
        paragraph.add_run(" text")
        self.assertTrue(compare_inventory(before, inventory(doc))["passed"])
        doc.element.body.remove(doc.paragraphs[-1]._p)
        self.assertFalse(compare_inventory(before, inventory(doc))["passed"])

    def test_table_link_and_textbox_content_are_protected(self):
        for kind in ("table", "link", "textbox"):
            doc = Document()
            if kind == "table":
                doc.add_table(rows=1, cols=1).cell(0, 0).text = "unique table value"
            else:
                parent = doc.add_paragraph()._p
                wrapper = element("w:hyperlink" if kind == "link" else "w:txbxContent")
                parent.append(wrapper)
                if kind == "textbox":
                    paragraph = element("w:p")
                    wrapper.append(paragraph)
                    wrapper = paragraph
                run = element("w:r")
                text = element("w:t")
                text.text = "unique display value"
                run.append(text)
                wrapper.append(run)
            before = inventory(doc)
            next(doc.element.iter(qn("w:t"))).text = "corrupt"
            self.assertFalse(compare_inventory(before, inventory(doc))["passed"], kind)

    def test_updateable_cache_may_change_but_neighboring_authored_text_may_not(self):
        doc = Document()
        paragraph = doc.add_paragraph("See ")
        _, cache = field(paragraph, "PAGE", "1")
        paragraph.add_run(" for the result.")
        before = inventory(doc)
        cache.text = "88"
        self.assertTrue(compare_inventory(before, inventory(doc))["passed"])
        paragraph.runs[-1].text = " deleted result."
        self.assertFalse(compare_inventory(before, inventory(doc))["passed"])

    def test_explicit_allowed_rewrite_does_not_allow_unrelated_edit(self):
        doc = Document()
        first = doc.add_paragraph("第1章 绪论")
        second = doc.add_paragraph("Retain this text.")
        before = inventory(doc)
        first.runs[0].text = "绪论"
        changes = [dict(content_before="第1章 绪论", content_after="绪论")]
        doc.add_paragraph("目录")
        self.assertTrue(compare_inventory(before, inventory(doc), changes)["passed"])
        second.runs[0].text = "Unexpected edit."
        self.assertFalse(compare_inventory(before, inventory(doc), changes)["passed"])

    def test_partial_body_run_loss_blocks_workflow_publication(self):
        template = load_template("master")
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "source.docx"
            doc = Document()
            doc.add_heading("第1章 绪论", 1)
            doc.add_paragraph("Protected prose.")
            doc.save(source)
            original = source.read_bytes()

            def corrupt(paragraph, key, template):
                apply_style(paragraph, key, template)
                if key == "body":
                    paragraph.runs[0].text = "Unexpected deletion."

            with patch("word_formatter.academic.workflow.apply_style", side_effect=corrupt):
                with self.assertRaisesRegex(RuntimeError, "body_text"):
                    run(source, template)
            self.assertEqual(original, source.read_bytes())
            self.assertEqual([source], list(Path(folder).iterdir()))

    def test_office_body_loss_keeps_saved_pre_update_copy(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "source.docx"
            doc = Document()
            doc.add_heading("第1章 绪论", 1)
            doc.add_paragraph("Protected prose.")
            doc.save(source)

            def corrupt(stage, host, pdf):
                document = Document(stage)
                paragraph = next(p for p in document.paragraphs if p.text == "Protected prose.")
                paragraph.runs[0].text = "Unexpected edit."
                document.save(stage)
                return {"success": True}

            with patch("word_formatter.academic.office_io.finalize", side_effect=corrupt):
                result = run(source, load_template("master"), host="word")
            self.assertFalse(result["office"]["success"])
            self.assertIn("正文文字", result["office"]["error"])
            self.assertIn("Protected prose.", [p.text for p in Document(result["output"]).paragraphs])

    def test_caption_reference_citation_and_regeneration_allow_only_recorded_edits(self):
        from word_formatter.academic.xmlutil import bookmark

        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "source.docx"
            refs = Path(folder) / "refs.json"
            glossary = Path(folder) / "glossary.csv"
            doc = Document()
            doc.add_heading("第1章 绪论", 1)
            doc.add_paragraph("正文 {{ref:fig:sample}} 和 {{cite:demo}} 保留。")
            doc.add_paragraph("{{fig:sample}} 研究流程")
            doc.add_paragraph("Figure 1 Research flow")
            doc.add_paragraph("{{eq:plain}} x = 1")
            doc.add_paragraph("参考文献")
            # Original, manual bibliography must survive alongside imported entries.
            manual = doc.add_paragraph("Original manual reference.")
            bookmark(manual, "manual-reference")
            doc.save(source)
            refs.write_text('[{"key":"demo","title":"A study","authors":["A. Author"],'
                            '"year":"2024","type":"book","publisher":"Example"}]', encoding="utf-8")
            glossary.write_text("symbol,meaning,unit\nx,source value,m\n", encoding="utf-8")
            template = load_template("master")
            template["glossary_path"] = str(glossary)
            for style in ("gb7714-numeric", "apa"):
                with self.subTest(style=style):
                    current = deepcopy(template)
                    current["references"]["style"] = style
                    first = run(source, current, reference_path=refs)
                    self.assertTrue(first["integrity"]["passed"])
                    self.assertTrue(first["integrity"]["content_checked"])
                    self.assertGreater(first["integrity"]["allowed_body_changes"], 3)
                    self.assertIn("Original manual reference.",
                                  [p.text for p in Document(first["output"]).paragraphs])
                    second = run(first["output"], current, reference_path=refs)
                    self.assertTrue(second["integrity"]["passed"])
                    self.assertTrue(any(i["action"] == "重建本工具生成的符号表条目" for i in second["changes"]))

    def test_paragraph_snapshot_keeps_tabs_and_line_breaks(self):
        from docx.shared import Cm

        doc = Document()
        paragraph = doc.add_paragraph("a\tb\nc")
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(5))
        self.assertEqual("a\tb\nc", paragraph_text(paragraph))

    def test_internal_numbering_corruption_is_not_accepted_as_expected_text(self):
        from word_formatter.academic.layout import number_headings
        from word_formatter.academic.structure import scan

        doc = Document()
        doc.add_heading("第1章 绪论中的真实正文标题", 1)
        template = load_template("master")
        setup_styles(doc, template)
        before = inventory(doc)
        changes = []

        def erase(runs, text):
            for text_run in runs:
                text_run.text = ""

        with patch("word_formatter.academic.layout.replace_text_nodes", side_effect=erase):
            number_headings(doc, scan(doc), template, changes)
        self.assertEqual("绪论中的真实正文标题", changes[0]["content_after"])
        self.assertFalse(compare_inventory(before, inventory(doc), changes)["passed"])

    def test_internal_reference_corruption_cannot_discard_neighboring_prose(self):
        from word_formatter.academic.fields import cross_references, substitute
        from word_formatter.academic.xmlutil import bookmark

        doc = Document()
        doc.add_paragraph("保留前文 {{ref:target}} 保留后文")
        bookmark(doc.add_paragraph("目标"), "target-bookmark")
        before = inventory(doc)
        changes = []

        def corrupt(paragraph, start, end, nodes):
            success = substitute(paragraph, start, end, nodes)
            paragraph.runs[0].text = ""
            return success

        with patch("word_formatter.academic.fields.substitute", side_effect=corrupt):
            cross_references(doc, {"target": "target-bookmark"}, [], changes)
        self.assertEqual("保留前文  保留后文", changes[0]["content_after"])
        self.assertFalse(compare_inventory(before, inventory(doc), changes)["passed"])

    def test_internal_caption_corruption_cannot_discard_caption_title(self):
        from word_formatter.academic.fields import number_captions
        from word_formatter.academic.structure import scan

        doc = Document()
        doc.add_heading("第1章 绪论", 1)
        paragraph = doc.add_paragraph("{{fig:sample}} 真实题注标题")
        template = load_template("master")
        setup_styles(doc, template)
        before = inventory(doc)
        changes = []
        original_add_run = paragraph.__class__.add_run

        def drop_title(p, text=None, style=None):
            if text == "　真实题注标题":
                text = ""
            return original_add_run(p, text, style)

        with patch("docx.text.paragraph.Paragraph.add_run", new=drop_title):
            number_captions(doc, scan(doc), template, changes, [])
        self.assertIn("真实题注标题", changes[-1]["content_after"])
        self.assertFalse(compare_inventory(before, inventory(doc), changes)["passed"])


if __name__ == "__main__":
    unittest.main()
