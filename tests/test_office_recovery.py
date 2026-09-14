"""Office failures must not hide a usable formatted copy or damage the source."""

import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from unittest.mock import Mock, patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from word_formatter.academic.office_io import close_office, finalize
from word_formatter.academic.templates import load_template
from word_formatter.academic.workflow import run
from word_formatter.storage import publish_file


class OfficeRecoveryTests(unittest.TestCase):
    def test_office_dropped_equation_keeps_pre_update_copy(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory)/'formula.docx'
            doc = Document()
            p = doc.add_paragraph('测试正文')
            math = OxmlElement('m:oMath')
            run_node, text_node = OxmlElement('m:r'), OxmlElement('m:t')
            text_node.text = 'x=1'
            run_node.append(text_node)
            math.append(run_node)
            p._p.append(math)
            doc.save(source)

            def drop_formula(stage, host, pdf):
                changed = Document(stage)
                for node in list(changed.element.iter(qn('m:oMath'))):
                    node.getparent().remove(node)
                changed.save(stage)
                return dict(success=True)

            with patch('word_formatter.academic.office_io.finalize', drop_formula):
                result = run(source, load_template('master'), host='wps')
            self.assertFalse(result['office']['success'])
            self.assertTrue(result['warnings'])
            self.assertEqual(len(list(Document(result['output']).element.iter(qn('m:oMath')))), 1)

    def test_close_failure_still_quits_and_preserves_original_error(self):
        doc, app = Mock(), Mock()
        doc.Close.side_effect = RuntimeError('RPC unavailable')
        app.Quit.side_effect = RuntimeError('server stopped')
        report = dict(success=False, error='field update failed')
        close_office(doc, app, report)
        app.Quit.assert_called_once()
        self.assertEqual(report['error'], 'field update failed')
        self.assertEqual(len(report['cleanup_errors']), 2)

    def test_close_failure_invalidates_success(self):
        doc, app = Mock(), Mock()
        doc.Close.side_effect = RuntimeError('busy')
        report = dict(success=True)
        close_office(doc, app, report)
        self.assertFalse(report['success'])
        app.Quit.assert_called_once()

    def test_timeout_returns_actionable_partial_result(self):
        with patch('word_formatter.academic.office_io.subprocess.run',
                   side_effect=subprocess.TimeoutExpired('office', 600)):
            result = finalize('temporary.docx')
        self.assertFalse(result['success'])
        self.assertIn('排版副本已保留', result['error'])

    def test_publish_preserves_bytes_and_no_clobber(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory)/'stage', Path(directory)/'result'
            source.write_bytes(b'updated archive bytes')
            output.write_bytes(b'existing')
            with self.assertRaises(FileExistsError):
                publish_file(source, output)
            self.assertEqual(output.read_bytes(), b'existing')
            publish_file(source, output, overwrite=True)
            self.assertEqual(output.read_bytes(), source.read_bytes())
            self.assertEqual(set(Path(directory).iterdir()), {source, output})

    @unittest.skipUnless(os.name == 'nt', 'Requires a real Windows sharing lock')
    def test_locked_office_stage_keeps_output_and_source(self):
        import ctypes
        from ctypes import wintypes
        kernel = ctypes.WinDLL('kernel32', use_last_error=True)
        kernel.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD,
            wintypes.DWORD, ctypes.c_void_p, wintypes.DWORD, wintypes.DWORD,
            wintypes.HANDLE]
        kernel.CreateFileW.restype = wintypes.HANDLE
        kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        handles, stages = [], []

        def locked_finalize(stage, host, pdf):
            stages.append(stage)
            handle = kernel.CreateFileW(str(stage), 0x80000000, 0, None, 3, 0, None)
            self.assertNotEqual(handle, ctypes.c_void_p(-1).value)
            handles.append(handle)
            return dict(success=False, error='Office timeout')

        try:
            with tempfile.TemporaryDirectory() as directory:
                source = Path(directory)/'原文.docx'
                doc = Document()
                doc.add_paragraph('第一章 绪论')
                doc.add_paragraph('锁定测试正文。')
                doc.save(source)
                original = source.read_bytes()
                with patch('word_formatter.academic.office_io.finalize', locked_finalize):
                    result = run(source, load_template('master'), host='wps')
                self.assertEqual(source.read_bytes(), original)
                self.assertTrue(result['warnings'])
                self.assertIn('锁定测试正文。',
                    [p.text for p in Document(result['output']).paragraphs])
                self.assertEqual(set(Path(directory).iterdir()),
                                 {source, Path(result['output'])})
                self.assertNotEqual(stages[0].parent.parent, source.parent)
        finally:
            for handle in handles:
                kernel.CloseHandle(handle)
            for stage in stages:
                shutil.rmtree(stage.parent)
