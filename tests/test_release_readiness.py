"""Compatibility regressions discovered during release-readiness review."""
import io
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from word_formatter.academic.office_io import finalize
from word_formatter.academic.preflight import font_issues
from word_formatter.academic.structure import scan
from word_formatter.academic.templates import load_template
from word_formatter.academic.workflow import run
from word_formatter.updates import latest_release


class ReadinessTests(unittest.TestCase):
    def test_numbered_prose_and_lists_do_not_become_chapters(self):
        doc = Document()
        doc.add_heading('1. Explicit heading.', 1)
        doc.add_paragraph('1. Short list item', style='List Number')
        doc.add_paragraph('2. This is a complete sentence.')
        doc.add_paragraph('3. ' + '编号正文测试' * 30)
        doc.add_paragraph('1.1 研究方法')
        self.assertEqual([item.kind for item in scan(doc)], ['h1', 'body', 'body', 'body', 'h2'])

    def test_font_aliases_scope_and_unknown_environment(self):
        template = load_template('fudan-master-2026')
        required = template['styles']['body']
        template['styles']['body'].update(font='宋体', latin='Times New Roman')
        self.assertEqual(font_issues(template, ['SimSun', 'Times New Roman']), [])
        issues = font_issues(template, [])
        self.assertIn(required['font'], issues[0]['message'])
        self.assertEqual(font_issues(template)[0]['code'], 'fonts-unchecked')

    def test_report_keeps_fonts_and_source_when_output_is_saved(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / '中文 空格.docx'
            doc = Document()
            doc.add_paragraph('正文测试。')
            doc.save(source)
            original = source.read_bytes()
            result = run(source, load_template('master'), save_report=True, available_fonts=[])
            self.assertTrue(Path(result['report']).is_file())
            self.assertIn('missing-font', [i['code'] for i in result['after_issues']])
            self.assertEqual(source.read_bytes(), original)
            with patch('word_formatter.academic.workflow.write_report', side_effect=PermissionError('locked')):
                result = run(source, load_template('master'), save_report=True)
            self.assertTrue(Path(result['output']).exists())
            self.assertIsNone(result['report'])
            self.assertIn('report-save', [i['code'] for i in result['after_issues']])

    def test_old_office_report_cannot_mask_a_crashed_worker(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'stage.docx'
            old = source.with_suffix('.office.json')
            old.write_text('{"success":true}')
            with patch('word_formatter.academic.office_io.subprocess.run',
                       return_value=subprocess.CompletedProcess([], 1, '', 'worker crashed')):
                with self.assertRaisesRegex(RuntimeError, 'worker crashed'):
                    finalize(source)
            self.assertEqual(json.loads(old.read_text()), {'success': True})

    def test_crashed_worker_cannot_claim_success(self):
        def fail(command, **kwargs):
            Path(command[command.index('--report') + 1]).write_text('{"success":true}')
            return subprocess.CompletedProcess(command, 1, '', '')
        with patch('word_formatter.academic.office_io.subprocess.run', side_effect=fail):
            self.assertFalse(finalize('stage.docx')['success'])

    def test_release_channels_and_download_origin(self):
        records = [dict(tag_name='v4.2.1-beta.1', prerelease=True, html_url='https://evil.invalid'),
                   dict(tag_name='v4.0.0', prerelease=False, html_url='https://evil.invalid')]
        for preview, expected in ((True, 'v4.2.1-beta.1'), (False, 'v4.0.0')):
            with patch('word_formatter.updates.urlopen', return_value=io.BytesIO(json.dumps(records).encode())):
                release = latest_release(preview)
            self.assertEqual(release['tag'], expected)
            self.assertEqual(release['url'], 'https://github.com/Studyer-Tang/ThesisCraft/releases/tag/' + expected)


if __name__ == '__main__':
    unittest.main()
