"""Opt-in source emphasis removal and template reapplication across document parts."""

from copy import deepcopy
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from word_formatter.academic.cleanup import clear_source_emphasis
from word_formatter.academic.templates import load_template, validate_template
from word_formatter.academic.workflow import run


def emphasize(run):
    run.bold = run.italic = run.underline = True
    return run


class EmphasisCleanupTests(unittest.TestCase):
    def test_default_and_individual_switches(self):
        doc = Document()
        r = emphasize(doc.add_paragraph().add_run('保留强调'))
        self.assertEqual(clear_source_emphasis(doc, {}), 0)
        self.assertTrue(r.bold and r.italic and r.underline)
        clear_source_emphasis(doc, {'underline': True})
        self.assertTrue(r.bold and r.italic)
        self.assertFalse(r.underline)

    def test_inherited_format_hyperlinks_fields_tables_headers_and_notes(self):
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        doc = Document()
        style = doc.styles.add_style('Source emphasis', WD_STYLE_TYPE.CHARACTER)
        style.font.bold = style.font.italic = style.font.underline = True
        p = doc.add_paragraph()
        r = p.add_run('继承样式')
        r.style = style
        link = OxmlElement('w:hyperlink')
        link.set(qn('w:anchor'), 'target')
        linked_run = emphasize(p.add_run('链接'))
        link.append(linked_run._r)
        p._p.append(link)
        from word_formatter.academic.xmlutil import field
        field(p, 'ADDIN ZOTERO_ITEM CSL_CITATION test', '引用')
        emphasize(doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run('表格'))
        emphasize(doc.sections[0].header.paragraphs[0].add_run('页眉'))
        for kind in ('footnote', 'endnote', 'comment'):
            part = Part(PackURI('/word/' + kind + 's.xml'),
                        'application/vnd.openxmlformats-officedocument.wordprocessingml.' + kind + 's+xml',
                        (f'<w:{kind}s xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                         f'<w:{kind} w:id="1"><w:p><w:r><w:rPr><w:b/><w:i/><w:u w:val="single"/></w:rPr>'
                         f'<w:t>注释正文</w:t></w:r></w:p></w:{kind}></w:{kind}s>').encode(), doc.part.package)
            doc.part.relate_to(part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/' + kind + 's')
        before_text = ''.join(doc.element.itertext())
        clear_source_emphasis(doc, dict(bold=True, italic=True, underline=True))
        self.assertEqual(before_text, ''.join(doc.element.itertext()))
        self.assertEqual(link.get(qn('w:anchor')), 'target')
        for part in doc.part.package.parts:
            if not str(part.partname).endswith('.xml'):
                continue
            root = etree.fromstring(part.blob)
            for node in root.iter(qn('w:r')):
                props = node.find(qn('w:rPr'))
                self.assertEqual(props.find(qn('w:b')).get(qn('w:val')), '0')
                self.assertEqual(props.find(qn('w:i')).get(qn('w:val')), '0')
                self.assertEqual(props.find(qn('w:u')).get(qn('w:val')), 'none')

    def test_full_workflow_reapplies_heading_table_bold_preserves_math_and_revisions(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory)/'source.docx'
            doc = Document()
            emphasize(doc.add_paragraph('论文标题').runs[0])
            emphasize(doc.add_paragraph('第一章 绪论', 'Heading 1').runs[0])
            emphasize(doc.add_paragraph('正文带有异常格式。').runs[0])
            table = doc.add_table(rows=2, cols=1)
            for row in table.rows:
                emphasize(row.cells[0].paragraphs[0].add_run('表格文字'))
            p = doc.add_paragraph()
            math = OxmlElement('m:oMath')
            mr, mt = OxmlElement('m:r'), OxmlElement('m:t')
            mt.text = 'x=1'
            mr.append(mt)
            math.append(mr)
            p._p.append(math)
            expected_math = etree.tostring(math)
            insertion = OxmlElement('w:ins')
            insertion.set(qn('w:id'), '7')
            insertion.set(qn('w:author'), 'Author')
            insertion.append(emphasize(p.add_run('修订正文'))._r)
            p._p.append(insertion)
            doc.save(source)
            original = source.read_bytes()
            template = load_template('master')
            template['cleanup'] = dict(bold=True, italic=True, underline=True)
            result = run(source, template)
            self.assertTrue(result['integrity']['passed'])
            self.assertEqual(source.read_bytes(), original)
            out = Document(result['output'])
            body = next(p for p in out.paragraphs if '正文带有异常格式' in p.text).runs[0]
            self.assertFalse(body.bold or body.italic or body.underline)
            heading = next(p for p in out.paragraphs if p.style.name == 'Heading 1').runs[0]
            self.assertTrue(heading.bold)
            self.assertFalse(heading.italic or heading.underline)
            self.assertTrue(out.tables[0].cell(0, 0).paragraphs[0].runs[0].bold)
            self.assertFalse(out.tables[0].cell(1, 0).paragraphs[0].runs[0].bold)
            self.assertEqual(etree.tostring(next(out.element.iter(qn('m:oMath')))), expected_math)
            self.assertEqual(next(out.element.iter(qn('w:ins'))).get(qn('w:id')), '7')

    def test_check_only_and_legacy_templates(self):
        template = load_template('master')
        legacy = deepcopy(template)
        del legacy['cleanup']
        self.assertEqual(validate_template(legacy)['cleanup'], dict(bold=False, italic=False, underline=False))
        with self.assertRaises(ValueError):
            validate_template({'cleanup': {'bold': 'yes'}})
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory)/'source.docx'
            doc = Document()
            emphasize(doc.add_paragraph('测试正文').runs[0])
            doc.save(source)
            original = source.read_bytes()
            template['cleanup'] = dict(bold=True, italic=True, underline=True)
            result = run(source, template, check_only=True)
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual(result['changes'], [])
            self.assertIsNone(result['output'])
