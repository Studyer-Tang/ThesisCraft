import base64
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from word_formatter.config import DEFAULT_CONFIG, normalize_config
from word_formatter.cli import load_config
from word_formatter.engine import WordProcessor
from word_formatter.jobs import build_jobs, Job
from word_formatter.service import run_jobs
from word_formatter.conversion import LegacyConversionUnavailable

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aX1cAAAAASUVORK5CYII="
)


class PreservationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def format(self, doc, **config):
        source, target = self.root / "source.docx", self.root / "out.docx"
        doc.save(source)
        before = source.read_bytes()
        engine = WordProcessor(config)
        engine.format_document(source, target)
        self.assertEqual(before, source.read_bytes())
        self.assertEqual([], engine.temp_files)
        return Document(target)

    def test_same_path_cannot_overwrite_original_even_when_enabled(self):
        source = self.root / "source.docx"
        Document().save(source)
        original = source.read_bytes()
        with self.assertRaises(ValueError):
            WordProcessor({}).format_document(source, source, overwrite=True)
        self.assertEqual(original, source.read_bytes())

    def test_existing_output_untouched_without_opt_in(self):
        source, target = self.root / "source.docx", self.root / "out.docx"
        Document().save(source)
        target.write_bytes(b"previous output")
        with self.assertRaises(FileExistsError):
            WordProcessor({}).format_document(source, target)
        self.assertEqual(target.read_bytes(), b"previous output")

    def test_failed_save_preserves_previous_output(self):
        source, target = self.root / "source.docx", self.root / "out.docx"
        Document().save(source)
        target.write_bytes(b"previous output")
        with patch(
            "word_formatter.storage.os.replace", side_effect=OSError("disk error")
        ):
            with self.assertRaises(OSError):
                WordProcessor({}).format_document(source, target, overwrite=True)
        self.assertEqual(target.read_bytes(), b"previous output")
        self.assertFalse(list(self.root.glob(".wfp-*")))

    def test_picture_survives_punctuation(self):
        doc = Document()
        p = doc.add_paragraph("中文,测试")
        p.add_run().add_picture(io.BytesIO(PNG))
        out = self.format(doc, normalize_punctuation=True)
        self.assertEqual(len(out.inline_shapes), 1)
        self.assertEqual(out.paragraphs[0].text, "中文，测试")

    def test_emphasis_survives_length_change(self):
        doc = Document()
        p = doc.add_paragraph("中文...")
        p.add_run("强调").bold = True
        out = self.format(doc, normalize_punctuation=True)
        self.assertEqual(out.paragraphs[0].text, "中文……强调")
        self.assertTrue(
            next(r for r in out.paragraphs[0].runs if r.text == "强调").bold
        )

    def test_inline_heading_preserves_link_bookmark_and_format(self):
        doc = Document()
        p = doc.add_paragraph("（一）标题。正文")
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), "sample")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "链接文字"
        run.append(text)
        link.append(run)
        p._p.append(link)
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), "0")
        start.set(qn("w:name"), "sample")
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), "0")
        p._p.append(start)
        p._p.append(end)
        out = self.format(doc)
        self.assertEqual(out.paragraphs[0].text, p.text)
        self.assertEqual(len(out.paragraphs[0]._p.findall(qn("w:hyperlink"))), 1)
        self.assertEqual(len(out.paragraphs[0]._p.findall(qn("w:bookmarkStart"))), 1)
        self.assertEqual(
            out.paragraphs[0].runs[0].font.size.pt, DEFAULT_CONFIG["h2_size"]
        )

    def test_field_survives_inline_heading(self):
        doc = Document()
        p = doc.add_paragraph("（一）标题。正文")
        run = p.add_run()
        field = OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), "begin")
        run._r.append(field)
        instruction = OxmlElement("w:instrText")
        instruction.text = "PAGE"
        p.add_run()._r.append(instruction)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        p.add_run()._r.append(end)
        out = self.format(doc)
        self.assertEqual(len(list(out.paragraphs[0]._p.iter(qn("w:fldChar")))), 2)

    def test_color_and_pagination_preserved_by_default(self):
        doc = Document()
        p = doc.add_paragraph("正文")
        p.runs[0].font.color.rgb = RGBColor(255, 0, 0)
        p.paragraph_format.page_break_before = True
        p.paragraph_format.keep_with_next = True
        out = self.format(doc)
        p = out.paragraphs[0]
        self.assertEqual(str(p.runs[0].font.color.rgb), "FF0000")
        self.assertTrue(p.paragraph_format.page_break_before)
        self.assertTrue(p.paragraph_format.keep_with_next)

    def test_black_text_is_explicit(self):
        doc = Document()
        doc.add_paragraph("正文").runs[0].font.color.rgb = RGBColor(255, 0, 0)
        out = self.format(doc, force_black_text=True)
        self.assertEqual(str(out.paragraphs[0].runs[0].font.color.rgb), "000000")

    def test_footer_content_preserved_and_number_not_duplicated(self):
        doc = Document()
        doc.add_paragraph("正文")
        doc.sections[0].footer.paragraphs[0].text = "保密说明"
        doc.settings.odd_and_even_pages_header_footer = True
        processor = WordProcessor({"page_number_align": "居中"})
        processor._apply_page_setup(doc)
        processor._apply_page_setup(doc)
        self.assertFalse(doc.settings.odd_and_even_pages_header_footer)
        self.assertIn("保密说明", [p.text for p in doc.sections[0].footer.paragraphs])
        fields = list(doc.sections[0].footer._element.iter(qn("w:instrText")))
        self.assertEqual(len(fields), 1)

    def test_title_inherits_style_alignment(self):
        doc = Document()
        doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("标题", style="Title")
        self.assertEqual(
            WordProcessor({})._find_title_and_subtitle_paragraphs(doc, False), ([0], [])
        )

    def test_default_does_not_accept_revisions_via_com(self):
        doc = Document()
        doc.add_paragraph("正文")
        with patch.object(WordProcessor, "_preprocess_com_tasks") as preprocess:
            self.format(doc)
        preprocess.assert_not_called()


class ConfigurationTests(unittest.TestCase):
    def test_legacy_json_migrated_before_defaults(self):
        config, _ = load_config(
            config_json='{"remove_blank_lines":false,"use_times_new_roman":true}'
        )
        self.assertTrue(config["use_custom_english_font"])
        self.assertIn("保留单个空行", config["blank_line_mode"])

    def test_invalid_values_rejected(self):
        for invalid in (
            {"body_size": -1},
            {"force_a4": "false"},
            {"typo": 1},
            {"body_size": float("nan")},
            {"line_spacing_unit": "px"},
            {"table_width_percent": 101},
        ):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                normalize_config(invalid)


class JobTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def source(self, name):
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("标题\n正文", encoding="utf-8")
        return path

    def test_same_names_have_distinct_outputs(self):
        first, second = self.source("a/same.txt"), self.source("b/same.md")
        jobs = build_jobs([first, second], self.root / "out")
        self.assertNotEqual(jobs[0].output, jobs[1].output)
        result = run_jobs(WordProcessor({}), jobs)
        self.assertEqual(len(result.outputs), 2)
        self.assertEqual(result.exit_code, 0)

    def test_existing_output_is_renamed(self):
        source = self.source("file.txt")
        self.source("file_formatted.docx")
        self.assertEqual(build_jobs([source])[0].output.name, "file_formatted_2.docx")

    def test_single_file_directory_retains_relative_path(self):
        self.source("tree/sub/one.txt")
        job = build_jobs([self.root / "tree"], self.root / "out")[0]
        self.assertEqual(job.output, (self.root / "out/sub/one_formatted.docx").resolve())

    def test_output_subdirectory_excluded(self):
        self.source("tree/one.txt")
        self.source("tree/out/old.docx")
        self.assertEqual(
            len(build_jobs([self.root / "tree"], self.root / "tree/out")), 1
        )

    def test_overlapping_inputs_deduplicated(self):
        source = self.source("tree/one.txt")
        self.assertEqual(len(build_jobs([self.root / "tree", source])), 1)

    def test_skips_return_incomplete_status(self):
        source = self.source("old.wps")
        with patch.object(
            WordProcessor,
            "format_document",
            side_effect=LegacyConversionUnavailable("missing"),
        ):
            result = run_jobs(WordProcessor({}), [Job(source, self.root / "out.docx")])
        self.assertEqual(result.exit_code, 2)

    def test_cancellation_stops_between_files(self):
        from threading import Event

        cancel = Event()
        cancel.set()
        jobs = build_jobs([self.source("one.txt")])
        result = run_jobs(WordProcessor({}), jobs, cancel_event=cancel)
        self.assertTrue(result.cancelled)
        self.assertFalse(jobs[0].output.exists())
