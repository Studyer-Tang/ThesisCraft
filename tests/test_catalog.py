"""Real-document coverage for source-backed presets, export and partial formatting."""

import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from contextlib import redirect_stdout

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt

from word_formatter.academic import catalog
from word_formatter.academic.audit import audit
from word_formatter.academic.layout import paginate
from word_formatter.academic.templates import load_template, validate_template
from word_formatter.academic.workflow import run


class CatalogTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)

    def test_all_profiles_validate_and_new_documents_are_offline(self):
        entries = catalog.profiles()
        self.assertEqual(len(entries), len({p["id"] for p in entries}))
        self.assertEqual({p["category"] for p in entries}, set(catalog.CATEGORIES))
        for school in ("北京大学", "清华大学", "上海交通大学", "复旦大学"):
            self.assertTrue(catalog.profiles(school))
        for profile in entries:
            with self.subTest(profile=profile["id"]):
                template = load_template(profile["id"])
                self.assertEqual(template["catalog_id"], profile["id"])
                self.assertTrue(template["manual_checks"])
                records = catalog.sources(profile)
                self.assertTrue(records)
                public = [s for s in records if s.get("kind") != "user-provided"]
                self.assertEqual(bool(template["sources"]), bool(public))
                for source in records:
                    if source.get("kind") == "user-provided":
                        self.assertEqual(source["distribution"], "metadata-only")
                        self.assertIsNone(source.get("url"))
                        self.assertIsNone(source.get("page"))
                        self.assertTrue(source["summary"])
                        self.assertRegex(source["sha256"], r"^[0-9a-f]{64}$")
                    else:
                        self.assertEqual(source["distribution"], "link-only")
                        self.assertTrue(source["url"].startswith("https://"))
                copy = catalog.new_document(profile, self.root / (profile["id"] + ".docx"))
                self.assertTrue(Document(copy).paragraphs)

    def test_search_returns_independent_profiles_and_exact_category(self):
        self.assertEqual(len(catalog.profiles("Stanford")), 1)
        self.assertEqual(len(catalog.profiles("pku philosophy", "doctor")), 1)
        self.assertEqual(catalog.profiles("not-a-university"), [])
        profile = catalog.profiles("Stanford")[0]
        profile["template"]["name"] = "mutated"
        self.assertNotEqual(
            catalog.profiles("Stanford")[0]["template"]["name"], "mutated"
        )

    def test_export_is_self_contained_and_roundtrips_to_template(self):
        profile = catalog.get_profile("fudan-master-2026")
        path = catalog.export_bundle(profile, self.root / "规范.zip")
        with zipfile.ZipFile(path) as bundle:
            config = json.loads(bundle.read("template.json"))
            self.assertEqual(validate_template(config), load_template(profile["id"]))
            sources = json.loads(bundle.read("sources.json"))
            for source in sources:
                self.assertEqual(source["distribution"], "link-only")
                self.assertTrue(source["url"].startswith("https://"))
            self.assertFalse(any(name.startswith("sources/") for name in bundle.namelist()))
            self.assertIn("NOTICE.md", bundle.namelist())

    def test_community_course_is_searchable_without_public_or_private_paths(self):
        identifier = "pku-marxism-course"
        for query in ("北京大学 思政", "马原", "PKU community"):
            with self.subTest(query=query):
                self.assertIn(identifier, [p["id"] for p in catalog.profiles(query, "course")])
        profile = catalog.get_profile(identifier)
        template = load_template(identifier)
        self.assertIsNone(profile["word_source"])
        self.assertEqual(template["document_type"], "course")
        self.assertEqual(template["sources"], [])
        self.assertEqual(template["structure_overrides"], {})
        self.assertTrue(all(not value for value in template["metadata"].values()))
        self.assertEqual(template["cleanup"], dict(bold=False, italic=False, underline=False))
        self.assertEqual(template["enabled"], ["styles", "pages", "tables", "figures"])
        self.assertEqual(load_template("master")["name"], "硕士通用")
        source, = catalog.sources(profile)
        self.assertEqual(source["kind"], "user-provided")
        self.assertEqual(source["distribution"], "metadata-only")
        self.assertEqual(source["file"], Path(source["file"]).name)
        self.assertNotIn("/", source["file"])
        self.assertNotIn("\\", source["file"])
        self.assertIsNone(source["resolved_url"])
        data = json.dumps(profile, ensure_ascii=False) + json.dumps(source, ensure_ascii=False)
        for private in ("/Users/", "file://"):
            self.assertNotIn(private, data)

    def test_community_course_skeleton_bundle_and_cli_roundtrip(self):
        from word_formatter.cli import main

        identifier = "pku-marxism-course"
        profile = catalog.get_profile(identifier)
        template = load_template(identifier)
        output = catalog.new_document(profile, self.root / "课程起草.docx")
        document = Document(output)
        self.assertTrue(document.paragraphs)
        text = "\n".join(p.text for p in document.paragraphs)
        for thesis_only in ("英文摘要", "原创性声明", "学术成果"):
            self.assertNotIn(thesis_only, text)
        self.assertFalse(any("TOC" in (n.text or "") for n in document.element.xpath("//w:instrText")))
        archive = catalog.export_bundle(profile, self.root / "课程格式.zip")
        with zipfile.ZipFile(archive) as bundle:
            self.assertEqual(set(bundle.namelist()), {"template.json", "sources.json", "NOTICE.md"})
            self.assertEqual(validate_template(json.loads(bundle.read("template.json"))), template)
            self.assertEqual(json.loads(bundle.read("sources.json")), catalog.sources(profile))
            self.assertFalse(any(n.endswith((".pdf", ".docx", ".doc")) for n in bundle.namelist()))

        stream = io.StringIO()
        with redirect_stdout(stream):
            self.assertEqual(main(["thesis", "templates", "--search", "马原", "--category", "course"]), 0)
        self.assertEqual([p["id"] for p in json.loads(stream.getvalue())], [identifier])
        generated = self.root / "cli-course.docx"
        exported = self.root / "cli-course.json"
        packaged = self.root / "cli-course.zip"
        with redirect_stdout(io.StringIO()):
            self.assertEqual(main(["thesis", "new", "--template", identifier, "-o", str(generated)]), 0)
            self.assertEqual(main(["thesis", "template", "--template", identifier, "-o", str(exported)]), 0)
            self.assertEqual(main(["thesis", "template-bundle", "--template", identifier, "-o", str(packaged)]), 0)
        self.assertTrue(Document(generated).paragraphs)
        self.assertEqual(load_template(exported), template)
        with zipfile.ZipFile(packaged) as bundle:
            self.assertEqual(validate_template(json.loads(bundle.read("template.json"))), template)

    def test_local_original_requires_matching_hash_and_never_overwrites_source(self):
        original = self.root / "school.docx"
        doc = Document()
        doc.add_paragraph("Official-style fixture")
        doc.save(original)
        data = original.read_bytes()
        profile = catalog.get_profile("sjtu-master")
        destination = self.root / "copy.docx"
        with self.assertRaises(ValueError):
            catalog.new_document(profile, destination, original_path=original)
        records = catalog._catalog()
        source = records["sources"][profile["word_source"]]
        with patch.dict(source, sha256=hashlib.sha256(data).hexdigest()):
            catalog.new_document(profile, destination, original_path=original)
            self.assertEqual(destination.read_bytes(), data)
            with self.assertRaises(FileExistsError):
                catalog.new_document(profile, destination, original_path=original)
            with self.assertRaises(ValueError):
                catalog.new_document(profile, original, overwrite=True, original_path=original)
        self.assertEqual(original.read_bytes(), data)
        with self.assertRaises(ValueError):
            catalog.new_document(profile, catalog.DATA / "source.docx")

    def test_letter_format_retains_sections_headers_and_unselected_headings(self):
        source = self.root / "input.docx"
        document = Document()
        document.sections[0].header.paragraphs[0].text = "Keep this custom header"
        document.add_heading("Introduction", 1).runs[0].font.size = Pt(19)
        document.add_paragraph("This is ordinary text.").runs[0].font.size = Pt(9)
        document.add_heading("References", 1)
        document.add_paragraph("A source.")
        document.save(source)
        before = source.read_bytes()
        result = run(source, load_template("stanford-doctor"))
        doc = Document(result["output"])
        self.assertEqual(len(doc.sections), 1)
        self.assertEqual(
            doc.sections[0].header.paragraphs[0].text, "Keep this custom header"
        )
        self.assertAlmostEqual(doc.sections[0].page_width.cm, 21.59, places=2)
        self.assertAlmostEqual(doc.sections[0].page_height.cm, 27.94, places=2)
        self.assertAlmostEqual(doc.sections[0].left_margin.cm, 3.81, places=2)
        self.assertEqual(doc.paragraphs[0].runs[0].font.size.pt, 19)
        self.assertEqual(doc.paragraphs[1].runs[0].font.size.pt, 12)
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(
            result["manual_checks"], load_template("stanford-doctor")["manual_checks"]
        )

    def test_unspecified_margins_stay_unchanged_and_are_not_reported_as_errors(self):
        template = load_template("fudan-master-2026")
        doc = Document()
        doc.sections[0].left_margin = Cm(4.1)
        paginate(doc, [], template, [])
        self.assertAlmostEqual(doc.sections[0].left_margin.cm, 4.1, places=2)
        self.assertNotIn("margin", [i["code"] for i in audit(doc, template)[0]])

    def test_landscape_letter_and_explicit_portrait(self):
        template = load_template("stanford-doctor")
        doc = Document()
        doc.sections[0].orientation = WD_ORIENT.LANDSCAPE
        paginate(doc, [], template, [])
        self.assertAlmostEqual(doc.sections[0].page_width.cm, 27.94, places=2)
        template["page"]["preserve_landscape"] = False
        paginate(doc, [], template, [])
        self.assertEqual(doc.sections[0].orientation, WD_ORIENT.PORTRAIT)
        self.assertAlmostEqual(doc.sections[0].page_width.cm, 21.59, places=2)

    def test_english_course_and_thesis_skeletons_are_distinct(self):
        for identifier in ("cambridge-socanth-course", "stanford-doctor"):
            template = load_template(identifier)
            path = catalog.new_document(
                catalog.get_profile(identifier), self.root / (identifier + ".docx")
            )
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertNotIn("中文摘要", text)
            self.assertNotIn("第1章", text)
            self.assertIn("Introduction", text)
            self.assertNotIn("Theme", doc.styles["Heading 1"].element.xml)
            self.assertEqual(str(doc.styles["Heading 1"].font.color.rgb), "000000")
            codes = doc.element.xpath("//w:instrText")
            self.assertEqual(
                any("TOC" in n.text for n in codes),
                template["document_type"] == "thesis",
            )
            self.assertIn("placeholder", [i["code"] for i in audit(doc, template)[0]])


@unittest.skipUnless(
    sys.platform in ("win32", "darwin") or os.environ.get("DISPLAY"), "Display required"
)
class CatalogWindowTests(unittest.TestCase):
    def test_search_empty_state_and_apply_to_existing_paper(self):
        # Keep one Tk root per process, just as the desktop app does. macOS Tk 8
        # can retain Cocoa event state after destroying and recreating Tk roots.
        result = subprocess.run(
            [
                sys.executable,
                "-m",
                "unittest",
                "tests.test_catalog.CatalogWindowTests.exercise_window",
            ],
            cwd=Path(__file__).resolve().parents[1],
            capture_output=True,
            text=True,
            timeout=45,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def exercise_window(self):
        import tkinter as tk
        from word_formatter.academic.catalog_dialog import CatalogDialog
        from word_formatter.academic.gui import AcademicWindow

        with (
            tempfile.TemporaryDirectory() as folder,
            patch(
                "word_formatter.academic.templates.template_path",
                return_value=Path(folder) / "default.json",
            ),
        ):
            root = tk.Tk()
            root.withdraw()
            try:
                app = AcademicWindow(root)
                app.source.set("my paper.docx")
                app.open_catalog()
                dialog = next(
                    w for w in root.winfo_children() if isinstance(w, CatalogDialog)
                )
                dialog.query.set("pku-marxism-course")
                self.assertEqual(len(dialog.rows), 1)
                self.assertIn("disabled", dialog.link_button.state())
                self.assertIn("disabled", dialog.import_button.state())
                self.assertNotIn("disabled", dialog.word_button.state())
                self.assertNotIn("disabled", dialog.bundle_button.state())
                details = dialog.details.get("1.0", "end")
                self.assertIn("用户提供", details)
                self.assertIn("社区参考", details)
                self.assertIn("SHA-256", details)
                self.assertNotIn("None", details)
                with patch("word_formatter.academic.catalog_dialog.webbrowser.open") as browser:
                    dialog.open_source()
                    browser.assert_not_called()
                dialog.apply()
                self.assertEqual(app.source.get(), "my paper.docx")
                self.assertEqual(app.collect()["catalog_id"], "pku-marxism-course")
                app.open_catalog()
                dialog = next(
                    w for w in root.winfo_children() if isinstance(w, CatalogDialog)
                )
                dialog.query.set("no-result-xyz")
                self.assertIsNone(dialog.selected())
                self.assertIn("disabled", dialog.apply_button.state())
                dialog.query.set("Fudan")
                self.assertEqual(len(dialog.rows), 2)
                self.assertNotIn("disabled", dialog.link_button.state())
                dialog.apply()
                self.assertEqual(app.source.get(), "my paper.docx")
                self.assertEqual(app.collect()["style_keys"], ["body"])
                self.assertFalse(app.collect()["page"]["set_margins"])
                self.assertEqual(app.collect()["school"], "复旦大学")
            finally:
                poll_id = app._poll_id
                root.destroy()
                self.assertNotIn(poll_id, root.tk.call("after", "info"))
