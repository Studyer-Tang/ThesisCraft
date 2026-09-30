"""Reproducible synthetic workload. No user documents, third-party data or extra deps."""
import argparse
import json
from pathlib import Path
import platform
import sys
import tempfile
import time

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from word_formatter.academic.workflow import run  # noqa: E402
from word_formatter.academic.templates import load_template  # noqa: E402
from docx import Document  # noqa: E402


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--paragraphs', type=int, default=1000)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix='thesiscraft-benchmark-') as folder:
        source = Path(folder) / 'large.docx'
        document = Document()
        document.add_heading('第1章 性能验证', 1)
        for n in range(args.paragraphs):
            document.add_paragraph(f'{n}. 中文和 English 排版性能测试。' * 4)
        document.save(source)
        start = time.perf_counter()
        result = run(source, load_template('master'), save_report=True)
        print(json.dumps(dict(system=platform.platform(), python=platform.python_version(),
            paragraphs=args.paragraphs, seconds=round(time.perf_counter()-start, 3),
            input_bytes=source.stat().st_size, output_bytes=Path(result['output']).stat().st_size,
            integrity=result['integrity']['passed']), ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
