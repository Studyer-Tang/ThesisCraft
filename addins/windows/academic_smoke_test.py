"""Explicit host integration check using synthetic documents only."""

import argparse
import json
from pathlib import Path
import sys
import time
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('host',choices=['word','wps'])
    args=parser.parse_args()
    import pythoncom
    import win32com.client
    from docx import Document
    from word_formatter.office_toolbar import OfficeToolbar
    from word_formatter.academic.templates import load_template
    root=ROOT/'build'/'academic-validation'
    root.mkdir(exist_ok=True,parents=True)
    source=root/(args.host+'-toolbar.docx')
    d=Document();d.add_paragraph('插件验证论文')
    d.add_paragraph('摘要');d.add_paragraph('验证摘要。')
    d.add_paragraph('Abstract');d.add_paragraph('Test abstract.')
    d.add_paragraph('第1章 绪论','Heading 1');d.add_paragraph('1.1 背景','Heading 2')
    d.add_paragraph('正文。');d.add_paragraph('参考文献');d.add_paragraph('致谢')
    d.save(source);before=source.read_bytes()
    pythoncom.CoInitialize();app=None;toolbar=None;opened=[]
    try:
        app=win32com.client.DispatchEx('Word.Application' if args.host=='word' else 'KWPS.Application')
        app.Visible=False
        opened.append(app.Documents.Open(str(source),ReadOnly=True,AddToRecentFiles=False))
        errors=[]
        toolbar=OfficeToolbar(app,args.host,notify=errors.append)
        template=load_template('pku-master')
        with patch('word_formatter.academic.templates.load_template',return_value=template),patch('webbrowser.open',return_value=True):
            toolbar.buttons[1].Execute()
            deadline=time.monotonic()+150
            while not errors and toolbar.last_result is None and time.monotonic()<deadline:
                pythoncom.PumpWaitingMessages();toolbar.poll();time.sleep(.05)
        assert not errors,errors
        assert toolbar.last_result, 'Academic toolbar did not return a result'
        result=toolbar.last_result
        opened.append(app.ActiveDocument)
        assert Path(app.ActiveDocument.FullName)==Path(result['output'])
        assert source.read_bytes()==before
        assert result['office']['success'],result['office']
        summary=dict(host=args.host,success=True,buttons=len(toolbar.buttons),source_unchanged=True,
                     opened_result=True,office=result['office'])
        (root/(args.host+'-toolbar-result.json')).write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
        print(json.dumps(summary,ensure_ascii=True))
    finally:
        if toolbar:toolbar.close()
        for doc in reversed(opened):doc.Close(False)
        if app is not None and app.Documents.Count==0:app.Quit()
        pythoncom.CoUninitialize()


if __name__=='__main__':main()
