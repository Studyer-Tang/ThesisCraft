"""Check real host loading and the add-in backend using synthetic files only."""
import json
import subprocess
import tempfile
from pathlib import Path
import pythoncom
import win32com.client
from docx import Document


def process_exists(name):
    output = subprocess.check_output(['tasklist', '/FI', f'IMAGENAME eq {name}', '/FO', 'CSV'],
        creationflags=subprocess.CREATE_NO_WINDOW)
    return name.lower().encode() in output.lower()


def main():
    results = []
    pythoncom.CoInitialize()
    try:
        for program, executable in [('Word.Application', 'WINWORD.EXE'), ('KWPS.Application', 'wps.exe')]:
            existed = process_exists(executable)
            app = None
            try:
                app = win32com.client.DispatchEx(program)
                with tempfile.TemporaryDirectory(prefix='studytang_addin_') as temp:
                    source = Path(temp)/'插件测试.docx'
                    doc = Document(); doc.add_paragraph('测试标题'); doc.add_paragraph('（一）标题。正文')
                    doc.save(source)
                    active = app.Documents.Open(str(source), ReadOnly=True)
                    try:
                        app.COMAddIns.Update()
                        item = app.COMAddIns.Item('StudyTang.WordFormatter')
                        item.Connect = True
                        api = item.Object
                        status = api.Status()
                    finally:
                        active.Close(SaveChanges=0)
                    original = source.read_bytes()
                    output = Path(api.FormatSavedFile(str(source)))
                    formatted = Document(output)
                    assert output.exists() and source.read_bytes() == original
                    assert formatted.paragraphs[1].text == '（一）标题。正文'
                    opened = app.Documents.Open(str(output), ReadOnly=True)
                    opened.Close(SaveChanges=0)
                    results.append({'host': program, 'version': str(app.Version),
                        'connected': bool(item.Connect), 'status': status, 'format_and_open': True})
            except Exception as exc:
                results.append({'host': program, 'error': str(exc)})
            finally:
                # WPS can reuse an existing application despite DispatchEx.
                # Never quit an application that was present before this check.
                if app is not None and not existed and app.Documents.Count == 0:
                    app.Quit()
                app = None
    finally:
        pythoncom.CoUninitialize()
    target = Path(__file__).parent/'dist/host-smoke-results.json'
    target.parent.mkdir(exist_ok=True)
    target.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(results, ensure_ascii=False, indent=2), flush=True)
    return int(any('error' in result for result in results))


if __name__ == '__main__':
    raise SystemExit(main())
