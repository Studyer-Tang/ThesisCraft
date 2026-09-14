"""Exercise actual host button events, formatting, and opening the new document."""

import io
import json
from pathlib import Path
import tempfile
import time
import pythoncom
import win32com.client
from docx import Document
from word_formatter.office_toolbar import OfficeToolbar

PNG = __import__("base64").b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aX1cAAAAASUVORK5CYII="
)


def main():
    results = []
    pythoncom.CoInitialize()
    try:
        for host, program in [
            ("word", "Word.Application"),
            ("wps", "KWPS.Application"),
        ]:
            app = win32com.client.DispatchEx(program)
            toolbar = None
            active = None
            opened = None
            errors = []
            workspace = tempfile.TemporaryDirectory(prefix="studytang_toolbar_")
            try:
                temp = workspace.name
                if True:
                    source = Path(temp) / "插件验证.docx"
                    doc = Document()
                    doc.add_paragraph("插件验证标题")
                    doc.add_paragraph("（一）标题。正文").add_run().add_picture(
                        io.BytesIO(PNG)
                    )
                    doc.save(source)
                    before = source.read_bytes()
                    active = app.Documents.Open(str(source), ReadOnly=True)
                    toolbar = OfficeToolbar(app, host, notify=errors.append)
                    # Execute goes through the real COM button's event source.
                    toolbar.buttons[0].Execute()
                    deadline = time.monotonic() + 60
                    while (
                        (toolbar.busy or toolbar.last_result is None)
                        and time.monotonic() < deadline
                        and not errors
                    ):
                        pythoncom.PumpWaitingMessages()
                        toolbar.poll()
                        time.sleep(0.05)
                    assert not errors, errors
                    assert toolbar.last_result and toolbar.last_result.exit_code == 0
                    output = toolbar.last_result.outputs[0]
                    opened = app.ActiveDocument
                    assert Path(opened.FullName) == Path(output)
                    assert source.read_bytes() == before
                    assert len(Document(output).inline_shapes) == 1
                    results.append(
                        {
                            "host": host,
                            "version": str(app.Version),
                            "path": str(app.Path),
                            "button_event": True,
                            "format_and_open": True,
                            "original_unchanged": True,
                            "image_preserved": True,
                        }
                    )
                    opened.Close(SaveChanges=0)
                    opened = None
                    active.Close(SaveChanges=0)
                    active = None
                    toolbar.close()
                    toolbar = None
            except Exception as exc:
                results.append({"host": host, "error": str(exc)})
            finally:
                if toolbar:
                    toolbar.close()
                if opened:
                    opened.Close(SaveChanges=0)
                if active:
                    active.Close(SaveChanges=0)
                # DispatchEx created the test instance; only quit it when empty.
                if app.Documents.Count == 0:
                    app.Quit()
                try:
                    workspace.cleanup()
                except OSError:
                    pass
    finally:
        pythoncom.CoUninitialize()
    target = Path(__file__).parent / "dist/toolbar-smoke-results.json"
    target.parent.mkdir(exist_ok=True)
    target.write_text(
        json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(results, ensure_ascii=False, indent=2))
    return int(any("error" in result for result in results))


if __name__ == "__main__":
    raise SystemExit(main())
