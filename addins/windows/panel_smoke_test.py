"""Verify visible document selection, panel actions and cleanup after host exit.

Creates synthetic documents only; never closes a pre-existing document.
"""
import argparse
import json
from pathlib import Path
import sys
import time
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('host', choices=['word', 'wps'])
    args = parser.parse_args()
    import pythoncom
    import win32com.client
    from docx import Document
    from word_formatter.gui import _create_root
    from word_formatter.office_connection import connect_application
    from word_formatter.office_toolbar import OfficePanel

    folder = ROOT / 'build' / 'panel-validation'
    folder.mkdir(parents=True, exist_ok=True)
    pythoncom.CoInitialize()
    instances, opened, panel, root = [], [], None, None
    sources = []
    try:
        for number in range(2):
            source = folder / f'{args.host}-instance-{number}.docx'
            document = Document()
            document.add_paragraph('学研排版插件连接验证')
            document.add_paragraph('仅用于验证连接、排版副本和退出行为。')
            document.save(source)
            sources.append((source, source.read_bytes()))
            app = win32com.client.DispatchEx('Word.Application' if args.host == 'word' else 'KWPS.Application')
            instances.append(app)
            doc = app.Documents.Open(str(source), ReadOnly=True, AddToRecentFiles=False)
            opened.append(doc)
            app.Visible = True
            doc.Activate()
        root, _ = _create_root()
        root.update()
        root.focus_force()
        for _ in range(12):
            root.update()
            pythoncom.PumpWaitingMessages()
            time.sleep(.1)
        selected = connect_application(args.host)
        assert Path(selected.ActiveDocument.FullName) == sources[-1][0], selected.ActiveDocument.FullName
        errors = []
        with patch.object(OfficePanel, 'notify', side_effect=errors.append):
            panel = OfficePanel(root, args.host)
            root.update()
            assert panel.toolbar, panel.status.get()
            assert sources[-1][0].name in panel.status.get(), panel.status.get()
            # Invoke the visible companion panel's generic-format button.
            panel.action_buttons[5].invoke()
            deadline = time.monotonic() + 90
            while panel.toolbar.last_result is None and time.monotonic() < deadline and not errors:
                root.update()
                time.sleep(.05)
            assert not errors, errors
            result = panel.toolbar.last_result
            assert result and result.exit_code == 0, result
            output = panel.toolbar.application.ActiveDocument
            opened.append(output)
            assert Path(output.FullName) == Path(result.outputs[0])
            assert all(path.read_bytes() == before for path, before in sources)
            # Close only our own documents, then the now-empty test instances.
            for doc in reversed(opened):
                doc.Close(False)
            opened.clear()
            for app in instances:
                if app.Documents.Count == 0:
                    app.Quit()
            instances.clear()
            # This used to fail at cp.Unadvise(cookie) with RPC unavailable.
            panel.toolbar.close()
            panel.toolbar.close()
            panel.close()
            panel, root = None, None
        summary = dict(host=args.host, multi_instance_selection=True, panel_format_and_open=True,
                       original_unchanged=True, cleanup_after_host_exit=True)
        (folder / f'{args.host}-result.json').write_text(json.dumps(summary, indent=2), encoding='utf-8')
        print(json.dumps(summary))
    finally:
        if panel:
            panel.close()
        elif root:
            root.destroy()
        for doc in reversed(opened):
            try:
                doc.Close(False)
            except Exception:
                pass
        for app in instances:
            try:
                if app.Documents.Count == 0:
                    app.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()


if __name__ == '__main__':
    main()
