"""Exercise the actual frozen application without touching user documents or settings."""

import argparse
import hashlib
import json
from pathlib import Path
import platform
import tempfile
import traceback

from .version import __version__


def check(gui=True):
    from docx import Document
    from .academic.csl import render
    from .academic.templates import load_template
    from .academic.workflow import run

    with tempfile.TemporaryDirectory(prefix="thesiscraft-check-") as folder:
        source = Path(folder) / "论文 smoke test.docx"
        doc = Document()
        doc.add_heading("第1章 绪论", 1)
        doc.add_paragraph("本地排版验证。Local document formatting.")
        doc.save(source)
        digest = hashlib.sha256(source.read_bytes()).digest()
        result = run(source, load_template("master"))
        if hashlib.sha256(source.read_bytes()).digest() != digest:
            raise RuntimeError("Source document changed")
        if (
            not result["integrity"]["passed"]
            or not Document(result["output"]).paragraphs
        ):
            raise RuntimeError("Output verification failed")
        report = run(source, load_template("master"), check_only=True)
        if not Path(report["report"]).is_file():
            raise RuntimeError("Audit report missing")
        from .academic.catalog import get_profile, new_document, export_bundle

        profile = get_profile("sjtu-master")
        new_document(profile, Path(folder) / "school.docx")
        export_bundle(profile, Path(folder) / "school.zip")
        international = run(source, load_template("stanford-doctor"))
        page = Document(international["output"]).sections[0]
        if abs(page.page_width.cm - 21.59) > 0.01:
            raise RuntimeError("Letter paper setup failed")
        record = dict(
            key="smoke",
            authors=["Study-Tang"],
            year="2026",
            type="book",
            title="ThesisCraft self-check",
        )
        for style in ("gb7714-numeric", "apa", "ieee"):
            texts, cite = render([record], ["smoke"], style)
            if not texts["smoke"] or not cite(["smoke"]):
                raise RuntimeError("CSL rendering failed: " + style)
        if gui:
            import tkinter as tk
            from .academic.gui import AcademicWindow

            root = tk.Tk()
            root.withdraw()
            try:
                window = AcademicWindow(root, str(source))
                root.update_idletasks()
                window.collect()
                window.toggle_advanced()
                window.toggle_advanced()
                window.lock_inputs(True)
                window.lock_inputs(False)
                from .academic.catalog_dialog import CatalogDialog

                dialog = CatalogDialog(root, lambda _: None, lambda _: None)
                dialog.query.set("Stanford")
                root.update_idletasks()
                if len(dialog.rows) != 1:
                    raise RuntimeError("Catalog search failed")
                dialog.destroy()
            finally:
                root.destroy()
    return dict(
        success=True,
        version=__version__,
        system=platform.system(),
        architecture=platform.machine(),
        gui=gui,
        checks=["docx", "source-preservation", "audit", "csl", "university-catalog", "letter-paper"]
        + (["gui"] if gui else []),
    )


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Validate this build and write a JSON report"
    )
    parser.add_argument("report", type=Path)
    parser.add_argument("--no-gui", action="store_true")
    args = parser.parse_args(argv)
    try:
        result = check(gui=not args.no_gui)
    except Exception:
        result = dict(success=False, version=__version__, error=traceback.format_exc())
    args.report.write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return 0 if result["success"] else 1
