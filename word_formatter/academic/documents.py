"""Chapter assembly, split copies, continuation tables and PDF hand-in inspection."""

from copy import deepcopy
from pathlib import Path
import re
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from ..storage import save_document, ensure_distinct_paths
from .structure import scan
from .xmlutil import element, field


def landscape_tables(doc, template):
    if not template["tables"]["landscape_wide"]:
        return
    from docx.shared import Cm

    portrait_width = Cm(
        21
        - template["page"]["left"]
        - template["page"]["right"]
        - template["page"]["gutter"]
    )
    for table in doc.tables:
        width = sum(c.width or 0 for c in table.columns)
        if width <= portrait_width and len(table.columns) < 8:
            continue
        following = table._tbl.getnext()
        current = doc.sections[-1]._sectPr
        node = following
        while node is not None:
            match = node.find(qn("w:pPr") + "/" + qn("w:sectPr"))
            if match is not None:
                current = match
                break
            node = node.getnext()
        size = current.find(qn("w:pgSz"))
        if size is not None and size.get(qn("w:orient")) == "landscape":
            continue
        before = element("w:p")
        before_pr = element("w:pPr")
        before.append(before_pr)
        before_pr.append(deepcopy(current))
        anchor = table._tbl
        prior = anchor.getprevious()
        if (
            prior is not None
            and prior.tag == qn("w:p")
            and Paragraph(prior, doc._body).style.name == "PS caption_table"
        ):
            anchor = prior
        anchor.addprevious(before)
        after = element("w:p")
        after_pr = element("w:pPr")
        after.append(after_pr)
        landscape = deepcopy(current)
        from .xmlutil import set_child

        set_child(
            landscape,
            "w:pgSz",
            w=round(Cm(29.7).twips),
            h=round(Cm(21).twips),
            orient="landscape",
        )
        set_child(landscape, "w:type", val="nextPage")
        after_pr.append(landscape)
        table._tbl.addnext(after)


def merge_chapters(paths, output):
    from docxcompose.composer import Composer

    paths = [Path(p).resolve() for p in paths]
    if not paths:
        raise ValueError("请至少选择一个章节文件。")
    for p in paths:
        ensure_distinct_paths(p, output)
    # Duplicate named bookmarks can change the target of REF; fail instead of guessing.
    seen = set()
    for path in paths:
        doc = Document(path)
        names = {
            n.get(qn("w:name"))
            for n in doc.element.iter(qn("w:bookmarkStart"))
            if not n.get(qn("w:name"), "").startswith("_")
        }
        if seen & names:
            raise ValueError(
                "章节中有重复书签名称，请先在 Word/WPS 中调整："
                + ", ".join(sorted(seen & names))
            )
        seen.update(names)
    composer = Composer(Document(paths[0]))
    for path in paths[1:]:
        composer.doc.add_page_break()
        composer.append(Document(path))
    save_document(composer.doc, output)
    return str(output)


def split_chapters(source, directory):
    source, directory = Path(source), Path(directory)
    doc = Document(source)
    items = scan(doc)
    starts = [doc.paragraphs[i.index]._p for i in items if i.kind in ("h1", "appendix")]
    if not starts:
        raise ValueError("没有找到章标题，请先设置标题样式。")
    nodes = list(doc.element.body)
    positions = [nodes.index(n) for n in starts]
    if positions[0] != 0:
        positions.insert(0, 0)
    ends = positions[1:] + [len(nodes) - 1]
    directory.mkdir(parents=True, exist_ok=True)
    outputs = []
    for i, (start, end) in enumerate(zip(positions, ends)):
        copy = Document(source)
        body = copy.element.body
        for j, node in enumerate(list(body)):
            if node.tag != qn("w:sectPr") and not start <= j < end:
                body.remove(node)
        path = directory / f"章节_{i + 1:02d}.docx"
        ensure_distinct_paths(source, path)
        save_document(copy, path)
        outputs.append(str(path))
    return outputs


def continue_tables(doc, template, changes, warnings):
    per_part = template["tables"].get("rows_per_part", 0)
    if not per_part:
        return
    for table in list(doc.tables):
        rows = list(table._tbl.tr_lst)
        if len(rows) <= per_part + 1:
            continue
        if any(table._tbl.iter(qn("w:vMerge"))):
            warnings.append(
                dict(
                    code="merged-table",
                    index=-1,
                    message="纵向合并单元格表格保留，续表需人工拆分。",
                )
            )
            continue
        preceding = table._tbl.getprevious()
        names = (
            list(preceding.iter(qn("w:bookmarkStart"))) if preceding is not None else []
        )
        bookmark_name = names[0].get(qn("w:name")) if names else None
        tail = table._tbl
        for start in range(per_part + 1, len(rows), per_part):
            caption = Paragraph(element("w:p"), doc._body)
            caption.add_run("续")
            if bookmark_name:
                field(caption, f"REF {bookmark_name} \\h", "表")
            else:
                caption.add_run("表（请核对表号）")
            caption.paragraph_format.page_break_before = True
            caption.paragraph_format.keep_with_next = True
            tail.addnext(caption._p)
            copy = deepcopy(table._tbl)
            for row in list(copy.tr_lst):
                copy.remove(row)
            copy.append(deepcopy(rows[0]))
            for row in rows[start : start + per_part]:
                copy.append(deepcopy(row))
            caption._p.addnext(copy)
            tail = copy
        for row in rows[per_part + 1 :]:
            table._tbl.remove(row)
        changes.append(dict(index=-1, action=f"按每段 {per_part} 行拆分续表并重复表头"))


def inspect_pdf(path):
    from pypdf import PdfReader

    pdf = PdfReader(path)
    missing, blank, errors = set(), [], []
    for index, page in enumerate(pdf.pages):
        text = page.extract_text() or ""
        if not text.strip() and not page.images:
            blank.append(index + 1)
        if re.search("未找到引用源|Error!|Bookmark not defined", text):
            errors.append(index + 1)
        resources = page.get("/Resources")
        if not resources:
            continue
        fonts = resources.get_object().get("/Font", {})
        for ref in fonts.values():
            font = ref.get_object()
            entries = font.get("/DescendantFonts", [font])
            for item in entries:
                f = item.get_object()
                descriptor = f.get("/FontDescriptor")
                if not descriptor or not any(
                    k in descriptor.get_object()
                    for k in ("/FontFile", "/FontFile2", "/FontFile3")
                ):
                    missing.add(str(f.get("/BaseFont", "Unknown")))
    return dict(
        pages=len(pdf.pages),
        bookmarks=bool(pdf.outline),
        fonts_without_embedded_data=sorted(missing),
        possibly_blank_pages=blank,
        pages_with_field_errors=errors,
        note="字体和空白页为待确认项；本检查不等同于 PDF/A 认证。",
    )


def add_glossary(doc, path, template):
    import csv
    from .layout import apply_style
    from .xmlutil import set_child

    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if not {"symbol", "meaning"}.issubset(reader.fieldnames or []):
            raise ValueError("符号表 CSV 需要 symbol、meaning 列，可选 unit 列。")
        rows = list(reader)
    if not rows:
        raise ValueError("符号表为空。")
    for table in list(doc.tables):
        caption = table._tbl.tblPr.find(qn("w:tblCaption"))
        if caption is not None and caption.get(qn("w:val")) == "PaperStudioGlossary":
            table._tbl.getparent().remove(table._tbl)
    heading = next(
        (
            p
            for p in doc.paragraphs
            if p.text.strip() in ("符号表", "主要符号对照表", "缩略语表")
        ),
        None,
    )
    if heading is None:
        body = next((p for p in doc.paragraphs if p.style.name == "Heading 1"), None)
        heading = (
            body.insert_paragraph_before("主要符号对照表")
            if body is not None
            else doc.add_paragraph("主要符号对照表")
        )
    apply_style(heading, "front_heading", template)
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ("符号 / 缩略语", "含义", "单位")):
        cell.text = text
    for row in rows:
        cells = table.add_row().cells
        for cell, key in zip(cells, ("symbol", "meaning", "unit")):
            cell.text = row.get(key, "")
    set_child(table._tbl.tblPr, "w:tblCaption", val="PaperStudioGlossary")
    heading._p.addnext(table._tbl)
    return len(rows)
