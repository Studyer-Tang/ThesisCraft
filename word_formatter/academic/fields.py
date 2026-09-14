"""Updateable captions, equations, tables of contents and bookmark references."""

from copy import deepcopy
import re
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from .xmlutil import (
    element,
    field,
    bookmark,
    bookmark_name,
    plain,
    has_field,
    set_font,
    set_child,
)
from .templates import STYLE_NAMES
from .structure import MARKER, CAPTION
from .layout import apply_style
from ..ooxml import paragraph_runs, replace_text_nodes


def substitute(paragraph, start, end, nodes):
    """Replace an explicit marker across plain runs without flattening adjacent formatting."""
    if paragraph.text != "".join(r.text for r in paragraph.runs):
        return False
    depth, position = 0, 0
    for run in paragraph.runs:
        flags = list(run._r.iter(qn("w:fldChar")))
        for flag in flags:
            if flag.get(qn("w:fldCharType")) == "begin":
                depth += 1
        if min(end, position + len(run.text)) > max(start, position):
            if depth or any(
                child.tag not in (qn("w:rPr"), qn("w:t")) for child in run._r
            ):
                return False
        for flag in flags:
            if flag.get(qn("w:fldCharType")) == "end":
                depth = max(0, depth - 1)
        position += len(run.text)
    position, anchor, tail = 0, None, None
    for run in list(paragraph.runs):
        text = run.text
        left, right = max(0, start - position), min(len(text), end - position)
        if right > left:
            if anchor is None:
                anchor = run._r
                if right < len(text):
                    tail = deepcopy(run._r)
                    from docx.text.run import Run

                    Run(tail, paragraph).text = text[right:]
                run.text = text[:left]
            else:
                run.text = text[:left] + text[right:]
        position += len(text)
    if anchor is None:
        return False
    for node in nodes:
        anchor.addnext(node)
        anchor = node
    if tail is not None:
        anchor.addnext(tail)
    return True


def field_nodes(instruction, cached, spec=None):
    p = Paragraph(element("w:p"), None)
    field(p, instruction, cached, spec)
    return list(p._p)


def number_captions(doc, items, template, changes, warnings):
    chapter, appendix = 0, 0
    counters, targets = {}, {}
    paragraphs = list(doc.paragraphs)
    config = template["numbering"]
    numeric = dict(
        font=config["east_asia"],
        latin=config["font"],
        size=config["size"],
        bold=config["bold"],
    )
    for item in items:
        p = paragraphs[item.index]
        if item.kind == "h1":
            if config["by_chapter"] and not has_field_in_paragraph(p, "SEQ PSChapter"):
                field(p, "SEQ PSChapter \\h", "")
            chapter += 1
            appendix = 0
            if config["by_chapter"]:
                counters.clear()
        elif item.kind == "appendix":
            if config["by_chapter"] and not has_field_in_paragraph(p, "SEQ PSAppendix"):
                field(p, "SEQ PSAppendix \\h", "")
            appendix += 1
            if config["by_chapter"]:
                counters.clear()
        if item.kind not in ("caption_figure", "caption_table", "equation"):
            continue
        token = {"caption_figure": "fig", "caption_table": "table", "equation": "eq"}[
            item.kind
        ]
        if config["match_caption"]:
            numeric = dict(template["styles"][item.kind])
        marker = MARKER.match(item.text)
        caption = CAPTION.match(item.text)
        key = item.key or f"{token}-{item.index + 1}"
        name = bookmark_name(token + ":" + key)
        # Generated captions retain existing fields and bookmarks on subsequent runs.
        existing = [
            n.get(qn("w:name"))
            for n in p._p.iter(qn("w:bookmarkStart"))
            if n.get(qn("w:name"), "").startswith("PS_")
        ]
        if has_field_in_paragraph(p, "SEQ PS"):
            if existing:
                targets[key] = existing[0]
            continue
        is_equation = token == "eq"
        if not is_equation and not plain(p):
            warnings.append(
                dict(
                    code="protected-caption",
                    index=item.index,
                    message="复杂题注已保留；可用插件插入新的题注。",
                )
            )
            continue
        # English caption directly below its Chinese counterpart shares the same number.
        english = bool(
            caption and caption.group(1).lower() in ("figure", "fig.", "table")
        )
        if (
            english
            and item.index > 0
            and items[item.index - 1].kind == item.kind
            and not re.match("[A-Za-z]", items[item.index - 1].text)
        ):
            prior_key = items[item.index - 1].key or f"{token}-{item.index}"
            prior = targets.get(prior_key)
            if prior:
                replace_text_nodes(list(paragraph_runs(p)), "")
                field(p, f"REF {prior} \\h", "", numeric)
                p.add_run(" " + caption.group(3))
                apply_style(p, "caption_en", template)
                continue
        counters[token] = counters.get(token, 0) + 1
        number = counters[token]
        chap = chr(64 + appendix) if appendix else str(chapter)
        if config["by_chapter"] and chapter == 0 and appendix == 0:
            warnings.append(
                dict(
                    code="caption-no-chapter",
                    index=item.index,
                    message="题注前没有章标题，使用全篇编号。",
                )
            )
        chapter_numbered = config["by_chapter"] and (chapter > 0 or appendix > 0)
        if not is_equation:
            body = (
                marker.group(3)
                if marker
                else caption.group(3)
                if caption
                else item.text
            )
            replace_text_nodes(list(paragraph_runs(p)), "")
        else:
            body = ""
            if marker and plain(p):
                replace_text_nodes(list(paragraph_runs(p)), marker.group(3))
            # Inline formulas are not made into separately numbered display formulas.
            if (
                not marker
                and item.confidence != "manual"
                and p.text.strip()
                and not re.fullmatch(r"[\s\d().（）-]*", p.text)
            ):
                warnings.append(
                    dict(
                        code="inline-equation",
                        index=item.index,
                        message="行内公式保留；独立公式可通过结构标记或插件编号。",
                    )
                )
                continue
            if has_field_in_paragraph(p, "SEQ"):
                continue
            width = (
                21
                - template["page"]["left"]
                - template["page"]["right"]
                - template["page"]["gutter"]
            )
            p.paragraph_format.tab_stops.add_tab_stop(
                Cm(width / 2), WD_TAB_ALIGNMENT.CENTER
            )
            p.paragraph_format.tab_stops.add_tab_stop(Cm(width), WD_TAB_ALIGNMENT.RIGHT)
            # Keep the original OMML/OLE element itself unchanged.
            first = p._p.find(qn("w:r"))
            leading = element("w:r")
            leading.append(element("w:tab"))
            if first is None:
                p._p.insert(1 if p._p.pPr is not None else 0, leading)
            else:
                first.addprevious(leading)
            p.add_run("\t")
        label = "" if is_equation else ("表 " if token == "table" else "图 ")
        first_run = p.add_run(config["equation_brackets"][0] if is_equation else label)
        set_font(first_run, numeric)
        if chapter_numbered:
            counter = "PSAppendix" if appendix else "PSChapter"
            field(
                p,
                f"SEQ {counter} \\c \\* " + ("ALPHABETIC" if appendix else "ARABIC"),
                chap,
                numeric,
            )
            set_font(p.add_run(config["separator"]), numeric)
        seq = f"PS{token}" + ("Appendix" if appendix else "")
        reset = " \\s 1" if chapter_numbered else ""
        field(p, f"SEQ {seq}{reset} \\* ARABIC", str(number), numeric)
        if is_equation:
            set_font(p.add_run(config["equation_brackets"][1]), numeric)
        last_run = p._p[-1]
        bookmark(p, name, first_run._r, last_run)
        if body:
            p.add_run("　" + body)
        apply_style(p, item.kind, template)
        targets[key] = name
        # Explicit markers use their user key; stable generated bookmarks make Word REF editable.
        changes.append(
            dict(
                index=item.index,
                action="建立题注/公式编号与引用目标",
                key=key,
                bookmark=name,
            )
        )
        if template["figures"]["keep_caption"]:
            if token == "fig" and item.index:
                prior = paragraphs[item.index - 1]
                if next(prior._p.iter(qn('w:drawing')), None) is not None:
                    prior.paragraph_format.keep_with_next = True
            elif token == "table":
                p.paragraph_format.keep_with_next = True
    return targets


def has_field_in_paragraph(p, code):
    return any(code in (n.text or "") for n in p._p.iter(qn("w:instrText"))) or any(
        code in n.get(qn("w:instr"), "") for n in p._p.iter(qn("w:fldSimple"))
    )


def cross_references(doc, targets, warnings):
    pattern = re.compile(r"\{\{ref:(?:(fig|table|eq):)?([\w.-]+)\}\}")
    names = {n.get(qn("w:name")) for n in doc.element.iter(qn("w:bookmarkStart"))}
    for index, p in enumerate(doc.paragraphs):
        for match in reversed(list(pattern.finditer(p.text))):
            token, key = match.groups()
            name = bookmark_name(token + ":" + key) if token else targets.get(key)
            if not name or name not in names:
                warnings.append(
                    dict(
                        code="missing-reference",
                        index=index,
                        message="找不到引用目标：" + key,
                    )
                )
            elif not substitute(
                p, *match.span(), field_nodes(f"REF {name} \\h", "更新引用")
            ):
                warnings.append(
                    dict(
                        code="protected-reference",
                        index=index,
                        message="引用标记位于复杂段落，请通过插件插入。",
                    )
                )


def add_contents(doc, template, changes):
    contents = template["contents"]
    candidates = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    if not candidates:
        return
    anchor = candidates[0]
    if not any(
        n.get(qn("w:name")) == "PSBodyContents"
        for n in doc.element.iter(qn("w:bookmarkStart"))
    ):
        bookmark(anchor, "PSBodyContents")
        start = next(
            n
            for n in anchor._p.iter(qn("w:bookmarkStart"))
            if n.get(qn("w:name")) == "PSBodyContents"
        )
        end = next(
            n
            for n in anchor._p.iter(qn("w:bookmarkEnd"))
            if n.get(qn("w:id")) == start.get(qn("w:id"))
        )
        doc.paragraphs[-1]._p.append(end)
    fields = [
        ("toc", "目录", f'TOC \\o "1-{contents["depth"]}" \\b PSBodyContents \\h \\z'),
        ("figures", "图目录", 'TOC \\t "PS caption_figure,1" \\h \\z'),
        ("tables", "表目录", 'TOC \\t "PS caption_table,1" \\h \\z'),
    ]
    for key, title, code in fields:
        needle = (
            "TOC \\o"
            if key == "toc"
            else ("PS caption_figure" if key == "figures" else "PS caption_table")
        )
        if not contents[key] or has_field(doc, needle):
            continue
        heading = anchor.insert_paragraph_before(title, STYLE_NAMES["front_heading"])
        heading.paragraph_format.page_break_before = True
        # The contents heading itself must not enter the contents.
        set_child(heading._p.get_or_add_pPr(), "w:outlineLvl", val=9)
        p = anchor.insert_paragraph_before()
        field(p, code, "请点击插件“更新目录与引用”生成页码。")
        changes.append(dict(index=-1, action="插入可更新的" + title))
    set_child(doc.settings.element, "w:updateFields", val="true")
