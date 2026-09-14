"""Shared thesis workflow for desktop, CLI and Office plugins."""

from datetime import datetime
import hashlib
from pathlib import Path
import re
import shutil
import time
import tempfile
import uuid

from docx import Document

from ..jobs import unique_path
from ..storage import save_document, publish_file
from .templates import validate_template, STYLE_NAMES
from .audit import audit, inventory, compare_inventory, write_report
from .layout import (
    setup_styles,
    apply_style,
    number_headings,
    paginate,
    format_tables,
    format_images,
    format_notes,
)
from .fields import number_captions, cross_references, add_contents
from .bibliography import load_references, apply_references


def run(
    source,
    template,
    output_dir=None,
    check_only=False,
    reference_path=None,
    host=None,
    pdf=False,
    cancel=None,
    progress=None,
):
    started = time.monotonic()
    source = Path(source).resolve()
    if source.suffix.lower() != ".docx" or not source.is_file():
        raise ValueError("论文模式需要 DOCX；旧格式请先使用桌面通用模式转换。")
    template = validate_template(template)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    doc = Document(source)
    before = inventory(doc)
    before_issues, items = audit(doc, template)
    parent = Path(output_dir).resolve() if output_dir else source.parent
    safe_stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", source.stem)[:70]
    directory = parent
    if check_only:
        directory = parent / f"{safe_stem}_检查_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}"
        directory.mkdir(parents=True, exist_ok=False)
    else:
        parent.mkdir(parents=True, exist_ok=True)
    report = dict(
        source=str(source),
        source_sha256=digest,
        template_name=template["name"],
        template_scope=template["scope"],
        before_issues=before_issues,
        structure=[i.to_dict() for i in items],
        changes=[],
        output=None,
        report=None,
        warnings=[],
    )

    def step(message):
        if cancel and cancel.is_set():
            raise InterruptedError("已取消；原件未修改。")
        if progress:
            progress(message)

    try:
        step("正在检查文档结构和对象完整性…")
        if not check_only:
            from .cleanup import clear_source_emphasis

            cleaned = clear_source_emphasis(doc, template["cleanup"])
            if cleaned:
                # Selected emphasis changes inside tracked insertions are intentional;
                # establish the preservation baseline after only that explicit edit.
                before = inventory(doc)
                labels = {"bold": "加粗", "italic": "斜体", "underline": "下划线"}
                report["changes"].append(dict(index=-1, action="全文清理原有"
                    + "、".join(labels[k] for k, v in template["cleanup"].items() if v)
                    + "，再应用所选模板；原生公式保留"))
            groups = set(template["enabled"])
            if groups:
                setup_styles(doc, template, update_existing="styles" in groups)
            paragraphs = list(doc.paragraphs)
            if "styles" in groups:
                for item in items:
                    key = (
                        item.kind
                        if item.kind in STYLE_NAMES
                        else "front_heading"
                        if item.kind.endswith("_heading")
                        else None
                    )
                    if key:
                        apply_style(paragraphs[item.index], key, template)
                format_notes(doc, template)
                report["changes"].append(
                    dict(index=-1, action="应用分区字体、字号、行距、样式和脚注设置")
                )
            step("正在设置编号、公式和图表…")
            if "numbering" in groups:
                number_headings(doc, items, template, report["changes"])
            if "captions" in groups:
                targets = number_captions(
                    doc, items, template, report["changes"], report["warnings"]
                )
                cross_references(doc, targets, report["warnings"])
            if "tables" in groups:
                if template["glossary_path"]:
                    from .documents import add_glossary

                    count = add_glossary(doc, template["glossary_path"], template)
                    report["changes"].append(
                        dict(index=-1, action=f"生成 {count} 条符号/缩略语释义")
                    )
                format_tables(doc, template)
                from .documents import continue_tables, landscape_tables

                continue_tables(doc, template, report["changes"], report["warnings"])
                landscape_tables(doc, template)
            if "references" in groups and reference_path:
                apply_references(
                    doc,
                    load_references(reference_path),
                    template,
                    report["warnings"],
                    report["changes"],
                )
            step("正在配置分页与目录…")
            # Re-scan after insertions. Index-based user overrides are source-only.
            from .structure import scan

            if "contents" in groups:
                add_contents(doc, template, report["changes"])
            if "pages" in groups:
                # Keep source overrides attached to their original XML nodes after insertions.
                kinds = {
                    paragraphs[int(index)]._p: kind
                    for index, kind in template["structure_overrides"].items()
                    if int(index) < len(paragraphs)
                }
                final_overrides = {
                    str(i): kinds[p._p]
                    for i, p in enumerate(doc.paragraphs)
                    if p._p in kinds
                }
                paginate(doc, scan(doc, final_overrides), template, report["changes"])
            if "figures" in groups:
                count = format_images(doc, template)
                report["changes"].append(
                    dict(index=-1, action=f"按页面与单元格等比缩小 {count} 张图片，调整独立图片与题注间距")
                )
            from .xmlutil import normalize_order

            normalize_order(doc)
            after = inventory(doc)
            report["integrity"] = compare_inventory(before, after)
            if not report["integrity"]["passed"]:
                raise RuntimeError(
                    "对象完整性检查未通过，未发布排版副本："
                    + ", ".join(report["integrity"]["lost_categories"])
                )
            step("正在保存排版副本…")
            base = directory / f"{safe_stem}_排版.docx"
            seen = set()
            while True:
                destination = unique_path(base, seen)
                if pdf and destination.with_suffix('.pdf').exists():
                    continue
                try:
                    save_document(doc, destination)
                    break
                except FileExistsError:
                    # Another job may publish the same name after planning.
                    continue
            report["output"] = str(destination)
            if host:
                step("正在更新目录与交叉引用" + ("并导出 PDF…" if pdf else "…"))
                from .office_io import finalize

                # Office works on a separate staging copy; a failed update never damages the formatted output.
                # A timed-out Office server may retain its file lock after the worker
                # exits. Cleanup must never turn an already saved document into a
                # failed job. Keep those intermediates in the OS temp folder.
                with tempfile.TemporaryDirectory(
                    prefix="thesiscraft-office-", ignore_cleanup_errors=True
                ) as temporary:
                    stage = Path(temporary) / "office-stage.docx"
                    shutil.copy2(destination, stage)
                    try:
                        office = finalize(stage, host, pdf)
                        report["office"] = office
                        if office.get("success"):
                            checked = compare_inventory(before, inventory(Document(stage)))
                            # Office legitimately rewrites XML; binary media loss is still unacceptable.
                            binaries_lost = [
                                s for s in checked["lost_categories"] if s.startswith("/")
                            ]
                            host_inventory = inventory(Document(stage))
                            # Office rewrites equation/revision XML, so byte hashes
                            # differ legitimately; disappearing objects do not.
                            for tag, label in (("m:oMath", "公式"),
                                               ("w:ins", "插入修订"),
                                               ("w:del", "删除修订")):
                                if sum(host_inventory[tag].values()) < sum(before[tag].values()):
                                    binaries_lost.append(label)
                            if set(before["bookmarks"]) - set(host_inventory["bookmarks"]):
                                binaries_lost.append("原有书签")
                            original_manager_fields = [
                                c.strip()
                                for c in before["field_codes"]
                                if "ADDIN" in c.upper()
                            ]
                            if any(
                                c not in {v.strip() for v in host_inventory["field_codes"]}
                                for c in original_manager_fields
                            ):
                                binaries_lost.append("文献管理器域")
                            if binaries_lost:
                                raise RuntimeError(
                                    "Office 更新后嵌入对象发生变化："
                                    + ", ".join(binaries_lost)
                                )
                            publish_file(stage, destination, overwrite=True)
                            if office.get("pdf"):
                                pdf_path = destination.with_suffix(".pdf")
                                publish_file(office["pdf"], pdf_path)
                                office["pdf"] = str(pdf_path)
                                from .documents import inspect_pdf

                                office["pdf_check"] = inspect_pdf(pdf_path)
                            doc = Document(destination)
                        else:
                            report["warnings"].append(
                                dict(
                                    code="office-update",
                                    index=-1,
                                    message="Office 更新未完成，已保留未更新的排版副本："
                                    + str(office.get("error", office.get("errors"))),
                                )
                            )
                    except Exception as exc:
                        report.setdefault("office", {}).update(
                            success=False, error=str(exc), pdf=None
                        )
                        report["warnings"].append(
                            dict(code="office-update", index=-1, message=
                                 "排版副本已生成；目录、交叉引用自动更新未完成，可在 Word/WPS 中更新域。"
                                 + str(exc))
                        )
            elif pdf:
                report["warnings"].append(
                    dict(
                        code="pdf-host",
                        index=-1,
                        message="请选择 Word 或 WPS 来生成 PDF。",
                    )
                )
            report["after_issues"], final_structure = audit(
                doc, {**template, "structure_overrides": {}}
            )
            report["source_structure"] = report["structure"]
            report["structure"] = [item.to_dict() for item in final_structure]
            # Exact identity of the original is verified at the end of the operation.
            if hashlib.sha256(source.read_bytes()).hexdigest() != digest:
                report["warnings"].append(
                    dict(
                        code="source-changed",
                        index=-1,
                        message="处理期间原件被其他程序修改；当前副本基于启动时版本。",
                    )
                )
            report["after_issues"].extend(report["warnings"])
        report["duration_seconds"] = round(time.monotonic() - started, 2)
        if check_only:
            report["report"] = str(directory / "检查报告.html")
            write_report(directory, report)
        step("检查报告已生成。" if check_only else "排版副本已生成。")
        return report
    except Exception as exc:
        report["error"] = str(exc)
        report["duration_seconds"] = round(time.monotonic() - started, 2)
        if check_only:
            write_report(directory, report)
        raise
