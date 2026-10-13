"""Opt-in course-paper layout; never insert research content or citation records."""

from copy import deepcopy
import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.run import Run

from .templates import STYLE_NAMES
from .xmlutil import field, visible_text


def _heading_text_is_safe(paragraph):
    protected = ('fldChar', 'fldSimple', 'instrText', 'hyperlink', 'drawing',
                 'object', 'pict', 'ins', 'del', 'sdt')
    if any(next(paragraph._p.iter(qn('w:' + tag)), None) is not None for tag in protected):
        return False
    if any(next(paragraph._p.iter(qn(tag)), None) is not None
           for tag in ('m:oMath', 'm:oMathPara')):
        return False
    allowed = {qn('w:' + tag) for tag in ('rPr', 't', 'footnoteReference', 'endnoteReference')}
    return all(child.tag in allowed for run in paragraph.runs for child in run._r)


def _split_run(run, offset, paragraph):
    """Split eligible text while retaining each note anchor exactly once in place."""
    copy = deepcopy(run._r)
    left, right, consumed = [], [], 0
    for child in run._r:
        if child.tag == qn('w:rPr'):
            continue
        if child.tag == qn('w:t'):
            text = child.text or ''
            point = max(0, min(len(text), offset - consumed))
            if point:
                node = deepcopy(child)
                node.text = text[:point]
                node.set(qn('xml:space'), 'preserve')
                left.append(node)
            if point < len(text):
                node = deepcopy(child)
                node.text = text[point:]
                node.set(qn('xml:space'), 'preserve')
                right.append(node)
            consumed += len(text)
        else:
            (left if consumed < offset else right).append(deepcopy(child))
    for node, children in ((run._r, left), (copy, right)):
        for child in list(node):
            if child.tag != qn('w:rPr'):
                node.remove(child)
        node.extend(children)
    run._r.addnext(copy)
    return Run(copy, paragraph)


def subtitle_boundary(paragraph, template):
    """Only split ordinary text; preserve fields, links and tracked changes intact."""
    separator = template["course"]["heading_subtitle_separator"]
    if template["document_type"] != "course" or not separator or not _heading_text_is_safe(paragraph):
        return None
    text = visible_text(paragraph)
    before, found, after = text.partition(separator)
    if not found or not before.strip() or not after.strip():
        return None
    return len(before) + len(found)


def heading_run_sizes(paragraph, template):
    """Expected sizes per text run, including an unsplit run that crosses the boundary."""
    boundary = subtitle_boundary(paragraph, template)
    if boundary is None:
        return {}
    main = template["styles"]["h1"]["size"]
    secondary = template["course"]["heading_subtitle_size"]
    result, offset = {}, 0
    for run in paragraph.runs:
        end = offset + len(run.text)
        sizes = []
        if offset < boundary and run.text[:max(0, boundary - offset)].strip():
            sizes.append(main)
        if end > boundary and run.text[max(0, boundary - offset):].strip():
            sizes.append(secondary)
        if sizes:
            result[run._r] = sizes
        offset = end
    return result


def format_heading_subtitle(paragraph, template):
    boundary = subtitle_boundary(paragraph, template)
    if boundary is None:
        return
    offset = 0
    secondary = template["course"]["heading_subtitle_size"]
    for run in list(paragraph.runs):
        text, end = run.text, offset + len(run.text)
        if offset < boundary < end:
            remainder = _split_run(run, boundary - offset, paragraph)
            remainder.font.size = Pt(secondary)
        elif offset >= boundary and text:
            run.font.size = Pt(secondary)
        offset = end


def format_inline_label(paragraph, key):
    label = "摘要" if key == "abstract" else "(?:关键词|关键字)"
    match = re.match(r"^\s*" + label + r"\s*[:：]", visible_text(paragraph))
    if not match or not _heading_text_is_safe(paragraph):
        return
    boundary, offset = match.end(), 0
    for run in list(paragraph.runs):
        end = offset + len(run.text)
        if offset < boundary < end:
            _split_run(run, boundary - offset, paragraph)
        if offset < boundary:
            run.bold = True
        offset = end


def compact_skeleton(doc, template):
    """A short editable course draft: inline abstract, identity fields, no cover or TOC."""
    meta = template["metadata"]
    doc.add_paragraph(meta["title"] or "【论文题目】", STYLE_NAMES["title"])
    doc.add_paragraph(meta["title_en"] or "【副标题（不需要时可删除）】",
                      STYLE_NAMES["subtitle"])
    empty = "________"
    doc.add_paragraph(f"姓名：{meta['author'] or empty}    学号：{meta['student_id'] or empty}",
                      STYLE_NAMES["cover_info"])
    doc.add_paragraph(f"院系：{template['department'] or empty}    班级：{empty}    "
                      f"授课教师：{meta['supervisor'] or empty}", STYLE_NAMES["cover_info"])
    for key, label, text in (
        ("abstract", "摘要：", "【填写摘要】"),
        ("keywords", "关键词：", meta["keywords"] or "【关键词一；关键词二】"),
    ):
        p = doc.add_paragraph(style=STYLE_NAMES[key])
        p.add_run(label).bold = True
        p.add_run(text)
    for title in ("一、引言", "二、分析与讨论", "三、结论"):
        doc.add_paragraph(title, STYLE_NAMES["h1"])
        doc.add_paragraph("【在这里填写正文】", STYLE_NAMES["body"])
    doc.add_paragraph("参考文献", STYLE_NAMES["front_heading"])
    doc.add_paragraph("【按实际使用的文献填写，不得将格式示例作为真实引文】",
                      STYLE_NAMES["bibliography"])
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field(footer, "PAGE", "1", template["styles"]["footer"])
    # Built-in Word's Subtitle style is italic. The course style explicitly
    # chooses upright text, without clearing italic text in the user's paper.
    doc.styles[STYLE_NAMES["subtitle"]].font.italic = False
