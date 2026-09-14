"""Structured bibliography import and citation fields; existing manager fields stay intact."""

import json
import re
from pathlib import Path
from .xmlutil import bookmark_name, bookmark, field
from .fields import substitute, field_nodes
from .layout import apply_style


def load_references(path):
    if not path:
        return []
    path = Path(path)
    text = path.read_text(encoding="utf-8-sig")
    if path.suffix.lower() == ".json":
        records = json.loads(text)
        if not isinstance(records, list):
            raise ValueError("参考文献 JSON 必须是数组。")
    elif path.suffix.lower() == ".bib":
        import bibtexparser

        parsed = bibtexparser.loads(text)
        records = []
        for item in parsed.entries:
            records.append(
                dict(
                    key=item["ID"],
                    type={
                        "article": "journal",
                        "book": "book",
                        "inproceedings": "conference",
                        "phdthesis": "thesis",
                        "mastersthesis": "thesis",
                        "techreport": "report",
                        "misc": "web",
                    }.get(item["ENTRYTYPE"], "unsupported"),
                    authors=item.get("author", "").split(" and "),
                    title=item.get("title", "").replace("{", "").replace("}", ""),
                    year=item.get("year", ""),
                    journal=item.get("journal", ""),
                    publisher=item.get("publisher", item.get("school", "")),
                    place=item.get("address", ""),
                    volume=item.get("volume", ""),
                    issue=item.get("number", ""),
                    pages=item.get("pages", "").replace("--", "-"),
                    doi=item.get("doi", ""),
                    url=item.get("url", ""),
                    booktitle=item.get("booktitle", ""),
                    accessed=item.get("urldate", ""),
                )
            )
    elif path.suffix.lower() == ".ris":
        records, raw, last = [], {}, None
        for line in text.splitlines():
            match = re.match(r"^([A-Z0-9]{2})\s{2}-\s?(.*)$", line)
            if match:
                tag, value = match.groups()
                if tag == "ER":

                    def get(*tags):
                        return next((raw[t][0] for t in tags if t in raw), "")

                    records.append(
                        dict(
                            key=get("ID") or f"ref{len(records) + 1}",
                            type={
                                "JOUR": "journal",
                                "BOOK": "book",
                                "CONF": "conference",
                                "CPAPER": "conference",
                                "THES": "thesis",
                                "RPRT": "report",
                                "ELEC": "web",
                            }.get(get("TY"), "unsupported"),
                            authors=raw.get("AU", raw.get("A1", [])),
                            title=get("TI", "T1"),
                            year=get("PY", "Y1")[:4],
                            journal=get("JO", "JF", "T2"),
                            publisher=get("PB"),
                            place=get("CY"),
                            volume=get("VL"),
                            issue=get("IS"),
                            pages=get("SP") + ("-" + get("EP") if get("EP") else ""),
                            doi=get("DO"),
                            url=get("UR"),
                            booktitle=get("T2"),
                            accessed=get("Y2"),
                        )
                    )
                    raw, last = {}, None
                else:
                    raw.setdefault(tag, []).append(value)
                    last = tag
            elif last and line.strip():
                raw[last][-1] += " " + line.strip()
        if raw:
            raise ValueError("RIS 最后一个记录缺少 ER 结束标记。")
    else:
        raise ValueError("文献库支持 JSON、RIS、BibTeX。")
    seen = set()
    for row in records:
        if not isinstance(row, dict) or not all(
            row.get(k) for k in ("key", "title", "year", "authors")
        ):
            raise ValueError("每条文献需要 key、title、year、authors。")
        if not re.fullmatch(r"[\w.-]+", row["key"]) or row["key"] in seen:
            raise ValueError("文献 key 不合法或重复：" + row["key"])
        seen.add(row["key"])
        if not isinstance(row["authors"], list) or any(
            not isinstance(a, str) or not a for a in row["authors"]
        ):
            raise ValueError("authors 必须是作者姓名数组。")
        row.setdefault("type", "journal")
        if row["type"] not in (
            "journal",
            "book",
            "conference",
            "thesis",
            "web",
            "report",
        ):
            raise ValueError(
                "不支持的文献类型：" + row["type"] + "；请保留文献管理器排版。"
            )
        for k, v in row.items():
            if k != "authors" and not isinstance(v, str):
                raise ValueError("文献字段须为文本：" + k)
    return records


def apply_references(doc, records, template, warnings, changes):
    citation = re.compile(r"\{\{cite:([\w., -]+)\}\}")
    lookup = {r["key"]: r for r in records}
    from docx.oxml.ns import qn

    bookmark_keys = {bookmark_name("bib:" + key): key for key in lookup}
    generated = []
    ordered = []
    for p in doc.paragraphs:
        owned = [
            n.get(qn("w:name"))
            for n in p._p.iter(qn("w:bookmarkStart"))
            if n.get(qn("w:name")) in bookmark_keys
        ]
        if owned and p.style.name == "PS bibliography":
            generated.append(p)
            continue
        for node in p._p.iter(qn("w:instrText")):
            match = re.match(r"\s*REF\s+(PS_\w+)", node.text or "")
            if match and match.group(1) in bookmark_keys:
                key = bookmark_keys[match.group(1)]
                if key not in ordered:
                    ordered.append(key)
        for match in citation.finditer(p.text):
            for key in match.group(1).replace(" ", "").split(","):
                if key in lookup and key not in ordered:
                    ordered.append(key)
                elif key not in lookup:
                    warnings.append(
                        dict(
                            code="missing-citation",
                            index=-1,
                            message="文献库缺少：" + key,
                        )
                    )
    if not ordered:
        return
    style = template["references"]["style"]
    if style == "apa":
        ordered.sort(
            key=lambda key: (
                lookup[key]["authors"][0],
                lookup[key]["year"],
                lookup[key]["title"],
            )
        )
    from .csl import render

    rendered, render_citation = render(
        records, ordered, style, template["references"]["csl_path"]
    )
    # Rebuild only entries generated by Paper Studio; manual and manager entries stay intact.
    retained_bookmarks = []
    for p in generated:
        own_ids = {
            n.get(qn("w:id"))
            for n in p._p.iter(qn("w:bookmarkStart"))
            if n.get(qn("w:name")) in bookmark_keys
        }
        retained_bookmarks.extend(
            n
            for n in list(p._p)
            if n.tag in (qn("w:bookmarkStart"), qn("w:bookmarkEnd"))
            and n.get(qn("w:id")) not in own_ids
        )
        p._p.getparent().remove(p._p)
    heading = next(
        (
            p
            for p in doc.paragraphs
            if p.text.strip() in ("参考文献", "References", "Bibliography")
        ),
        None,
    )
    if heading is None:
        heading = doc.add_paragraph("参考文献")
    apply_style(heading, "front_heading", template)
    anchor = heading
    for index, key in enumerate(ordered, 1):
        row = lookup[key]
        from .xmlutil import insert_after

        p = insert_after(anchor)
        name = bookmark_name("bib:" + key)
        if style != "apa":
            p.add_run("[")
            first, display = field(
                p,
                "SEQ PSBibliography" + (" \\r 1" if index == 1 else "") + " \\* ARABIC",
                str(index),
            )
            bookmark(p, name, first._r, p._p[-1])
            p.add_run("] ")
        else:
            bookmark(p, name)
        p.add_run(rendered[key])
        apply_style(p, "bibliography", template)
        anchor = p
        for required in (
            ("journal",)
            if row["type"] == "journal"
            else ("url", "accessed")
            if row["type"] == "web"
            else ("publisher",)
        ):
            if not row.get(required):
                warnings.append(
                    dict(
                        code="reference-metadata",
                        index=-1,
                        message=f"{key} 缺少 {required}，请补全文献著录信息。",
                    )
                )
    for node in retained_bookmarks:
        anchor._p.append(node)
    for p in doc.paragraphs:
        for match in reversed(list(citation.finditer(p.text))):
            keys = match.group(1).replace(" ", "").split(",")
            if any(k not in ordered for k in keys):
                continue
            from docx.text.paragraph import Paragraph
            from .xmlutil import element

            temp = Paragraph(element("w:p"), None)
            if style == "apa":
                temp.add_run(render_citation(keys))
            else:
                temp.add_run("[")
                for i, key in enumerate(keys):
                    if i:
                        temp.add_run(", ")
                    for node in field_nodes(
                        f"REF {bookmark_name('bib:' + key)} \\h",
                        str(ordered.index(key) + 1),
                    ):
                        temp._p.append(node)
                temp.add_run("]")
            if template["references"].get("superscript", False):
                for run in temp.runs:
                    run.font.superscript = True
            substitute(p, *match.span(), list(temp._p))
    changes.append(dict(index=-1, action=f"生成 {len(ordered)} 条结构化参考文献与引用"))
