"""Conservative OOXML building blocks for academic fields and styles."""

from copy import deepcopy
import hashlib
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph


def element(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn("w:" + key), str(value))
    return node


def set_child(parent, tag, **attrs):
    for child in list(parent.findall(qn(tag))):
        parent.remove(child)
    node = element(tag, **attrs)
    parent.append(node)
    return node


def plain(paragraph):
    forbidden = (
        "fldChar",
        "fldSimple",
        "instrText",
        "hyperlink",
        "drawing",
        "object",
        "pict",
        "ins",
        "del",
        "sdt",
        "footnoteReference",
        "endnoteReference",
    )
    return not any(
        next(paragraph._p.iter(qn("w:" + tag)), None) is not None for tag in forbidden
    )


def visible_text(paragraph):
    return "".join(n.text or "" for n in paragraph._p.iter(qn("w:t")))


def set_font(run, spec):
    run.font.name = spec.get("latin", spec.get("font", "Times New Roman"))
    run.font.size = Pt(spec["size"])
    run.font.bold = spec.get("bold", False)
    if "color" in spec:
        run.font.color.rgb = RGBColor.from_string(spec["color"])
    fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), spec.get("font", "宋体"))
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + key), None)


def field(paragraph, instruction, cached="", spec=None):
    start = paragraph.add_run()
    start._r.append(element("w:fldChar", fldCharType="begin", dirty="true"))
    code = paragraph.add_run()
    instr = element("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " " + instruction + " "
    code._r.append(instr)
    paragraph.add_run()._r.append(element("w:fldChar", fldCharType="separate"))
    display = paragraph.add_run(str(cached))
    if spec:
        set_font(display, spec)
    paragraph.add_run()._r.append(element("w:fldChar", fldCharType="end"))
    return start, display


def bookmark_name(key):
    return "PS_" + hashlib.sha256(key.encode("utf-8")).hexdigest()[:28]


def bookmark(paragraph, name, start=None, end=None):
    root = paragraph._p.getroottree().getroot()
    if any(n.get(qn("w:name")) == name for n in root.iter(qn("w:bookmarkStart"))):
        return name
    ids = [int(n.get(qn("w:id"), "0")) for n in root.iter(qn("w:bookmarkStart"))]
    ident = max(ids, default=0) + 1
    first, last = (
        element("w:bookmarkStart", id=ident, name=name),
        element("w:bookmarkEnd", id=ident),
    )
    if start is None:
        ppr = paragraph._p.pPr
        if ppr is not None:
            ppr.addnext(first)
        else:
            paragraph._p.insert(0, first)
    else:
        start.addprevious(first)
    if end is None:
        paragraph._p.append(last)
    else:
        end.addnext(last)
    return name


def insert_before(anchor, text="", style=None):
    return anchor.insert_paragraph_before(text, style)


def insert_after(anchor, text="", style=None):
    node = element("w:p")
    anchor._p.addnext(node)
    paragraph = Paragraph(node, anchor._parent)
    if text:
        paragraph.add_run(text)
    if style:
        paragraph.style = style
    return paragraph


def field_instructions(root):
    """Read complete fields without changing runs, cached results or source XML.

    Complex instructions may span runs and paragraphs. Keep each field separate;
    a nested instruction is represented by {} without adding token separators.
    Its code is never concatenated with the parent's literal tokens, and nested
    instructions are yielded separately.
    This reads field syntax; it does not evaluate dynamic field arguments.
    """
    stack = []
    for node in root.iter():
        if node.tag == qn("w:fldSimple"):
            if stack and stack[-1]["code"]:
                stack[-1]["pieces"].append("{}")
            yield node.get(qn("w:instr"), "")
        elif node.tag == qn("w:fldChar"):
            kind = node.get(qn("w:fldCharType"))
            if kind == "begin":
                if stack and stack[-1]["code"]:
                    stack[-1]["pieces"].append("{}")
                stack.append(dict(pieces=[], code=True))
            elif kind == "separate" and stack:
                stack[-1]["code"] = False
            elif kind == "end" and stack:
                yield "".join(stack.pop()["pieces"])
        elif node.tag == qn("w:instrText") and stack and stack[-1]["code"]:
            stack[-1]["pieces"].append(node.text or "")


def has_field(doc, word):
    return any(word.casefold() in code.casefold() for code in field_instructions(doc.element))


def copy_section_end(doc, before, settings=None):
    """The inserted break stores properties of the section ENDING before the anchor."""
    previous = before._p.getprevious()
    if previous is not None and previous.tag == qn("w:p"):
        p = Paragraph(previous, before._parent)
        existing = p._p.get_or_add_pPr().find(qn("w:sectPr"))
        if existing is not None:
            return existing
    break_p = insert_before(before)
    section = deepcopy(settings if settings is not None else doc.sections[-1]._sectPr)
    break_p._p.get_or_add_pPr().append(section)
    return section


def normalize_order(doc):
    """Keep generated property children in OOXML schema order for Word and WPS."""
    orders = {
        "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
        "sectPr": "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings sectPrChange",
        "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath rPrChange",
        "lvl": "start numFmt lvlRestart pStyle isLgl suff lvlText lvlPicBulletId legacy lvlJc pPr rPr",
    }
    for part in doc.part.package.parts:
        root = getattr(part, "element", None)
        if root is None:
            continue
        for tag, order in orders.items():
            positions = {qn("w:" + key): i for i, key in enumerate(order.split())}
            for node in root.iter(qn("w:" + tag)):
                children = sorted(
                    list(node), key=lambda child: positions.get(child.tag, 999)
                )
                node[:] = children
