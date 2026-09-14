"""Real Word styles, multilevel lists, section pagination and academic objects."""

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from docx.text.paragraph import Paragraph

from .templates import STYLE_NAMES
from .xmlutil import element, set_child, set_font, plain, copy_section_end, field
from ..ooxml import paragraph_runs, replace_text_nodes, iter_tables

ALIGNS = dict(
    left=WD_ALIGN_PARAGRAPH.LEFT,
    center=WD_ALIGN_PARAGRAPH.CENTER,
    right=WD_ALIGN_PARAGRAPH.RIGHT,
    justify=WD_ALIGN_PARAGRAPH.JUSTIFY,
)


def setup_styles(doc, template, update_existing=True):
    for key, name in STYLE_NAMES.items():
        if not update_existing and name in doc.styles:
            continue
        style = (
            doc.styles[name]
            if name in doc.styles
            else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        )
        spec = template["styles"][key]
        style.font.name, style.font.size, style.font.bold = (
            spec["latin"],
            Pt(spec["size"]),
            spec["bold"],
        )
        style.font.color.rgb = RGBColor.from_string(spec["color"])
        style.element.get_or_add_rPr().get_or_add_rFonts().set(
            qn("w:eastAsia"), spec["font"]
        )
        fmt = style.paragraph_format
        fmt.alignment = ALIGNS[spec["align"]]
        fmt.line_spacing = (
            Pt(spec["spacing"])
            if spec["spacing_unit"] == "pt"
            else float(spec["spacing"])
        )
        fmt.space_before, fmt.space_after = Pt(spec["before"]), Pt(spec["after"])
        fmt.first_line_indent = Pt(spec["size"] * spec["indent"])
        if key.startswith("toc"):
            fmt.left_indent = Pt((int(key[-1]) - 1) * spec["size"])
        if key.startswith("h") and key in ("h1", "h2", "h3", "h4"):
            set_child(
                style.element.get_or_add_pPr(), "w:outlineLvl", val=int(key[-1]) - 1
            )
            fmt.keep_with_next = True
        if key in ("appendix", "front_heading"):
            set_child(style.element.get_or_add_pPr(), "w:outlineLvl", val=0)
            fmt.keep_with_next = True
        if key.startswith("caption"):
            fmt.keep_together = True


def apply_style(paragraph, key, template):
    if key not in STYLE_NAMES:
        return
    paragraph.style = STYLE_NAMES[key]
    spec = template["styles"][key]
    fmt = paragraph.paragraph_format
    fmt.alignment = ALIGNS[spec["align"]]
    fmt.space_before, fmt.space_after = Pt(spec["before"]), Pt(spec["after"])
    fmt.line_spacing = (
        Pt(spec["spacing"]) if spec["spacing_unit"] == "pt" else float(spec["spacing"])
    )
    # Exact line height clips inline pictures and tall equations even though
    # their XML/media survive. Keep the requested spacing as a minimum instead.
    if spec["spacing_unit"] == "pt" and any(
        next(paragraph._p.iter(qn(tag)), None) is not None
        for tag in ('wp:inline', 'w:object', 'w:pict', 'm:oMath')
    ):
        fmt.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    fmt.first_line_indent = Pt(spec["indent"] * spec["size"])
    fmt.left_indent = None
    fmt.right_indent = None
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is not None:
        for name in (
            "firstLineChars",
            "hangingChars",
            "leftChars",
            "rightChars",
            "hanging",
        ):
            ind.attrib.pop(qn("w:" + name), None)
    # Update ordinary text even next to fields, but never a manager field's cached result.
    depth = 0
    for run in paragraph.runs:
        flags = list(run._r.iter(qn("w:fldChar")))
        depth += sum(n.get(qn("w:fldCharType")) == "begin" for n in flags)
        if (
            not depth
            and run.text
            and all(child.tag in (qn("w:rPr"), qn("w:t")) for child in run._r)
        ):
            original_bold = run.bold
            set_font(run, spec)
            if key in ("body", "abstract", "abstract_en") and original_bold:
                run.bold = True
        depth = max(0, depth - sum(n.get(qn("w:fldCharType")) == "end" for n in flags))
    if key in ("h1", "h2", "h3", "h4", "front_heading", "appendix"):
        fmt.keep_with_next = True
    if key == "bibliography":
        fmt.left_indent = Cm(template["references"]["hanging_cm"])
        fmt.first_line_indent = -Cm(template["references"]["hanging_cm"])


def _number_definition(doc, template, appendix=False):
    root = doc.part.numbering_part.element
    label = "PaperStudioAppendix" if appendix else "PaperStudioChapters"
    retained_ids = None
    for node in root.findall(qn("w:abstractNum")):
        name = node.find(qn("w:name"))
        if name is not None and name.get(qn("w:val")) == label:
            abstract_id = int(node.get(qn("w:abstractNumId")))
            for num in root.findall(qn("w:num")):
                if int(num.find(qn("w:abstractNumId")).get(qn("w:val"))) == abstract_id:
                    # Rebuild the shared definition so template edits apply on a repeated run.
                    retained_ids = (abstract_id, int(num.get(qn("w:numId"))))
                    root.remove(node)
                    root.remove(num)
                    break
            break
    abstract_id = (
        max(
            [
                int(n.get(qn("w:abstractNumId")))
                for n in root.findall(qn("w:abstractNum"))
            ],
            default=-1,
        )
        + 1
    )
    num_id = (
        max([int(n.get(qn("w:numId"))) for n in root.findall(qn("w:num"))], default=0)
        + 1
    )
    if retained_ids is not None:
        abstract_id, num_id = retained_ids
    abstract = element("w:abstractNum", abstractNumId=abstract_id)
    abstract.append(element("w:multiLevelType", val="multilevel"))
    abstract.append(element("w:name", val=label))
    config = template["numbering"]
    chapter_patterns = dict(
        chapter="第%1章",
        chinese="第%1章",
        decimal="%1",
        english="Chapter %1",
        none="%1",
    )
    for level in range(1 if appendix else 4):
        lvl = element("w:lvl", ilvl=level)
        lvl.append(element("w:start", val=1))
        fmt = (
            config["appendix_format"]
            if appendix
            else (
                "chineseCounting"
                if level == 0 and config["scheme"] == "chinese"
                else "decimal"
            )
        )
        lvl.append(element("w:numFmt", val=fmt))
        lvl.append(
            element(
                "w:pStyle",
                val=doc.styles[
                    STYLE_NAMES["appendix" if appendix else f"h{level + 1}"]
                ].style_id,
            )
        )
        if level:
            lvl.append(element("w:isLgl", val="1"))
        lvl.append(element("w:suff", val="space"))
        pattern = (
            "附录 %1"
            if appendix
            else (
                chapter_patterns[config["scheme"]]
                if level == 0
                else ".".join("%" + str(i + 1) for i in range(level + 1))
            )
        )
        lvl.append(element("w:lvlText", val=pattern))
        lvl.append(element("w:lvlJc", val="left"))
        rpr = element("w:rPr")
        spec = (
            template["styles"]["appendix" if appendix else f"h{level + 1}"]
            if config["match_heading"]
            else dict(
                latin=config["font"],
                font=config["east_asia"],
                size=config["size"],
                bold=config["bold"],
            )
        )
        rpr.append(
            element(
                "w:rFonts",
                ascii=spec["latin"],
                hAnsi=spec["latin"],
                eastAsia=spec["font"],
            )
        )
        rpr.append(element("w:sz", val=round(spec["size"] * 2)))
        rpr.append(element("w:b", val=int(spec["bold"])))
        lvl.append(rpr)
        abstract.append(lvl)
    # abstractNum must precede num according to the OOXML content model.
    first_num = root.find(qn("w:num"))
    if first_num is None:
        root.append(abstract)
    else:
        first_num.addprevious(abstract)
    num = element("w:num", numId=num_id)
    num.append(element("w:abstractNumId", val=abstract_id))
    root.append(num)
    for level in range(1 if appendix else 4):
        style = doc.styles[STYLE_NAMES["appendix" if appendix else f"h{level + 1}"]]
        props = set_child(style.element.get_or_add_pPr(), "w:numPr")
        props.append(element("w:ilvl", val=level))
        props.append(element("w:numId", val=num_id))
    return num_id


def number_headings(doc, items, template, changes):
    if template["numbering"]["scheme"] == "none":
        return
    ids = {
        False: _number_definition(doc, template),
        True: _number_definition(doc, template, True),
    }
    paragraphs = list(doc.paragraphs)
    for item in items:
        if item.kind not in ("h1", "h2", "h3", "h4", "appendix"):
            continue
        p = paragraphs[item.index]
        if not plain(p):
            changes.append(
                dict(index=item.index, action="保留复杂标题编号，需人工确认")
            )
            continue
        if item.prefix:
            runs = list(paragraph_runs(p))
            old = "".join(r.text for r in runs)
            replace_text_nodes(runs, old[len(item.prefix) :])
        appendix = item.kind == "appendix"
        p.style = STYLE_NAMES[item.kind]
        num = set_child(p._p.get_or_add_pPr(), "w:numPr")
        num.append(element("w:ilvl", val=0 if appendix else int(item.kind[-1]) - 1))
        num.append(element("w:numId", val=ids[appendix]))
        counter = "PSAppendix" if appendix else "PSChapter"
        if item.kind in ("h1", "appendix"):
            field(p, f"SEQ {counter} \\h", "")
        changes.append(
            dict(
                index=item.index,
                action="应用可更新的多级编号",
                before=item.text,
                after=p.text,
            )
        )


def _managed_paragraph(part, style_name):
    return (
        next((p for p in part.paragraphs if p.style.name == style_name), None)
        or next((p for p in part.paragraphs if not p.text and len(p._p) == 0), None)
        or part.add_paragraph()
    )


def paginate(doc, items, template, changes):
    paragraphs = list(doc.paragraphs)
    regions = {p._p: item.region for p, item in zip(paragraphs, items)}
    boundaries = [
        paragraphs[i.index]
        for i in items
        if i.kind in ("h1", "appendix", "abstract_heading", "abstract_en_heading")
        or i.kind
        in (
            "bibliography_heading",
            "acknowledgements_heading",
            "symbols_heading",
            "publications_heading",
        )
    ]
    for p in boundaries:
        if p._p is not doc.element.body[0]:
            copy_section_end(doc, p)
        p.paragraph_format.page_break_before = False
    # Determine each section's logical region from its first meaningful source paragraph.
    section_info, region, title = [], None, ""
    for node in doc.element.body:
        if node in regions and region is None:
            region = regions[node]
            title = Paragraph(node, doc._body).text
        end = (
            node.find(qn("w:pPr") + "/" + qn("w:sectPr"))
            if node.tag == qn("w:p")
            else None
        )
        if end is not None or node.tag == qn("w:sectPr"):
            section_info.append((region or "body", title))
            region, title = None, ""
    page = template["page"]
    doc.settings.odd_and_even_pages_header_footer = page["odd_even"]
    set_child(doc.settings.element, "w:mirrorMargins", val=int(page["mirror_margins"]))
    first_front, first_body = True, True
    for index, section in enumerate(doc.sections):
        region, title = section_info[index]
        is_cover = region == "front"
        is_front = region in ("abstract", "abstract_en") or (
            region == "symbols" and first_body
        )
        for name in ("top", "bottom", "left", "right"):
            setattr(section, name + "_margin", Cm(page[name]))
        section.gutter = Cm(page["gutter"])
        section.header_distance, section.footer_distance = (
            Cm(page["header_distance"]),
            Cm(page["footer_distance"]),
        )
        if (
            not page["preserve_landscape"] and not template["tables"]["landscape_wide"]
        ) or section.orientation != WD_ORIENT.LANDSCAPE:
            section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.start_type = (
            WD_SECTION_START.ODD_PAGE
            if page["chapter_recto"] and not is_cover and not is_front
            else WD_SECTION_START.NEW_PAGE
        )
        section.different_first_page_header_footer = False
        fmt = page["front_number"] if is_front else page["body_number"]
        attrs = dict(fmt=fmt)
        if is_front and first_front:
            attrs["start"], first_front = 1, False
        elif not is_cover and not is_front and first_body:
            attrs["start"], first_body = 1, False
        set_child(section._sectPr, "w:pgNumType", **attrs)
        # Unlink all variants BEFORE modifying any footer, preserving previous sections.
        for part in (
            section.header,
            section.even_page_header,
            section.first_page_header,
            section.footer,
            section.even_page_footer,
            section.first_page_footer,
        ):
            part.is_linked_to_previous = False
        for even, footer in ((False, section.footer), (True, section.even_page_footer)):
            existing = next(
                (
                    p
                    for p in footer.paragraphs
                    if any(
                        "PAGE" == (n.text or "").strip().split(" ")[0]
                        for n in p._p.iter(qn("w:instrText"))
                    )
                ),
                None,
            )
            if is_cover:
                if existing is not None:
                    changes.append(
                        dict(index=-1, action="封面存在旧页码，保留原页脚并提示检查")
                    )
                continue
            p = (
                existing
                if existing is not None
                else _managed_paragraph(footer, STYLE_NAMES["footer"])
            )
            if existing is None and not any(p._p.iter(qn("w:fldChar"))):
                field(p, "PAGE", "1", template["styles"]["footer"])
            p.style = STYLE_NAMES["footer"]
            p.alignment = (
                (WD_ALIGN_PARAGRAPH.LEFT if even else WD_ALIGN_PARAGRAPH.RIGHT)
                if page["number_position"] == "outside"
                else WD_ALIGN_PARAGRAPH.CENTER
            )
        if is_cover:
            continue
        degree = {"bachelor": "本科", "master": "硕士", "doctor": "博士"}[
            template["degree"]
        ]
        text = (
            page["header_text"]
            .replace("{school}", template["school"])
            .replace("{degree}", degree)
        )
        for even, header in ((False, section.header), (True, section.even_page_header)):
            p = _managed_paragraph(header, STYLE_NAMES["header"])
            p.clear()
            p.style = STYLE_NAMES["header"]
            if (
                page["chapter_header"]
                and not is_front
                and not even
                and region == "body"
            ):
                field(p, "STYLEREF 1", title, template["styles"]["header"])
            else:
                set_font(
                    p.add_run(title if not even and title else text.strip()),
                    template["styles"]["header"],
                )
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if page["header_line_pt"]:
                borders = set_child(p._p.get_or_add_pPr(), "w:pBdr")
                borders.append(
                    element(
                        "w:bottom",
                        val="single",
                        sz=round(page["header_line_pt"] * 8),
                        color="000000",
                        space=1,
                    )
                )
    changes.append(
        dict(index=-1, action=f"配置 {len(doc.sections)} 个分节、页眉与页码")
    )


def format_tables(doc, template):
    config = template["tables"]
    for table in iter_tables(doc):
        if config["three_line"]:
            borders = set_child(table._tbl.tblPr, "w:tblBorders")
            for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
                borders.append(
                    element(
                        "w:" + side,
                        val="single" if side in ("top", "bottom") else "nil",
                        sz=round(config["outer_pt"] * 8),
                        color="000000",
                    )
                )
        seen = set()
        for ri, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            set_child(trpr, "w:tblHeader", val=int(ri == 0 and config["repeat_header"]))
            set_child(trpr, "w:cantSplit", val=int(config["prevent_row_split"]))
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                if config["three_line"]:
                    borders = set_child(cell._tc.get_or_add_tcPr(), "w:tcBorders")
                    for side in ("top", "bottom", "left", "right"):
                        width = (
                            config["outer_pt"]
                            if (ri == 0 and side == "top")
                            or (ri == len(table.rows) - 1 and side == "bottom")
                            else config["inner_pt"]
                        )
                        visible = (ri == 0 and side in ("top", "bottom")) or (
                            ri == len(table.rows) - 1 and side == "bottom"
                        )
                        borders.append(
                            element(
                                "w:" + side,
                                val="single" if visible else "nil",
                                sz=round(width * 8),
                                color="000000",
                            )
                        )
                for p in cell.paragraphs:
                    apply_style(p, "table_header" if ri == 0 else "table", template)


def format_images(doc, template):
    limit = Cm(template["figures"]["max_width_cm"])
    resized = 0
    for shape in doc.inline_shapes:
        if shape.width > limit:
            ratio = limit / shape.width
            shape.height, shape.width = round(shape.height * ratio), limit
            resized += 1
    return resized


def format_notes(doc, template):
    props = doc.settings.element.find(qn("w:footnotePr"))
    if props is None:
        props = element("w:footnotePr")
        doc.settings.element.append(props)
    set_child(props, "w:numFmt", val=template["notes"]["number_format"])
    set_child(props, "w:numRestart", val=template["notes"]["restart"])
    for part in doc.part.package.parts:
        key = (
            "footnote"
            if str(part.partname).endswith("/footnotes.xml")
            else "endnote"
            if str(part.partname).endswith("/endnotes.xml")
            else None
        )
        if key is None:
            continue
        from lxml import etree
        from docx.oxml import parse_xml

        # Paragraph expects python-docx's typed CT_P nodes, not plain lxml nodes.
        root = parse_xml(part.blob)
        spec = template["styles"][key]
        for note in root:
            if int(note.get(qn("w:id"), "-1")) <= 0:
                continue
            for node in note.iter(qn("w:p")):
                apply_style(Paragraph(node, doc._body), key, template)
            for r in note.iter(qn("w:r")):
                if r.find(qn("w:t")) is None:
                    continue
                rpr = r.find(qn("w:rPr"))
                if rpr is None:
                    rpr = element("w:rPr")
                    r.insert(0, rpr)
                set_child(
                    rpr,
                    "w:rFonts",
                    eastAsia=spec["font"],
                    ascii=spec["latin"],
                    hAnsi=spec["latin"],
                )
                set_child(rpr, "w:sz", val=round(spec["size"] * 2))
        if hasattr(part, "_element"):
            part._element = root
        else:
            part._blob = etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )


def create_skeleton(template, destination):
    from docx import Document
    from ..storage import save_document

    doc = Document()
    setup_styles(doc, template)
    meta = template["metadata"]
    doc.add_paragraph(meta["title"] or "【填写论文题目】", STYLE_NAMES["title"])
    doc.add_paragraph(meta["title_en"] or "【填写英文题目】", STYLE_NAMES["subtitle"])
    for label, value in [
        ("学校", template["school"]),
        ("学院", template["department"]),
        ("姓名", meta["author"]),
        ("学号", meta["student_id"]),
        ("导师", meta["supervisor"]),
        ("日期", meta["date"]),
    ]:
        doc.add_paragraph(
            f"{label}：{value or '【待填写】'}", STYLE_NAMES["cover_info"]
        )
    doc.add_paragraph(
        "【请按学校当年要求替换正式封面、原创性声明及使用授权书；本页为编辑占位。】"
    )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    for heading, text in [
        ("摘要", "【填写中文摘要】"),
        ("Abstract", "【填写英文摘要】"),
    ]:
        doc.add_paragraph(heading, STYLE_NAMES["front_heading"])
        doc.add_paragraph(text)
        doc.add_paragraph(
            ("关键词：" + (meta["keywords"] or "【关键词】"))
            if heading == "摘要"
            else "Keywords: " + (meta["keywords_en"] or "【keywords】")
        )
    for title in ("第1章 绪论", "第2章 方法", "第3章 结果与讨论", "第4章 结论"):
        doc.add_paragraph(title, "Heading 1")
        doc.add_paragraph("【在这里填写正文】")
    for heading in ("参考文献", "附录 A 补充材料", "符号表", "致谢", "学术成果"):
        doc.add_paragraph(heading, STYLE_NAMES["front_heading"])
        doc.add_paragraph("【按实际内容填写；不需要的章节可删除】")
    if template["front_matter_path"]:
        from docxcompose.composer import Composer
        from ..storage import ensure_distinct_paths

        ensure_distinct_paths(template["front_matter_path"], destination)
        front = Document(template["front_matter_path"])
        # Template text substitutions are explicit; school logos and layout stay in the source package.
        replacements = {
            **meta,
            "school": template["school"],
            "department": template["department"],
        }
        parts = list(front.paragraphs)
        for table in iter_tables(front):
            for row in table.rows:
                for cell in row.cells:
                    parts.extend(cell.paragraphs)
        for p in parts:
            if plain(p):
                text = p.text
                for key, value in replacements.items():
                    text = text.replace("{{" + key + "}}", value)
                replace_text_nodes(list(paragraph_runs(p)), text)
        for p in list(doc.paragraphs):
            if p.text == "摘要":
                break
            p._p.getparent().remove(p._p)
        front.add_page_break()
        composer = Composer(front)
        composer.append(doc)
        doc = composer.doc
    save_document(doc, destination)
    return destination
