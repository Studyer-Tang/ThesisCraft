"""Tables operations for the shared document engine."""

from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm
from .constants import RE_CURRENCY_PREFIX, RE_CURRENCY_SUFFIX, RE_NUMERIC_TABLE_TEXT
from .ooxml import iter_tables, paragraph_runs


class TableMixin:
    def _get_or_add_table_pr(self, table):
        tbl = table._tbl
        tbl_pr = tbl.tblPr
        if tbl_pr is None:
            tbl_pr = OxmlElement("w:tblPr")
            tbl.insert(0, tbl_pr)
        return tbl_pr

    def _set_table_borders(self, table, size_pt=0.5, color="000000"):
        size = max(1, int(float(size_pt) * 8))
        tbl_pr = self._get_or_add_table_pr(table)
        borders = tbl_pr.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            tbl_pr.append(borders)
        else:
            for child in list(borders):
                borders.remove(child)

        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            elem = OxmlElement(f"w:{edge}")
            elem.set(qn("w:val"), "single")
            elem.set(qn("w:sz"), str(size))
            elem.set(qn("w:space"), "0")
            elem.set(qn("w:color"), color)
            borders.append(elem)

    def _set_cell_borders(self, cell, size_pt=0.5, color="000000"):
        size = max(1, int(float(size_pt) * 8))
        tc = cell._tc
        tc_pr = tc.tcPr
        if tc_pr is None:
            tc_pr = OxmlElement("w:tcPr")
            tc.insert(0, tc_pr)

        borders = tc_pr.find(qn("w:tcBorders"))
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tc_pr.append(borders)
        else:
            for child in list(borders):
                borders.remove(child)

        for edge in ("top", "left", "bottom", "right"):
            elem = OxmlElement(f"w:{edge}")
            elem.set(qn("w:val"), "single")
            elem.set(qn("w:sz"), str(size))
            elem.set(qn("w:space"), "0")
            elem.set(qn("w:color"), color)
            borders.append(elem)

    def _set_table_cell_margins(
        self, table, top_cm=0.0, bottom_cm=0.0, left_cm=0.05, right_cm=0.05
    ):
        tbl_pr = self._get_or_add_table_pr(table)
        cell_mar = tbl_pr.find(qn("w:tblCellMar"))
        if cell_mar is None:
            cell_mar = OxmlElement("w:tblCellMar")
            tbl_pr.append(cell_mar)

        def set_side(tag, cm_value):
            node = cell_mar.find(qn(f"w:{tag}"))
            if node is None:
                node = OxmlElement(f"w:{tag}")
                cell_mar.append(node)
            node.set(qn("w:type"), "dxa")
            node.set(qn("w:w"), str(int(Cm(float(cm_value)).twips)))

        set_side("top", top_cm)
        set_side("bottom", bottom_cm)
        set_side("left", left_cm)
        set_side("right", right_cm)

    def _set_table_width_percent(self, table, percent=100):
        percent = max(1, min(100, int(float(percent))))
        tbl_pr = self._get_or_add_table_pr(table)
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:type"), "pct")
        tbl_w.set(qn("w:w"), str(percent * 50))

    def _set_table_indent(self, table, indent_twips=0):
        tbl_pr = self._get_or_add_table_pr(table)
        tbl_ind = tbl_pr.find(qn("w:tblInd"))
        if tbl_ind is None:
            tbl_ind = OxmlElement("w:tblInd")
            tbl_pr.append(tbl_ind)
        tbl_ind.set(qn("w:type"), "dxa")
        tbl_ind.set(qn("w:w"), str(int(indent_twips)))

    @staticmethod
    def _table_text_weight(text):
        weight = 0.0
        for ch in text:
            weight += 0.5 if ord(ch) < 128 else 1.0
        return weight

    @staticmethod
    def _normalize_table_pcts(weights, min_pct, max_pct):
        total = sum(weights) or 1.0
        pcts = [w / total * 100 for w in weights]
        for i, value in enumerate(pcts):
            if value < min_pct:
                pcts[i] = min_pct
            elif value > max_pct:
                pcts[i] = max_pct
        total = sum(pcts) or 1.0
        return [value / total * 100 for value in pcts]

    def _set_table_col_widths_by_content(self, table, min_pct=8, max_pct=45):
        if not table.rows:
            return
        col_count = max(len(row.cells) for row in table.rows)
        if col_count == 0:
            return

        min_pct = max(1.0, float(min_pct))
        max_pct = max(min_pct, float(max_pct))
        max_weights = [1.0] * col_count
        for row in table.rows:
            for col_idx, cell in enumerate(row.cells):
                text = "".join(p.text for p in cell.paragraphs).strip()
                if text:
                    max_weights[col_idx] = max(
                        max_weights[col_idx], self._table_text_weight(text)
                    )

        pcts = self._normalize_table_pcts(max_weights, min_pct, max_pct)
        tbl = table._tbl
        tbl_grid = tbl.tblGrid
        if tbl_grid is None:
            tbl_grid = OxmlElement("w:tblGrid")
            tbl.insert(0, tbl_grid)
        else:
            for child in list(tbl_grid):
                tbl_grid.remove(child)

        for pct in pcts:
            grid_col = OxmlElement("w:gridCol")
            grid_col.set(qn("w:w"), str(int(pct * 50)))
            tbl_grid.append(grid_col)

        for row in table.rows:
            for col_idx, cell in enumerate(row.cells):
                tc = cell._tc
                tc_pr = tc.tcPr
                if tc_pr is None:
                    tc_pr = OxmlElement("w:tcPr")
                    tc.insert(0, tc_pr)
                tc_w = tc_pr.find(qn("w:tcW"))
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.append(tc_w)
                tc_w.set(qn("w:type"), "pct")
                tc_w.set(qn("w:w"), str(int(pcts[col_idx] * 50)))

    @staticmethod
    def _is_numeric_table_text(text):
        text = (text or "").strip()
        if not text:
            return False
        text = text.replace(",", "").replace("，", "").replace("％", "%")
        text = RE_CURRENCY_PREFIX.sub("", text)
        text = RE_CURRENCY_SUFFIX.sub("", text)
        return RE_NUMERIC_TABLE_TEXT.match(text) is not None

    @staticmethod
    def _is_short_table_text(text, max_len=4):
        text = (text or "").strip()
        return 0 < len(text) <= int(max_len)

    def _format_tables(self, doc, apply_color=True):
        if not self.config.get("enable_table_formatting", False):
            self._log("表格自动调整未启用，跳过表格内容格式化。")
            return

        tables = list(iter_tables(doc))
        if not tables:
            self._log("未发现表格，跳过表格内容格式化。")
            return

        table_font = self.config.get(
            "table_font", self.config.get("body_font", "仿宋_GB2312")
        )
        table_header_font = self.config.get("table_header_font", table_font)
        table_size = self._config_float(
            self.config, "table_size", self.config.get("body_size", 12)
        )
        row_height_cm = self._config_float(self.config, "table_row_height_cm", 0.7)
        border_size_pt = self._config_float(self.config, "table_border_size_pt", 0.5)
        width_percent = self._config_float(self.config, "table_width_percent", 100)
        col_min_pct = self._config_float(self.config, "table_col_min_pct", 8)
        col_max_pct = self._config_float(self.config, "table_col_max_pct", 45)
        short_text_len = self._config_float(self.config, "table_short_text_len", 4)
        auto_col_width = self.config.get("table_auto_col_width", True)
        header_bold = self.config.get("table_header_bold", True)
        smart_align = self.config.get("table_smart_align", False)
        unified_borders = self.config.get("table_unified_borders", True)

        self._log(f"开始格式化表格内容（共 {len(tables)} 个）...")
        for table_idx, table in enumerate(tables, start=1):
            self._log(f"  > 表格 {table_idx}: 调整宽度、行高、字体和单元格格式")
            table.autofit = not auto_col_width
            self._set_table_width_percent(table, width_percent)
            self._set_table_indent(table, 0)
            self._set_table_cell_margins(table)
            if unified_borders:
                self._set_table_borders(table, size_pt=border_size_pt)
            if auto_col_width:
                self._set_table_col_widths_by_content(
                    table, min_pct=col_min_pct, max_pct=col_max_pct
                )

            serial_col_idx = None
            if table.rows:
                for col_idx, cell in enumerate(table.rows[0].cells):
                    head_text = "".join(p.text for p in cell.paragraphs).strip()
                    if "序号" in head_text or head_text == "序":
                        serial_col_idx = col_idx
                        break

            for row_idx, row in enumerate(table.rows):
                if row_height_cm > 0:
                    row.height = Cm(row_height_cm)
                    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

                for col_idx, cell in enumerate(row.cells):
                    if unified_borders:
                        self._set_cell_borders(cell, size_pt=border_size_pt)

                    cell_text = "".join(p.text for p in cell.paragraphs).strip()
                    for para in cell.paragraphs:
                        if para.text.strip():
                            for run in paragraph_runs(para):
                                font_name = (
                                    table_header_font if row_idx == 0 else table_font
                                )
                                self._set_run_font(
                                    run, font_name, table_size, set_color=apply_color
                                )
                                if row_idx == 0 and header_bold:
                                    run.font.bold = True

                        para.paragraph_format.first_line_indent = Pt(0)
                        para.paragraph_format.space_before = Pt(0)
                        para.paragraph_format.space_after = Pt(0)
                        self._apply_line_spacing(
                            para,
                            "table_line_spacing",
                            "table_line_spacing_unit",
                            "table_line_spacing_multiple",
                            22,
                        )

                        if smart_align:
                            if row_idx == 0:
                                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            elif "合计" in cell_text or "总计" in cell_text:
                                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            elif (
                                serial_col_idx is not None and col_idx == serial_col_idx
                            ):
                                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            elif self._is_numeric_table_text(cell_text):
                                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                            elif self._is_short_table_text(cell_text, short_text_len):
                                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            else:
                                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
