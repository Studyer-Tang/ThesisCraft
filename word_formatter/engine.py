"""Engine operations for the shared document engine."""

import os
from pathlib import Path
import shutil
import tempfile
import uuid
from docx import Document
from docx.document import Document as _Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from .constants import (
    BLANK_LINE_MODE_DELETE_SINGLE,
    BLANK_LINE_MODE_KEEP_SINGLE,
    BLANK_LINE_MODE_OPTIONS,
    BLANK_LINE_MODE_PRESERVE,
    RE_ATTACHMENT,
    RE_H2_INLINE_TITLE,
    RE_HEADING_H1,
    RE_HEADING_H2,
    RE_HEADING_H3,
    RE_HEADING_H4,
    RE_SAFE_FILENAME_CHARS,
)
from .conversion import WPSAppManager, SofficeConverter, LegacyConversionUnavailable
from .text import TextMixin
from .tables import TableMixin
from .layout import LayoutMixin
from .recognition import RecognitionMixin
from .config import normalize_config
from .storage import save_document, ensure_distinct_paths
from .ooxml import paragraph_runs, split_run_at


class WordProcessor(TextMixin, TableMixin, LayoutMixin, RecognitionMixin):
    def __init__(
        self,
        config,
        log_callback=None,
        remove_blank_lines=True,
        blank_line_mode=None,
        com_manager=None,
        soffice_path=None,
        soffice_timeout=120,
    ):
        self.config = normalize_config(config)
        self.temp_files = []
        self.sys_temp_dir = tempfile.gettempdir()
        self.log_callback = log_callback
        self.com_manager = com_manager or WPSAppManager(log_callback)
        self._owns_com_manager = com_manager is None
        self.soffice_path = soffice_path
        self.soffice_timeout = soffice_timeout
        self.soffice_converter = None
        self.blank_line_mode = self._normalize_blank_line_mode(
            blank_line_mode or self.config.get("blank_line_mode"),
            remove_blank_lines=remove_blank_lines,
        )
        self.remove_blank_lines = self.blank_line_mode == BLANK_LINE_MODE_DELETE_SINGLE

    def _log(self, message):
        if self.log_callback:
            self.log_callback(message)

    @staticmethod
    def _com_available():
        return WPSAppManager._com_available()

    @staticmethod
    def _com_unavailable_message(file_ext=None):
        return WPSAppManager._com_unavailable_message(file_ext)

    @staticmethod
    def _paragraph_has_ooxml(para, tag_name):
        return para._p.find(".//" + qn(tag_name)) is not None

    @classmethod
    def _has_field_codes(cls, para):
        return cls._paragraph_has_ooxml(para, "w:fldChar") or cls._paragraph_has_ooxml(
            para, "w:instrText"
        )

    @classmethod
    def _has_drawing_or_pict(cls, para):
        return cls._paragraph_has_ooxml(para, "w:drawing") or cls._paragraph_has_ooxml(
            para, "w:pict"
        )

    @classmethod
    def _has_embedded_object(cls, para):
        return cls._paragraph_has_ooxml(para, "w:object")

    @staticmethod
    def _normalize_blank_line_mode(mode, remove_blank_lines=True):
        if mode in BLANK_LINE_MODE_OPTIONS:
            return mode
        if mode in ("preserve", "none"):
            return BLANK_LINE_MODE_PRESERVE
        if mode in ("delete_single", "remove_single"):
            return BLANK_LINE_MODE_DELETE_SINGLE
        if mode in ("keep_single", "compress"):
            return BLANK_LINE_MODE_KEEP_SINGLE
        if isinstance(mode, bool):
            return (
                BLANK_LINE_MODE_DELETE_SINGLE if mode else BLANK_LINE_MODE_KEEP_SINGLE
            )
        return (
            BLANK_LINE_MODE_DELETE_SINGLE
            if remove_blank_lines
            else BLANK_LINE_MODE_KEEP_SINGLE
        )

    def _cleanup_temp_files(self):
        if not self.temp_files:
            return
        self._log("正在清理本轮临时文件...")
        for f in self.temp_files:
            try:
                if os.path.exists(f):
                    os.remove(f)
                    self._log(f"  > 临时文件 {os.path.basename(f)} 已删除")
            except OSError as e:
                self._log(f"  > 警告：删除临时文件 {f} 失败: {e}")
        self.temp_files.clear()

    def _make_temp_docx_path(self, prefix, base_name):
        safe_base_name = RE_SAFE_FILENAME_CHARS.sub("_", base_name).strip(" ._")
        safe_base_name = (safe_base_name or "document")[:80]
        temp_name = (
            f"~temp_{prefix}_{safe_base_name}_{os.getpid()}_{uuid.uuid4().hex[:8]}.docx"
        )
        temp_path = os.path.join(self.sys_temp_dir, temp_name)
        self.temp_files.append(temp_path)
        return temp_path

    def _get_wps_app(self):
        return self.com_manager.get_app()

    def _get_soffice_converter(self):
        if self.soffice_converter is None:
            self.soffice_converter = SofficeConverter(
                self.soffice_path, self.soffice_timeout
            )
        return self.soffice_converter

    def _convert_legacy_with_com(self, input_path, temp_docx_path):
        app = self._get_wps_app()
        doc_com = None
        try:
            doc_com = app.Documents.Open(os.path.abspath(input_path), ReadOnly=1)
            doc_com.SaveAs2(os.path.abspath(temp_docx_path), FileFormat=12)
        finally:
            if doc_com is not None:
                doc_com.Close()
        self._log("文件格式转换完成。")

    def _convert_legacy_with_soffice(self, input_path, temp_docx_path):
        converter = self._get_soffice_converter()
        if not converter.available:
            raise LegacyConversionUnavailable(
                f"{self._com_unavailable_message(os.path.splitext(input_path)[1].lower())} "
                "已跳过该旧格式文件，继续处理其他可支持文件。"
            )

        converted_path, work_dir = converter.convert_to_docx(input_path, self._log)
        try:
            shutil.copy2(converted_path, temp_docx_path)
        finally:
            shutil.rmtree(work_dir, ignore_errors=True)
        self._log("LibreOffice 文件格式转换完成。")

    def quit_com_app(self):
        if self._owns_com_manager:
            self.com_manager.quit()

    def convert_to_docx(self, input_path):
        file_ext = os.path.splitext(input_path)[1].lower()
        is_from_txt = file_ext in (".txt", ".md")
        base_name = os.path.splitext(os.path.basename(input_path))[0]

        if file_ext == ".docx":
            self._log("检测到 .docx 文件，正在创建安全的处理副本...")
            temp_docx_path = self._make_temp_docx_path("copy", base_name)
            shutil.copy2(input_path, temp_docx_path)
            self._log(f"  > 副本创建成功: {os.path.basename(temp_docx_path)}")
            return temp_docx_path, False

        temp_docx_path = self._make_temp_docx_path("converted", base_name)

        if file_ext == ".txt":
            self._log("检测到 .txt 文件，正在创建 .docx...")
            text_content = self._read_text_file(input_path)
            text_content = self._normalize_text_blank_lines(text_content)
            self._log_blank_line_mode("TXT")
            doc = Document()
            for line in text_content.split("\n"):
                doc.add_paragraph(line.strip())
            doc.save(temp_docx_path)
            self._log("TXT转换完成。")
            return temp_docx_path, is_from_txt
        elif file_ext == ".md":
            self._log("检测到 .md 文件，正在清理 Markdown 标记并创建 .docx...")
            raw_text = self._read_text_file(input_path)
            cleaned_text = self._clean_markdown(raw_text)
            cleaned_text = self._normalize_text_blank_lines(cleaned_text)
            self._log_blank_line_mode("Markdown 文本")
            doc = Document()
            for line in cleaned_text.split("\n"):
                doc.add_paragraph(line.strip())
            doc.save(temp_docx_path)
            self._log("Markdown 转换完成。")
            return temp_docx_path, is_from_txt
        elif file_ext in [".wps", ".doc"]:
            self._log(f"正在转换 {file_ext} 文件为 .docx...")
            if self._com_available():
                try:
                    self._convert_legacy_with_com(input_path, temp_docx_path)
                    return temp_docx_path, is_from_txt
                except Exception as exc:
                    self._log(
                        f"  > WPS/Word 转换失败，尝试使用 LibreOffice 兜底: {exc}"
                    )
            else:
                self._log(f"  > {self._com_unavailable_message(file_ext)}")

            self._convert_legacy_with_soffice(input_path, temp_docx_path)
            return temp_docx_path, is_from_txt

        raise ValueError(f"不支持的文件格式: {file_ext}")

    def _preprocess_com_tasks(self, docx_path):
        if not self._com_available():
            self._log(f"  > {self._com_unavailable_message()}")
            return
        self._log("正在对副本执行预处理（接受所有修订、转换自动编号）...")
        doc_com = None
        try:
            app = self._get_wps_app()
            doc_com = app.Documents.Open(os.path.abspath(docx_path))

            doc_com.TrackRevisions = False
            self._log("  > 已关闭修订追踪。")

            if doc_com.Revisions.Count > 0:
                doc_com.AcceptAllRevisions()
                self._log("  > 已接受文档副本中的所有修订。")

            doc_com.Content.ListFormat.ConvertNumbersToText()
            self._log("  > 已将副本中的自动编号转换为文本。")

            if doc_com.Revisions.Count > 0:
                doc_com.AcceptAllRevisions()
                self._log("  > 已接受编号转换产生的修订。")

            doc_com.TrackRevisions = False

            doc_com.Save()
            self._log("预处理完成。")
        except Exception as e:
            self._log(f"警告：执行预处理任务时出错: {e}")
        finally:
            if doc_com is not None:
                try:
                    doc_com.Close()
                except Exception as e:
                    self._log(f"  > 警告：关闭预处理文档时发生异常: {e}")

    def _iter_block_items(self, parent):
        parent_elm = (
            parent.element.body if isinstance(parent, _Document) else parent._tc
        )
        for child in parent_elm.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield Table(child, parent)

    def format_document(self, input_path, output_path, *, overwrite=False):
        ensure_distinct_paths(input_path, output_path)
        if Path(output_path).exists() and not overwrite:
            raise FileExistsError(f"输出文件已存在: {output_path}")
        self._overwrite = overwrite
        try:
            return self._format_document(input_path, output_path)
        finally:
            self._cleanup_temp_files()
            self.quit_com_app()

    def _format_document(self, input_path, output_path):
        processing_path, is_from_txt = self.convert_to_docx(input_path)
        if not is_from_txt and self.config["preprocess_office"]:
            self._preprocess_com_tasks(processing_path)

        doc = Document(processing_path)

        if self.config.get("normalize_punctuation", False):
            symbol_changes = self._normalize_document_symbols(doc)
            self._log(f"符号标准化完成，共修复 {symbol_changes} 个段落/表格单元格。")

        all_blocks = list(self._iter_block_items(doc))
        processed_indices = set()

        apply_color = self.config["force_black_text"]

        if not is_from_txt:
            self._log("正在扫描图表标题...")
            for idx, block in enumerate(all_blocks):
                is_pic_para = isinstance(
                    block, Paragraph
                ) and self._has_drawing_or_pict(block)
                is_table = isinstance(block, Table)

                if not (is_pic_para or is_table):
                    continue

                for direction in [-1, 1]:
                    caption_found = False
                    for i in range(
                        idx + direction,
                        -1 if direction == -1 else len(all_blocks),
                        direction,
                    ):
                        if i in processed_indices:
                            continue
                        potential_caption = all_blocks[i]
                        if not isinstance(potential_caption, Paragraph):
                            break
                        text = potential_caption.text.strip()
                        if text:
                            if self._get_paragraph_alignment(
                                potential_caption
                            ) == WD_ALIGN_PARAGRAPH.CENTER and (
                                text.startswith("图") or text.startswith("表")
                            ):
                                detected_type = "图" if text.startswith("图") else "表"
                                self._log(
                                    f'  > 发现 {detected_type} 的标题: "{text[:30]}..." (在段落 {i + 1})'
                                )
                                config_font_key = f"{('figure' if detected_type == '图' else 'table')}_caption_font"
                                config_size_key = f"{('figure' if detected_type == '图' else 'table')}_caption_size"
                                config_font = self.config[config_font_key]
                                config_size = self.config[config_size_key]
                                self._apply_font_to_runs(
                                    potential_caption,
                                    config_font,
                                    config_size,
                                    set_color=apply_color,
                                )
                                processed_indices.add(i)
                                caption_found = True
                            break
                    if caption_found:
                        break

        # 查找主标题和副标题
        title_indices, subtitle_indices = self._find_title_and_subtitle_paragraphs(
            doc, is_from_txt
        )

        # 将标题和副标题索引加入已处理集合
        for idx in title_indices:
            processed_indices.add(idx)
        for idx in subtitle_indices:
            processed_indices.add(idx)

        self._log("预扫描完成，开始逐段格式化...")
        if self.config["set_outline"]:
            self._log("【大纲级别设置已启用】")
        else:
            self._log("【大纲级别设置已禁用】")

        self._format_title_group(all_blocks, title_indices, "title", apply_color)
        self._format_title_group(all_blocks, subtitle_indices, "subtitle", apply_color)

        block_idx = 0
        while block_idx < len(all_blocks):
            block = all_blocks[block_idx]

            if block_idx in processed_indices:
                if block_idx not in title_indices and block_idx not in subtitle_indices:
                    self._log(f"块 {block_idx + 1}: 已作为图表/附件标题处理 - 跳过")
                block_idx += 1
                continue

            current_block_num = block_idx + 1
            if isinstance(block, Table):
                self._log(f"块 {current_block_num}: 表格 - 跳过")
                block_idx += 1
                continue

            para = block
            if not para.text.strip():
                self._log(f"段落 {current_block_num}: 空白 - 跳过")
                block_idx += 1
                continue

            is_pic = self._has_drawing_or_pict(para)
            is_embedded_obj = self._has_embedded_object(para)
            if is_pic or is_embedded_obj:
                log_msg = "图片" if is_pic else "嵌入对象"
                self._log(f"段落 {current_block_num}: {log_msg} - 仅格式化文字")

                text_to_check = para.text.lstrip()
                para_text_preview = text_to_check[:30].replace("\n", " ")

                if RE_HEADING_H1.match(text_to_check):
                    self._log(f'  > 文字识别为一级标题: "{para_text_preview}..."')
                    self._apply_font_to_runs(
                        para,
                        self.config["h1_font"],
                        self.config["h1_size"],
                        set_color=apply_color,
                    )
                    self._apply_bold_to_runs(para, self.config.get("h1_bold", False))
                elif RE_HEADING_H2.match(text_to_check):
                    self._log(f'  > 文字识别为二级标题: "{para_text_preview}..."')
                    self._apply_font_to_runs(
                        para,
                        self.config["h2_font"],
                        self.config["h2_size"],
                        set_color=apply_color,
                    )
                    self._apply_bold_to_runs(para, self.config.get("h2_bold", False))
                elif RE_HEADING_H3.match(text_to_check):
                    self._log(f'  > 文字识别为三级标题: "{para_text_preview}..."')
                    self._apply_font_to_runs(
                        para,
                        self.config["body_font"],
                        self.config["body_size"],
                        set_color=apply_color,
                    )
                elif RE_HEADING_H4.match(text_to_check):
                    self._log(f'  > 文字识别为四级标题: "{para_text_preview}..."')
                    self._apply_font_to_runs(
                        para,
                        self.config["body_font"],
                        self.config["body_size"],
                        set_color=apply_color,
                    )
                elif text_to_check:
                    self._log(f'  > 文字识别为正文: "{para_text_preview}..."')
                    self._apply_font_to_runs(
                        para,
                        self.config["body_font"],
                        self.config["body_size"],
                        set_color=apply_color,
                    )

                block_idx += 1
                continue

            original_text, text_to_check = para.text, para.text.lstrip()
            text_to_check_stripped = para.text.strip()
            leading_space_count = len(original_text) - len(text_to_check)
            para_text_preview = text_to_check[:30].replace("\n", " ")

            spacing = para._p.get_or_add_pPr().get_or_add_spacing()
            spacing.set(qn("w:beforeAutospacing"), "0")
            spacing.set(qn("w:afterAutospacing"), "0")
            para.paragraph_format.space_before, para.paragraph_format.space_after = (
                Pt(0),
                Pt(0),
            )
            self._apply_line_spacing(
                para,
                "line_spacing",
                "line_spacing_unit",
                "line_spacing_multiple",
                28,
            )

            is_attachment_enabled = self.config.get(
                "enable_attachment_formatting", False
            )
            is_attachment_candidate = False
            if is_from_txt:
                if RE_ATTACHMENT.match(text_to_check_stripped):
                    is_attachment_candidate = True
            elif self._get_paragraph_alignment(para) in [
                WD_ALIGN_PARAGRAPH.LEFT,
                WD_ALIGN_PARAGRAPH.JUSTIFY,
                None,
            ] and RE_ATTACHMENT.match(text_to_check_stripped):
                is_attachment_candidate = True

            if is_attachment_enabled and is_attachment_candidate:
                self._log(
                    f'段落 {current_block_num}: 附件标识 - "{para_text_preview}..."'
                )
                self._strip_leading_whitespace(para)
                self._apply_font_to_runs(
                    para,
                    self.config["attachment_font"],
                    self.config["attachment_size"],
                    set_color=apply_color,
                )
                self._reset_pagination_properties(para)
                para.paragraph_format.page_break_before = True
                para.paragraph_format.left_indent = Pt(0)
                para.paragraph_format.first_line_indent = None

                ind = para._p.get_or_add_pPr().get_or_add_ind()
                ind.set(qn("w:firstLineChars"), "0")

                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                self._format_heading(para, 1)

                # 查找并格式化附件的标题和副标题
                search_idx = block_idx + 1

                # 查找附件的标题和副标题
                att_title_indices, att_subtitle_indices = (
                    self._find_title_and_subtitle_paragraphs(
                        doc, is_from_txt, search_idx
                    )
                )

                # 将附件的标题和副标题加入已处理集合
                for idx in att_title_indices:
                    processed_indices.add(idx)
                for idx in att_subtitle_indices:
                    processed_indices.add(idx)

                self._format_title_group(
                    all_blocks, att_title_indices, "title", apply_color, outline=True
                )
                self._format_title_group(
                    all_blocks, att_subtitle_indices, "subtitle", apply_color
                )

                # 计算下一个要处理的块索引。没有附件标题/副标题时，
                # 只跳过附件标识本段，避免漏处理紧随其后的正文段落。
                handled_indices = att_title_indices + att_subtitle_indices
                if handled_indices:
                    next_idx = max(handled_indices) + 1
                else:
                    next_idx = block_idx + 1

                block_idx = next_idx
                continue

            elif RE_HEADING_H1.match(text_to_check):
                self._log(
                    f'段落 {current_block_num}: 一级标题 - "{para_text_preview}..."'
                )
                self._strip_leading_whitespace(para)
                self._format_heading(para, 1)
                self._apply_font_to_runs(
                    para,
                    self.config["h1_font"],
                    self.config["h1_size"],
                    set_color=apply_color,
                )
                self._apply_bold_to_runs(para, self.config.get("h1_bold", False))
                self._apply_text_indent_and_align(para)
                self._reset_pagination_properties(para)

            elif RE_HEADING_H2.match(text_to_check):
                self._log(
                    f'段落 {current_block_num}: 二级标题 - "{para_text_preview}..."'
                )
                self._strip_leading_whitespace(para)

                parts = para.text.split("。", 1)

                if len(parts) == 2 and parts[1].strip():
                    self._log("  > 检测到二级标题与正文在同一段落，执行段内格式拆分。")
                    title_len = len(parts[0]) + 1

                    position = 0
                    for run in list(paragraph_runs(para)):
                        length = len(run.text)
                        right = None
                        if position < title_len < position + length:
                            right = split_run_at(run, title_len - position)
                        is_title = position < title_len
                        self._set_run_font(
                            run,
                            self.config["h2_font" if is_title else "body_font"],
                            self.config["h2_size" if is_title else "body_size"],
                            set_color=apply_color,
                        )
                        if right is not None:
                            self._set_run_font(
                                right,
                                self.config["body_font"],
                                self.config["body_size"],
                                set_color=apply_color,
                            )
                        position += length

                    self._format_heading(para, 2)
                    self._apply_bold_to_first_chars(
                        para, title_len, self.config.get("h2_bold", False)
                    )
                    self._apply_text_indent_and_align(para)
                    self._reset_pagination_properties(para)

                else:
                    match = RE_H2_INLINE_TITLE.match(text_to_check)
                    if match and not (
                        text_to_check.startswith("（")
                        and text_to_check.strip().endswith("）")
                    ):
                        self._log("  > 已将二级标题的括号统一为中文括号。")
                        for r in paragraph_runs(para):
                            for node in r._r.findall(qn("w:t")):
                                node.text = (
                                    (node.text or "")
                                    .replace("(", "（", 1)
                                    .replace(")", "）", 1)
                                )
                    self._format_heading(para, 2)
                    self._apply_font_to_runs(
                        para,
                        self.config["h2_font"],
                        self.config["h2_size"],
                        set_color=apply_color,
                    )
                    self._apply_bold_to_runs(para, self.config.get("h2_bold", False))
                    self._apply_text_indent_and_align(para)
                    self._reset_pagination_properties(para)

            elif RE_HEADING_H3.match(text_to_check):
                self._log(
                    f'段落 {current_block_num}: 三级标题 - "{para_text_preview}..."'
                )
                self._strip_leading_whitespace(para)
                self._format_heading(para, 3)
                self._apply_font_to_runs(
                    para,
                    self.config["body_font"],
                    self.config["body_size"],
                    set_color=apply_color,
                )
                self._apply_text_indent_and_align(para)
                self._reset_pagination_properties(para)

            elif RE_HEADING_H4.match(text_to_check):
                self._log(
                    f'段落 {current_block_num}: 四级标题 - "{para_text_preview}..."'
                )
                self._strip_leading_whitespace(para)
                self._format_heading(para, 4)
                self._apply_font_to_runs(
                    para,
                    self.config["body_font"],
                    self.config["body_size"],
                    set_color=apply_color,
                )
                self._apply_text_indent_and_align(para)
                self._reset_pagination_properties(para)

            elif not is_from_txt:
                para_alignment = self._get_paragraph_alignment(para)
                if para_alignment in [
                    WD_ALIGN_PARAGRAPH.CENTER,
                    WD_ALIGN_PARAGRAPH.RIGHT,
                ]:
                    align_text = (
                        "居中"
                        if para_alignment == WD_ALIGN_PARAGRAPH.CENTER
                        else "右对齐"
                    )
                    self._log(
                        f"段落 {current_block_num}: {align_text}正文 - 保留原对齐"
                    )
                    self._apply_font_to_runs(
                        para,
                        self.config["body_font"],
                        self.config["body_size"],
                        set_color=apply_color,
                    )
                    self._reset_pagination_properties(para)
                elif leading_space_count > 5:
                    self._log(
                        f'段落 {current_block_num}: 正文 (保留前导空格) - "{para_text_preview}..."'
                    )
                    self._apply_font_to_runs(
                        para,
                        self.config["body_font"],
                        self.config["body_size"],
                        set_color=apply_color,
                    )
                    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    self._reset_pagination_properties(para)
                elif (
                    para.paragraph_format.first_line_indent is None
                    or para.paragraph_format.first_line_indent.pt == 0
                ) and leading_space_count == 0:
                    self._log(
                        f'段落 {current_block_num}: 正文 (保留0缩进) - "{para_text_preview}..."'
                    )
                    self._apply_font_to_runs(
                        para,
                        self.config["body_font"],
                        self.config["body_size"],
                        set_color=apply_color,
                    )
                    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    self._reset_pagination_properties(para)
                else:
                    self._log(
                        f'段落 {current_block_num}: 正文 (应用标准缩进) - "{para_text_preview}..."'
                    )
                    self._strip_leading_whitespace(para)
                    self._apply_font_to_runs(
                        para,
                        self.config["body_font"],
                        self.config["body_size"],
                        set_color=apply_color,
                    )
                    self._apply_text_indent_and_align(para)
                    self._reset_pagination_properties(para)
            else:
                self._log(
                    f'段落 {current_block_num}: 正文 (源自TXT，强制缩进) - "{para_text_preview}..."'
                )
                self._strip_leading_whitespace(para)
                self._apply_font_to_runs(
                    para,
                    self.config["body_font"],
                    self.config["body_size"],
                    set_color=apply_color,
                )
                self._apply_text_indent_and_align(para)
                self._reset_pagination_properties(para)

            block_idx += 1

        self._format_tables(doc, apply_color=apply_color)
        self._apply_page_setup(doc, is_from_txt=is_from_txt)
        self._log("正在保存最终文档...")
        save_document(doc, output_path, overwrite=self._overwrite)
        return str(Path(output_path).resolve())
