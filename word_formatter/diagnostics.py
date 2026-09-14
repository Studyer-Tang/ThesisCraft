"""Explicit, synthetic checks that also run from the packaged Windows program."""

import argparse
import json
from pathlib import Path
import tempfile
import time

from docx import Document
from .engine import WordProcessor
from .version import __version__, AUTHOR, APP_NAME


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="将本机启动与排版诊断写入指定 JSON 文件"
    )
    parser.add_argument("report")
    parser.add_argument("--office", choices=("word", "wps"))
    parser.add_argument('--academic', action='store_true')
    args = parser.parse_args(argv)
    report = {"name": APP_NAME, "version": __version__, "author": AUTHOR}
    try:
        if args.office:
            import pythoncom
            import win32com.client
            from .office_toolbar import OfficeToolbar

            pythoncom.CoInitialize()
            app = win32com.client.DispatchEx(
                "Word.Application" if args.office == "word" else "KWPS.Application"
            )
            toolbar = None
            documents = []
            try:
                with tempfile.TemporaryDirectory(prefix="wfp_diagnostics_") as temp:
                    source = Path(temp) / "diagnostic.docx"
                    doc = Document()
                    doc.add_paragraph("排版验证标题")
                    doc.add_paragraph("（一）标题。正文")
                    doc.save(source)
                    original = source.read_bytes()
                    errors = []
                    try:
                        documents.append(app.Documents.Open(str(source), ReadOnly=True))
                        toolbar = OfficeToolbar(app, args.office, notify=errors.append)
                        toolbar.buttons[0].Execute()
                        deadline = time.monotonic() + 60
                        while (
                            not errors
                            and toolbar.last_result is None
                            and time.monotonic() < deadline
                        ):
                            pythoncom.PumpWaitingMessages()
                            toolbar.poll()
                            time.sleep(0.05)
                        if errors:
                            raise RuntimeError(str(errors))
                        if (
                            toolbar.last_result is None
                            or not toolbar.last_result.outputs
                        ):
                            raise RuntimeError("工具栏处理没有返回结果。")
                        output = Path(toolbar.last_result.outputs[0])
                        opened = app.ActiveDocument
                        if Path(opened.FullName) != output:
                            raise RuntimeError("宿主没有打开排版结果。")
                        documents.append(opened)
                        if original != source.read_bytes():
                            raise RuntimeError("原件发生变化。")
                        report.update(
                            host=args.office,
                            button_event=True,
                            format_and_open=True,
                            original_unchanged=True,
                        )
                    finally:
                        if toolbar:
                            toolbar.close()
                        for document in reversed(documents):
                            document.Close(SaveChanges=0)
            finally:
                if app.Documents.Count == 0:
                    app.Quit()
                pythoncom.CoUninitialize()
        elif args.academic:
            from .gui import _create_root
            from .academic.gui import AcademicWindow
            from .academic.templates import load_template
            from .academic.workflow import run
            root, _ = _create_root()
            root.withdraw()
            try:
                view = AcademicWindow(root)
                root.update_idletasks()
                view.collect()
                report.update(window_title=root.title(), academic_fields=len(view.vars)+len(view.style_vars))
                with tempfile.TemporaryDirectory(prefix='paper_studio_diagnostic_') as temp:
                    source=Path(temp)/'source.docx'
                    doc=Document();doc.add_paragraph('测试论文');doc.add_paragraph('第1章 绪论','Heading 1')
                    doc.add_paragraph('正文 {{cite:sample}}。');doc.add_paragraph('参考文献');doc.save(source)
                    references=Path(temp)/'refs.json'
                    references.write_text(json.dumps([dict(key='sample',type='journal',title='Test article',authors=['Test, A'],year='2024',journal='Test Journal')]),encoding='utf-8')
                    result=run(source,load_template('pku-master'),reference_path=references)
                    report['academic_format']=bool(result['output'])
                    report['integrity']=result['integrity']['passed']
            finally:
                root.destroy()
        else:
            from .gui import _create_root, WordFormatterGUI

            root, _ = _create_root()
            root.withdraw()
            try:
                view = WordFormatterGUI(root)
                root.update_idletasks()
                report["window_title"] = root.title()
                report["gui_fields"] = len(view.entries)
                view.collect_config()
                with tempfile.TemporaryDirectory(prefix="wfp_diagnostics_") as temp:
                    source = Path(temp) / "source.txt"
                    source.write_text("测试标题\n正文", encoding="utf-8")
                    output = Path(temp) / "output.docx"
                    WordProcessor({}).format_document(source, output)
                    report["format"] = bool(Document(output).paragraphs)
            finally:
                root.destroy()
        report["success"] = True
    except Exception as exc:
        report.update(success=False, error=str(exc))
    destination = Path(args.report).resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return 0 if report["success"] else 1
