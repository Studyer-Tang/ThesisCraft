"""Layout operations for the shared document engine."""

import re
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

from .ooxml import paragraph_runs


class LayoutMixin:
    def _format_title_group(self, blocks, indices, kind, apply_color, outline=False):
        for index in indices:
            para = blocks[index]
            self._strip_leading_whitespace(para)
            self._apply_font_to_runs(
                para,
                self.config[kind + "_font"],
                self.config[kind + "_size"],
                set_color=apply_color,
            )
            self._apply_bold_to_runs(para, self.config.get(kind + "_bold", False))
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = None
            indent = para._p.get_or_add_pPr().get_or_add_ind()
            for key in (
                "w:firstLine",
                "w:firstLineChars",
                "w:hanging",
                "w:hangingChars",
            ):
                indent.attrib.pop(qn(key), None)
            spacing = para._p.get_or_add_pPr().get_or_add_spacing()
            spacing.set(qn("w:beforeAutospacing"), "0")
            spacing.set(qn("w:afterAutospacing"), "0")
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
            self._apply_line_spacing(
                para,
                kind + "_line_spacing",
                kind + "_line_spacing_unit",
                kind + "_line_spacing_multiple",
                33,
            )
            self._reset_pagination_properties(para)
            if outline:
                self._format_heading(para, 1)

    def _create_page_number(self, paragraph, text):
        font_name = self.config["page_number_font"]
        font_size = self.config["page_number_size"]
        self._set_run_font(
            paragraph.add_run("— "), font_name, font_size, set_color=True
        )
        run_field = paragraph.add_run()
        self._set_run_font(run_field, font_name, font_size, set_color=True)
        fldChar1 = OxmlElement("w:fldChar")
        fldChar1.set(qn("w:fldCharType"), "begin")
        instrText = OxmlElement("w:instrText")
        instrText.set(qn("xml:space"), "preserve")
        instrText.text = text
        fldChar2 = OxmlElement("w:fldChar")
        fldChar2.set(qn("w:fldCharType"), "end")
        run_field._r.extend([fldChar1, instrText, fldChar2])
        self._set_run_font(
            paragraph.add_run(" —"), font_name, font_size, set_color=True
        )

    def _apply_page_setup(self, doc, is_from_txt=False):
        should_set_a4 = is_from_txt or self.config["force_a4"]
        odd_even = self.config["page_number_align"] == "奇偶分页"
        if self.config["page_numbers"]:
            doc.settings.odd_and_even_pages_header_footer = odd_even
        for section in doc.sections:
            for name in ("top", "bottom", "left", "right"):
                setattr(section, name + "_margin", Cm(self.config["margin_" + name]))
            section.footer_distance = Cm(self.config["footer_distance"])
            if should_set_a4:
                section.page_width, section.page_height = Cm(21), Cm(29.7)
            if not self.config["page_numbers"]:
                continue
            footers = [
                (
                    section.footer,
                    WD_ALIGN_PARAGRAPH.RIGHT if odd_even else WD_ALIGN_PARAGRAPH.CENTER,
                )
            ]
            if odd_even:
                footers.append((section.even_page_footer, WD_ALIGN_PARAGRAPH.LEFT))
            if section.different_first_page_header_footer:
                footers.append((section.first_page_footer, WD_ALIGN_PARAGRAPH.CENTER))
            for footer, alignment in footers:
                page_para = None
                for paragraph in footer.paragraphs:
                    instructions = [
                        n.text or "" for n in paragraph._p.iter(qn("w:instrText"))
                    ]
                    instructions += [
                        n.get(qn("w:instr"), "")
                        for n in paragraph._p.iter(qn("w:fldSimple"))
                    ]
                    if any(re.search(r"\bPAGE\b", text, re.I) for text in instructions):
                        page_para = paragraph
                        break
                if page_para is None:
                    empty = next(
                        (p for p in footer.paragraphs if not p.text and len(p._p) == 0),
                        None,
                    )
                    page_para = empty if empty is not None else footer.add_paragraph()
                    self._create_page_number(page_para, "PAGE")
                page_para.alignment = alignment
                self._apply_font_to_runs(
                    page_para,
                    self.config["page_number_font"],
                    self.config["page_number_size"],
                )

    def _set_run_font(self, run, font_name, size_pt, set_color=False):
        run.font.size = Pt(size_pt)
        if set_color:
            run.font.color.rgb = RGBColor(0, 0, 0)
        rPr = run._r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        for theme_attr in (
            "w:eastAsiaTheme",
            "w:asciiTheme",
            "w:hAnsiTheme",
            "w:cstheme",
            "w:csTheme",
        ):
            rFonts.attrib.pop(qn(theme_attr), None)
        rFonts.set(qn("w:eastAsia"), font_name)
        # 根据配置决定西文字体（数字、字母）
        en_font = (
            self.config.get("english_font")
            if self.config.get("use_custom_english_font", False)
            else font_name
        )
        en_font = en_font or font_name
        run.font.name = en_font
        rFonts.set(qn("w:ascii"), en_font)
        rFonts.set(qn("w:hAnsi"), en_font)

    def _apply_font_to_runs(self, para, font_name, size_pt, set_color=False):
        for run in paragraph_runs(para):
            self._set_run_font(run, font_name, size_pt, set_color=set_color)

    def _get_paragraph_font_info(self, para):
        run = next((r for r in paragraph_runs(para) if r.text.strip()), None)
        if run is None:
            return None, None
        fonts = [run.font]
        for starting_style in (run.style, para.style):
            style, seen = starting_style, set()
            while style is not None and style.style_id not in seen:
                seen.add(style.style_id)
                fonts.append(style.font)
                style = style.base_style
        name = next((f.name for f in fonts if f.name), None)
        size = next((f.size.pt for f in fonts if f.size), None)
        return name, size

    @staticmethod
    def _get_paragraph_alignment(para):
        try:
            if para.alignment is not None:
                return para.alignment
        except ValueError:
            p_pr = para._p.pPr
            jc = p_pr.jc if p_pr is not None else None
            raw = jc.get(qn("w:val")) if jc is not None else None
            return {
                "start": WD_ALIGN_PARAGRAPH.LEFT,
                "end": WD_ALIGN_PARAGRAPH.RIGHT,
                "both": WD_ALIGN_PARAGRAPH.JUSTIFY,
                "distribute": WD_ALIGN_PARAGRAPH.JUSTIFY,
            }.get(raw)
        style, seen = para.style, set()
        while style is not None and style.style_id not in seen:
            seen.add(style.style_id)
            if style.paragraph_format.alignment is not None:
                return style.paragraph_format.alignment
            style = style.base_style
        return None

    def _strip_leading_whitespace(self, para):
        for run in list(paragraph_runs(para)):
            for child in run._r:
                if child.tag == qn("w:rPr"):
                    continue
                if child.tag != qn("w:t"):
                    return
                text = child.text or ""
                child.text = text.lstrip()
                if child.text:
                    return
            # Empty ordinary runs can be removed, but never special nodes.
            if all(c.tag in (qn("w:rPr"), qn("w:t")) for c in run._r) and not run.text:
                run._r.getparent().remove(run._r)

    def _reset_pagination_properties(self, para):
        if self.config["preserve_pagination"]:
            return
        para.paragraph_format.widow_control = False
        para.paragraph_format.keep_with_next = False
        para.paragraph_format.page_break_before = False
        para.paragraph_format.keep_together = False

    def _get_outline_level(self, para):
        """
        读取段落的当前大纲级别
        返回: 0-8 表示级别1-9，None 表示未设置
        """
        pPr = para._p.get_or_add_pPr()
        outlineLvl = pPr.find(qn("w:outlineLvl"))
        if outlineLvl is not None:
            val = outlineLvl.get(qn("w:val"))
            if val is not None:
                return int(val)
        return None

    def _set_outline_level(self, para, level):
        """
        直接设置段落的大纲级别，不通过样式，不影响字体字号等格式
        level: 1-9 的整数，表示大纲级别
        返回: 原有的大纲级别 (0-8) 或 None
        """
        if level < 1 or level > 9:
            self._log(f"  > 警告：大纲级别 {level} 超出范围 (1-9)，已跳过设置")
            return None

        # 读取原有大纲级别
        original_level = self._get_outline_level(para)

        # 设置新的大纲级别 (Word内部用0-8表示1-9级)
        pPr = para._p.get_or_add_pPr()
        outlineLvl = pPr.find(qn("w:outlineLvl"))
        if outlineLvl is None:
            outlineLvl = OxmlElement("w:outlineLvl")
            pPr.append(outlineLvl)
        outlineLvl.set(qn("w:val"), str(level - 1))

        return original_level

    def _format_heading(self, para, level):
        """
        为段落设置大纲级别（仅设置大纲级别，不影响其他格式）
        """
        if not self.config["set_outline"]:
            self._log("  > 大纲级别设置已禁用，跳过")
            return

        # 获取段落文本预览用于日志
        text_preview = para.text.strip()[:30].replace("\n", " ")

        original_level = self._set_outline_level(para, level)

        if original_level is not None:
            self._log(
                f'  > 大纲级别: Lv{original_level + 1} → Lv{level} (覆盖) - "{text_preview}..."'
            )
        else:
            self._log(f'  > 大纲级别: 无 → Lv{level} (新设) - "{text_preview}..."')

    def _apply_text_indent_and_align(self, para):
        pf = para.paragraph_format
        indent_unit = self._normalize_indent_unit(
            self.config.get("paragraph_indent_unit", "cm")
        )
        pf.first_line_indent = None
        pf.left_indent = None
        pf.right_indent = None

        ind = para._p.get_or_add_pPr().get_or_add_ind()
        ind.attrib.pop(qn("w:left"), None)
        ind.attrib.pop(qn("w:right"), None)
        ind.attrib.pop(qn("w:leftChars"), None)
        ind.attrib.pop(qn("w:rightChars"), None)
        ind.attrib.pop(qn("w:hanging"), None)
        ind.attrib.pop(qn("w:hangingChars"), None)
        ind.attrib.pop(qn("w:firstLine"), None)
        ind.attrib.pop(qn("w:firstLineChars"), None)

        if indent_unit == "chars":
            left_chars = self._config_float(self.config, "left_indent_chars", 0.0)
            right_chars = self._config_float(self.config, "right_indent_chars", 0.0)
            ind.set(qn("w:leftChars"), self._chars_to_word_value(left_chars))
            ind.set(qn("w:rightChars"), self._chars_to_word_value(right_chars))
        else:
            pf.left_indent = Cm(self._config_float(self.config, "left_indent_cm", 0.0))
            pf.right_indent = Cm(
                self._config_float(self.config, "right_indent_cm", 0.0)
            )

        first_line_chars = self._config_float(
            self.config, "first_line_indent_chars", 2.0
        )
        if first_line_chars > 0:
            ind.set(qn("w:firstLineChars"), self._chars_to_word_value(first_line_chars))

        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    @staticmethod
    def _config_float(config, key, default):
        value = config.get(key, default)
        try:
            if value == "":
                return default
            return float(value)
        except (TypeError, ValueError):
            return default

    @staticmethod
    def _normalize_spacing_unit(value):
        value = str(value or "pt").strip().lower()
        if value in ("multiple", "倍数", "multi"):
            return "multiple"
        return "pt"

    @staticmethod
    def _normalize_indent_unit(value):
        value = str(value or "cm").strip().lower()
        if value in ("chars", "char", "character", "characters", "字符"):
            return "chars"
        return "cm"

    @staticmethod
    def _chars_to_word_value(value):
        return str(int(round(max(0.0, float(value)) * 100)))

    def _apply_line_spacing(
        self, para, pt_key, unit_key, multiple_key, default_pt, default_multiple=1.0
    ):
        unit = self._normalize_spacing_unit(self.config.get(unit_key, "pt"))
        spacing = para._p.get_or_add_pPr().get_or_add_spacing()
        if unit == "multiple":
            multiple = self._config_float(self.config, multiple_key, default_multiple)
            if multiple <= 0:
                multiple = default_multiple
            spacing.set(qn("w:line"), str(int(round(multiple * 240))))
            spacing.set(qn("w:lineRule"), "auto")
            return

        line_spacing_pt = self._config_float(self.config, pt_key, default_pt)
        if line_spacing_pt <= 0:
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            spacing.attrib.pop(qn("w:line"), None)
            return
        spacing.set(qn("w:line"), str(int(round(line_spacing_pt * 20))))
        spacing.set(qn("w:lineRule"), "exact")

    @staticmethod
    def _apply_bold_to_runs(para, enabled):
        if not enabled:
            return
        for run in paragraph_runs(para):
            run.font.bold = True

    @staticmethod
    def _apply_bold_to_first_chars(para, char_count, enabled):
        if not enabled:
            return
        consumed = 0
        for run in paragraph_runs(para):
            run_len = len(run.text or "")
            if run_len and consumed < char_count:
                run.font.bold = True
            consumed += run_len
